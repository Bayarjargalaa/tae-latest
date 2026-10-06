"""
Shop app views - E-commerce frontend
"""
from django.shortcuts import render, redirect
from django.core.paginator import Paginator
from django.contrib.auth import login
from django.contrib.auth.decorators import login_required
from django.contrib.admin.views.decorators import staff_member_required
from django.http import JsonResponse
from django.urls import reverse
from django.db import IntegrityError
from django.db.models import Sum, Count, Avg, Q, FloatField, DateField, Min, Max, F, ExpressionWrapper
from django.db.models.functions import Cast, TruncDate, ExtractYear, ExtractMonth, ExtractDay
from django.utils import timezone
from datetime import timedelta, datetime
from decimal import Decimal, InvalidOperation
from datamigration.models import OpenDataItem, OpenDataCustomer, OpenDataEmployee, OpenDataSale, OpenDataSaleHeader
from shop.models import ProductExtension
from shop.forms import CustomUserCreationForm


def _attach_base_channel_prices(products):
    """Бүтээгдэхүүн бүрт '001' (Үндсэн үнэ) сувгийн хамгийн сүүлийн огнооны
    үнийг saleprice/baseprice дээр нь бичиж өгнө. Тухайн барааны энэ сувагт
    үнэ бүртгэгдээгүй бол OpenDataItem дээрх анхны утгыг хэвээр үлдээнэ.
    """
    from datamigration.models import OpenDataItemPrice, OpenDataItemPriceChannel

    products = list(products)
    pkids = [p.pkid for p in products if p.pkid]
    if not pkids:
        return products

    base_channel = OpenDataItemPriceChannel.objects.filter(code='001').first()
    if not base_channel:
        return products

    latest_prices = (
        OpenDataItemPrice.objects
        .filter(customerpkid__isnull=True, pricechannelpkid=base_channel.pkid, itempkid__in=pkids)
        .order_by('itempkid', '-pricedate')
        .distinct('itempkid')
    )
    price_map = {p.itempkid: p for p in latest_prices}

    for product in products:
        entry = price_map.get(product.pkid)
        if entry:
            if entry.price is not None:
                product.saleprice = entry.price
            if entry.baseprice is not None:
                product.baseprice = entry.baseprice

    return products


def home(request):
    """Нүүр хуудас - Featured бүтээгдэхүүнүүд"""
    # Ажилтан нэвтэрсэн бол шууд хувийн мэдээлэл рүү нь чиглүүлнэ
    if request.user.is_authenticated and request.user.is_staff:
        return redirect('shop:profile')
    
    # Идэвхтэй бүтээгдэхүүнүүд (эхний 12)
    products = OpenDataItem.objects.filter(activestatus='Y').order_by('-pkid')[:12]
    products = _attach_base_channel_prices(products)

    # ProductExtension холбох
    product_ids = [p.id for p in products]
    extensions = {ext.item_id: ext for ext in ProductExtension.objects.filter(item_id__in=product_ids)}
    
    # Extension нэмэх
    for product in products:
        product.extension = extensions.get(product.id)
    
    # Категориуд (unique)
    categories = OpenDataItem.objects.filter(activestatus='Y').values_list('itemcategoryname', flat=True).distinct()[:12]
    
    context = {
        'products': products,
        'categories': [cat for cat in categories if cat],
    }
    
    return render(request, 'shop/home.html', context)


def product_list(request):
    """Бүтээгдэхүүний жагсаалт (pagination, filter)"""
    products = OpenDataItem.objects.filter(activestatus='Y')
    
    # Id баганаас 8 болон 9-р эхэлсэн бараануудыг харуулах
    from django.db.models import Q
    products = products.filter(
        Q(id__istartswith='8') | Q(id__istartswith='9')
    )
    
    # Агуулах жижигт 10+ үлдэгдэлтэй бараа шүүлт
    # Нэвтрээгүй хэрэглэгчдэд default-оор идэвхтэй
    if request.user.is_authenticated:
        jijig_min_qty = request.GET.get('jijig_min', 'false')
    else:
        jijig_min_qty = request.GET.get('jijig_min', 'true')  # Нэвтрээгүй бол 10+ л харуулах
    
    if jijig_min_qty == 'true':
        from shop.models_inventory import InventorySnapshot
        # InventorySnapshot-с jijig_qty >= 10 байгаа барааны нэрүүдийг авах
        eligible_items = InventorySnapshot.objects.filter(jijig_qty__gte=10).values_list('itemname', flat=True)
        products = products.filter(name__in=eligible_items)
    
    # Category шүүлт
    category = request.GET.get('category')
    if category:
        products = products.filter(itemcategoryname=category)
    
    # Хайлт
    search_query = request.GET.get('q')
    if search_query:
        products = products.filter(name__icontains=search_query)
    
    # Brand шүүлт
    brand = request.GET.get('brand')
    if brand:
        products = products.filter(brandname=brand)
    
    # Pagination
    paginator = Paginator(products, 20)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    page_obj.object_list = _attach_base_channel_prices(page_obj.object_list)

    # ProductExtension холбох
    product_ids = [p.id for p in page_obj]
    extensions = {ext.item_id: ext for ext in ProductExtension.objects.filter(item_id__in=product_ids)}
    
    for product in page_obj:
        product.extension = extensions.get(product.id)
    
    # Filters
    categories = OpenDataItem.objects.filter(activestatus='Y').values_list('itemcategoryname', flat=True).distinct()
    brands = OpenDataItem.objects.filter(activestatus='Y').values_list('brandname', flat=True).distinct()
    
    context = {
        'page_obj': page_obj,
        'categories': [cat for cat in categories if cat],
        'brands': [brand for brand in brands if brand],
        'current_category': category,
        'current_brand': brand,
        'search_query': search_query,
        'jijig_min_qty': jijig_min_qty,
    }
    
    return render(request, 'shop/products.html', context)


def product_detail(request, pk):
    """Бүтээгдэхүүний дэлгэрэнгүй"""
    from django.shortcuts import get_object_or_404
    
    product = get_object_or_404(OpenDataItem, pk=pk)
    _attach_base_channel_prices([product])

    # Extension холбох
    try:
        extension = ProductExtension.objects.get(item_id=product.id)
    except ProductExtension.DoesNotExist:
        extension = None

    product.extension = extension

    # Төстэй бүтээгдэхүүнүүд
    related = OpenDataItem.objects.filter(
        itemcategoryname=product.itemcategoryname,
        activestatus='Y'
    ).exclude(pk=pk)[:4]
    related = _attach_base_channel_prices(related)

    context = {
        'product': product,
        'related_products': related,
    }
    
    return render(request, 'shop/product_detail.html', context)


def product_edit(request, pk):
    """Бүтээгдэхүүний нэмэлт мэдээлэл засах (staff only)"""
    from django.shortcuts import get_object_or_404, redirect
    from django.contrib.admin.views.decorators import staff_member_required
    
    # Staff эрх шалгах
    if not request.user.is_staff:
        return redirect('shop:products')
    
    product = get_object_or_404(OpenDataItem, pk=pk)
    
    # Extension авах эсвэл үүсгэх
    extension, created = ProductExtension.objects.get_or_create(item_id=product.id)
    
    if request.method == 'POST':
        # Form өгөгдөл хадгалах
        extension.name_mn = request.POST.get('name_mn', '').strip()
        extension.country_of_origin = request.POST.get('country_of_origin', '').strip()
        extension.supplier = request.POST.get('supplier', '').strip()
        
        qty = request.POST.get('qty_per_box', '').strip()
        extension.qty_per_box = int(qty) if qty else None
        
        # Зураг хадгалах
        if 'image_main' in request.FILES:
            extension.image_main = request.FILES['image_main']
        
        extension.save()
        
        return redirect('shop:products')
    
    context = {
        'product': product,
        'extension': extension,
    }
    
    return render(request, 'shop/product_edit.html', context)


def register(request):
    """Бүртгүүлэх - OpenDataEmployee имэйл шалгана"""
    if request.method == 'POST':
        form = CustomUserCreationForm(request.POST)
        if form.is_valid():
            user = form.save()
            login(request, user)
            
            # Ажилтан эсэхийг мэдэгдэх
            if user.is_staff:
                from django.contrib import messages
                messages.success(request, f'Тавтай морилно уу, {user.username}! Та ажилтны эрхтэй нэвтэрлээ.')
            else:
                from django.contrib import messages
                messages.info(request, f'Амжилттай бүртгүүллээ, {user.username}!')
            
            return redirect('shop:home')
    else:
        form = CustomUserCreationForm()
    
    return render(request, 'shop/register.html', {'form': form})


@login_required
def dashboard(request):
    """Ажилтны хяналтын самбар (staff only)"""
    # Зөвхөн ажилтанд зөвшөөрөх
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')
    
    # Статистик - "15"-аар эхэлсэн код нь орцны түүхий эд тул жинхэнэ
    # бүтээгдэхүүний тооноос хасна
    from shop.models_inventory import InventorySnapshot

    products_qs = OpenDataItem.objects.exclude(id__istartswith='15')
    total_products = products_qs.count()
    active_products = products_qs.filter(activestatus='Y').count()
    in_stock_products = InventorySnapshot.objects.filter(total_qty__gt=0).count()
    total_customers = OpenDataCustomer.objects.count()
    total_employees = OpenDataEmployee.objects.count()

    context = {
        'total_products': total_products,
        'active_products': active_products,
        'in_stock_products': in_stock_products,
        'total_customers': total_customers,
        'total_employees': total_employees,
    }

    return render(request, 'shop/dashboard.html', context)


SALES_REPORT_TABS = ['summary', 'product', 'customer', 'channel', 'group', 'seller', 'warehouse', 'daily']


def _sales_report_filters(request):
    """Борлуулалтын тайлангийн шүүлт (бүх tab-д нэгэн адил). Огноо заагаагүй бол бүх хугацаа."""
    from shop.services.sales_report import MULTI_FILTERS

    def parse_date(name):
        try:
            return datetime.strptime(request.GET.get(name, ''), '%Y-%m-%d').date()
        except ValueError:
            return None

    filters = {'date_from': parse_date('date_from'), 'date_to': parse_date('date_to')}
    if filters['date_from'] and filters['date_to'] and filters['date_from'] > filters['date_to']:
        filters['date_from'], filters['date_to'] = filters['date_to'], filters['date_from']
    for key in MULTI_FILTERS:
        filters[key] = sorted({v.strip() for v in request.GET.getlist(key) if v.strip()})
    return filters


@login_required
def sales_report(request):
    """Борлуулалтын тайлан - OpenDataSale: нэгтгэл (он-сар), бараа/харилцагч/суваг/бүлэг/борлуулагч/агуулахаар
    задаргаа, өдрөөр. Бүх шүүлт бүх tab-д үйлчилнэ; үр дүн кэшлэгдэнэ (shop.services.sales_report-ийг үз)."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from shop.services import sales_report as sr

    tab = request.GET.get('tab', 'summary')
    tab = {'by_product': 'product', 'by_customer': 'customer', 'by_date': 'daily'}.get(tab, tab)  # хуучин холбоос
    if tab not in SALES_REPORT_TABS:
        tab = 'summary'
    filters = _sales_report_filters(request)

    data = {}
    if tab == 'summary':
        data['summary'] = sr.summary(filters)
    elif tab == 'daily':
        data['daily'] = sr.daily(filters)
    else:
        data['rows'] = sr.breakdown(filters, tab)

    options = sr.filter_options()
    query = request.GET.copy()
    query.pop('tab', None)
    labels = {'summary': ('Нэгтгэл', '📈'), 'daily': ('Өдрөөр', '📅')}
    labels.update({k: (v['label'], v['icon']) for k, v in sr.DIMENSIONS.items()})
    context = {
        'active_page': 'sales_report',
        'tab': tab,
        'tabs': [(k, *labels[k]) for k in SALES_REPORT_TABS],
        'dim': sr.DIMENSIONS.get(tab),
        'filters': filters,
        'options': options,
        'selected_customers': sr.customer_names(filters['customer']),
        'kpis': sr.kpis(filters),
        'data': data,
        'query_string': query.urlencode(),
        'is_filtered': any(filters[k] for k in sr.MULTI_FILTERS) or filters['date_from'] or filters['date_to'],
    }
    return render(request, 'shop/sales_report.html', context)


@login_required
def sales_report_customers(request):
    """Борлуулалтын тайлангийн харилцагчийн шүүлтийн хайлт (JSON)."""
    from django.http import JsonResponse
    if not request.user.is_staff:
        return JsonResponse({'results': []}, status=403)
    from shop.services.sales_report import sales_customers, search_customers
    q = request.GET.get('q', '').strip()
    # q-гүй бол борлуулалттай бүх харилцагч (шүүлтийн жагсаалтад нэг удаа ачаалагдана, кэштэй)
    results = search_customers(q) if q else sales_customers()
    # Нэрийн илүү зай, мөр шилжилтийг цэвэрлэнэ; кирилл үсгийг escape хийхгүй (хэмжээ 2 дахин бага)
    results = [{'id': r['id'], 'name': ' '.join((r['name'] or '').split())} for r in results]
    return JsonResponse({'results': results}, json_dumps_params={'ensure_ascii': False})


@login_required
def sales_report_export(request):
    """Борлуулалтын тайлангийн идэвхтэй tab-ыг шүүлтээр нь Excel (.xlsx) файлаар татна."""
    if not request.user.is_staff:
        return redirect('shop:home')

    import openpyxl
    from io import BytesIO
    from django.http import HttpResponse
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from shop.services import sales_report as sr

    tab = request.GET.get('tab', 'summary')
    if tab not in SALES_REPORT_TABS:
        tab = 'summary'
    filters = _sales_report_filters(request)
    metric_cols = [('revenue', 'Орлого'), ('nonvat', 'НӨАТ-гүй'), ('cost', 'Өртөг'), ('profit', 'Ашиг'), ('margin', 'Ашгийн %'),
                   ('orders', 'Захиалга'), ('customers', 'Харилцагч'), ('qty', 'Тоо ширхэг')]
    if tab == 'summary':
        headers = ['Он', 'Сар'] + [l for _, l in metric_cols]
        rows = []
        summary = sr.summary(filters)
        for y in summary['years']:
            for m in y['months']:
                rows.append([y['year'], m['month']] + [m[k] for k, _ in metric_cols])
            rows.append([y['year'], 'Нийт'] + [y['total'][k] for k, _ in metric_cols])
        if summary['total']:
            rows.append(['Бүгд', ''] + [summary['total'][k] for k, _ in metric_cols])
    elif tab == 'daily':
        headers = ['Огноо'] + [l for _, l in metric_cols]
        rows = [[r['date']] + [r[k] for k, _ in metric_cols] for r in sr.daily(filters)]
    else:
        cfg = sr.DIMENSIONS[tab]
        extra = list(cfg['extra'])
        headers = ['Код', 'Нэр'] + extra + [l for _, l in metric_cols] + [cfg['count_label'], 'Эзлэх %']
        rows = [[r['id'], r['name']] + [r.get(e, '') for e in extra] + [r[k] for k, _ in metric_cols] + [r['count'], r['share']]
                for r in sr.breakdown(filters, tab)]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Борлуулалт'
    period = f"{filters['date_from'] or 'эхнээс'} ~ {filters['date_to'] or 'одоог хүртэл'}"
    ws.cell(row=1, column=1, value=f'Борлуулалтын тайлан · {period}').font = Font(bold=True, size=13)
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = 40 if title == 'Нэр' else 15
    for i, values in enumerate(rows, start=4):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            if isinstance(value, float):
                cell.number_format = '0.0' if headers[col - 1] in ('Ашгийн %', 'Эзлэх %') else '#,##0'
    ws.freeze_panes = 'A4'
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="sales_{tab}.xlsx"'
    return response


def _get_available_years(sales_qs):
    """Боломжит жилүүдийг авах"""
    years = sales_qs.extra(
        select={'year': "EXTRACT(year FROM CAST(\"DocumentDate\" AS DATE))"}
    ).values_list('year', flat=True).distinct().order_by('-year')
    
    return [int(y) for y in years if y]


def inventory_report(request):
    """Барааны үлдэгдэл - InventorySnapshot хүснэгтээс (2 цаг тутамд шинэчлэгддэг)"""
    # Зөвхөн ажилтанд зөвшөөрөх
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')
    
    from shop.models_inventory import InventorySnapshot
    from django.db.models import Q
    from django.core.paginator import Paginator
    
    # Шүүлтийн параметрүүд
    warehouse_filter = request.GET.get('warehouse')
    product_filter = request.GET.get('product')
    brand_filter = request.GET.get('brand')
    show_stock_only = request.GET.get('stock_only', 'true')  # Default: үлдэгдэлтэй л харуулах
    jijig_min_qty = request.GET.get('jijig_min', 'false')  # Агуулах жижигт 10+ үлдэгдэлтэй
    
    # "None" болон хоосон утгуудыг цэвэрлэх
    def clean_param(value):
        if not value or value.strip() == '' or value.strip().lower() == 'none':
            return None
        return value.strip()
    
    warehouse_filter = clean_param(warehouse_filter)
    product_filter = clean_param(product_filter)
    brand_filter = clean_param(brand_filter)
    
    # Snapshot-с өгөгдөл авах (их хурдан!)
    inventory_qs = InventorySnapshot.objects.all()
    
    # Үлдэгдэлтэй барааг л харуулах эсэх
    if show_stock_only == 'true':
        inventory_qs = inventory_qs.filter(total_qty__gt=0)
    
    # Агуулах жижигт 10+ үлдэгдэлтэй шүүлт
    if jijig_min_qty == 'true':
        inventory_qs = inventory_qs.filter(jijig_qty__gte=10)
    
    # Шүүлт хэрэглэх
    if product_filter:
        inventory_qs = inventory_qs.filter(itemname__icontains=product_filter)
    
    if brand_filter:
        inventory_qs = inventory_qs.filter(brandname__icontains=brand_filter)
    
    # Нэрээр эрэмбэлэх (A → Z)
    inventory_qs = inventory_qs.order_by('itemname')
    
    # Агуулахаар шүүх (үлдэгдэлтэй агуулахаар)
    if warehouse_filter:
        if warehouse_filter == 'Агуулах Толгойт':
            inventory_qs = inventory_qs.filter(tolgoit_qty__gt=0)
        elif warehouse_filter == 'Агуулах жижиг':
            inventory_qs = inventory_qs.filter(jijig_qty__gt=0)
        elif warehouse_filter == 'Агуулах алтжин':
            inventory_qs = inventory_qs.filter(altjin_qty__gt=0)
    
    # Pagination (50 items per page)
    paginator = Paginator(inventory_qs, 50)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)
    
    # Dropdown жагсаалтууд
    target_warehouses = ['Агуулах Толгойт', 'Агуулах жижиг', 'Агуулах алтжин']
    
    # Dropdown-д хэт олон зүйл ачаалахгүйн тулд limit
    all_snapshots = InventorySnapshot.objects.all()
    if product_filter:
        all_snapshots = all_snapshots.filter(itemname__icontains=product_filter)
    if brand_filter:
        all_snapshots = all_snapshots.filter(brandname__icontains=brand_filter)
    
    products_list = all_snapshots.values_list('itemname', flat=True).distinct().order_by('itemname')[:500]
    brands_list = all_snapshots.values_list('brandname', flat=True).distinct().exclude(
        brandname__isnull=True
    ).exclude(brandname='').order_by('brandname')[:100]
    
    # Template-д илгээх өгөгдлийг бэлтгэх
    inventory_data = []
    for snapshot in page_obj:
        inventory_data.append({
            'itemname': snapshot.itemname,
            'brandname': snapshot.brandname or '-',
            'warehouses': {
                'Агуулах Толгойт': {
                    'endqty': float(snapshot.tolgoit_qty),
                    'date': snapshot.tolgoit_date
                },
                'Агуулах жижиг': {
                    'endqty': float(snapshot.jijig_qty),
                    'date': snapshot.jijig_date
                },
                'Агуулах алтжин': {
                    'endqty': float(snapshot.altjin_qty),
                    'date': snapshot.altjin_date
                }
            },
            'total_qty': float(snapshot.total_qty)
        })
    
    # Сүүлийн шинэчлэлтийн огноо
    latest_update = InventorySnapshot.objects.order_by('-updated_at').first()
    last_updated = latest_update.updated_at if latest_update else None
    
    context = {
        'inventory_data': inventory_data,
        'page_obj': page_obj,
        'target_warehouses': target_warehouses,
        'warehouses_list': target_warehouses,
        'products_list': [{'name': p} for p in products_list],
        'brands_list': list(brands_list),
        'selected_warehouse': warehouse_filter,
        'selected_product': product_filter,
        'selected_brand': brand_filter,
        'show_stock_only': show_stock_only,  # Үлдэгдэлтэй барааг л харуулах эсэх
        'jijig_min_qty': jijig_min_qty,  # Агуулах жижигт 10+ үлдэгдэл
        'total_items': inventory_qs.count(),
        'last_updated': last_updated,  # Сүүлийн шинэчлэлт
    }
    
    return render(request, 'shop/inventory.html', context)


def _get_item_price_data(request):
    """Барааны үнийн хуудасны шүүлт/тооцооллыг хийнэ. HTML харагдац болон Excel
    экспорт хоёулаа энэ функцийг ашиглаж, ижил дүрмээр өгөгдлөө бэлдэнэ."""
    from datamigration.models import OpenDataItem, OpenDataItemPrice, OpenDataItemPriceChannel

    search_query = request.GET.get('q', '').strip()
    stock_only = request.GET.get('stock_only') == 'true'
    below_base_only = request.GET.get('below_base_only') == 'true'
    show_all = request.GET.get('show_all') == 'true'
    selected_channel_pkids = [c for c in request.GET.getlist('channel') if c]

    # Суваг тус бүрийг кодоор нь эрэмбэлж авна (жиш: '001' = Үндсэн үнэ эхэнд орно)
    all_channels = list(OpenDataItemPriceChannel.objects.all().order_by('code'))
    if selected_channel_pkids:
        channels = [c for c in all_channels if str(c.pkid) in selected_channel_pkids]
    else:
        channels = all_channels
    channel_pkids = [c.pkid for c in channels]

    base_channel = next((c for c in all_channels if c.code == '001'), None)

    # Item+суваг хослол бүрийн хамгийн сүүлийн огнооны үнийг нэг query-ээр авна
    # (Postgres-ийн DISTINCT ON). Хэрэглэгч тусгайлсан үнэ (customerpkid) биш,
    # ерөнхий сувгийн үнийг л авна. Зөвхөн сонгосон (эсвэл бүх) сувгаар шүүнэ.
    latest_prices = (
        OpenDataItemPrice.objects
        .filter(customerpkid__isnull=True, pricechannelpkid__in=channel_pkids)
        .order_by('itempkid', 'pricechannelpkid', '-pricedate')
        .distinct('itempkid', 'pricechannelpkid')
    )

    price_map = {}
    item_pks_with_price = set()
    for p in latest_prices:
        price_map.setdefault(p.itempkid, {})[p.pricechannelpkid] = p.price
        item_pks_with_price.add(p.itempkid)

    # 001 (Үндсэн үнэ) сувгийн үнийг харьцуулалтад ашиглах зорилгоор, тухайн
    # суваг харагдаж буй баганад ороогүй ч гэсэн тусад нь авна
    base_price_map = {}
    if base_channel:
        base_prices_qs = (
            OpenDataItemPrice.objects
            .filter(customerpkid__isnull=True, pricechannelpkid=base_channel.pkid)
            .order_by('itempkid', '-pricedate')
            .distinct('itempkid')
        )
        base_price_map = {p.itempkid: p.price for p in base_prices_qs}

    def is_cell_below_base(ch, value, base_price):
        return (
            base_price is not None and value is not None
            and (not base_channel or ch.pkid != base_channel.pkid)
            and value < base_price
        )

    # Аль барааны ядаж нэг сувагт үнэ нь үндсэн үнээс доогуур болохыг тодорхойлно
    below_base_item_pks = set()
    for itempkid, channel_prices in price_map.items():
        base_price = base_price_map.get(itempkid)
        if base_price is None:
            continue
        for ch in channels:
            if is_cell_below_base(ch, channel_prices.get(ch.pkid), base_price):
                below_base_item_pks.add(itempkid)
                break

    items_qs = OpenDataItem.objects.filter(pkid__in=item_pks_with_price)
    if search_query:
        items_qs = items_qs.filter(name__icontains=search_query)
    if stock_only:
        from shop.models_inventory import InventorySnapshot
        in_stock_names = InventorySnapshot.objects.filter(total_qty__gt=0).values_list('itemname', flat=True)
        items_qs = items_qs.filter(name__in=in_stock_names)
    if below_base_only:
        items_qs = items_qs.filter(pkid__in=below_base_item_pks)
    items_qs = items_qs.order_by('name')

    def build_rows(items_iterable):
        rows = []
        for item in items_iterable:
            item_prices_by_channel = price_map.get(item.pkid, {})
            base_price = base_price_map.get(item.pkid)
            cells = []
            for ch in channels:
                value = item_prices_by_channel.get(ch.pkid)
                cells.append({'value': value, 'is_below_base': is_cell_below_base(ch, value, base_price)})
            rows.append({'item': item, 'cells': cells, 'base_price': base_price})
        return rows

    return {
        'search_query': search_query,
        'stock_only': stock_only,
        'below_base_only': below_base_only,
        'show_all': show_all,
        'selected_channel_pkids': selected_channel_pkids,
        'all_channels': all_channels,
        'channels': channels,
        'items_qs': items_qs,
        'build_rows': build_rows,
    }


def item_prices(request):
    """Барааны үнэ - OpenDataItemPrice + OpenDataItemPriceChannel-ээс.

    Мөр = бараа, багана = үнийн жагсаалт/суваг (кодоор эрэмбэлсэн), утга =
    тухайн бараа/сувгийн хамгийн сүүлийн огнооны үнэ.
    """
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    data = _get_item_price_data(request)
    items_qs = data['items_qs']

    page_obj = None
    if data['show_all']:
        display_items = items_qs
    else:
        from django.core.paginator import Paginator
        paginator = Paginator(items_qs, 50)
        page_obj = paginator.get_page(request.GET.get('page', 1))
        display_items = page_obj

    rows = data['build_rows'](display_items)

    querystring = request.GET.copy()
    querystring.pop('page', None)

    context = {
        'all_channels': data['all_channels'],
        'channels': data['channels'],
        'selected_channel_pkids': data['selected_channel_pkids'],
        'stock_only': data['stock_only'],
        'below_base_only': data['below_base_only'],
        'show_all': data['show_all'],
        'rows': rows,
        'page_obj': page_obj,
        'search_query': data['search_query'],
        'total_items': items_qs.count(),
        'querystring': querystring.urlencode(),
    }
    return render(request, 'shop/item_prices.html', context)


def item_prices_export(request):
    """Барааны үнийн хуудасны идэвхтэй шүүлтийг Excel (.xlsx) файлаар татаж авна.
    Үндсэн үнээс доогуур мөрүүдийг улаанаар тэмдэглэнэ."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from io import BytesIO
    from django.http import HttpResponse

    data = _get_item_price_data(request)
    channels = data['channels']
    rows = data['build_rows'](data['items_qs'])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Барааны үнэ'

    header_font = Font(bold=True)
    red_fill = PatternFill(start_color='FFCDD2', end_color='FFCDD2', fill_type='solid')

    ws.cell(row=1, column=1, value='Бараа').font = header_font
    for col_idx, ch in enumerate(channels, start=2):
        cell = ws.cell(row=1, column=col_idx, value=f"{ch.code} - {ch.name}")
        cell.font = header_font
        cell.alignment = Alignment(horizontal='right', wrap_text=True)

    for row_idx, row in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=row['item'].name)
        for col_idx, cell_data in enumerate(row['cells'], start=2):
            cell = ws.cell(row=row_idx, column=col_idx, value=cell_data['value'])
            if cell_data['is_below_base']:
                cell.fill = red_fill

    ws.column_dimensions['A'].width = 42
    for col_idx in range(2, len(channels) + 2):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14
    ws.freeze_panes = 'B2'

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="item_prices.xlsx"'
    return response


@login_required
def item_costs(request):
    """Барааны өртөг - татан авалтын өртөг (дараагийн татан авалт хүртэл) ба дундаж өртөг (OpenDataInventory)."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from shop.services import item_costs as service

    filters = service.parse_filters(request.GET)
    context = {
        'filters': filters,
        'date_from': filters['date_from'].isoformat(),
        'date_to': filters['date_to'].isoformat(),
        'item_options': service.get_item_options(),
        'costs': service.build_costs(filters),
        'querystring': request.GET.urlencode(),
    }
    return render(request, 'shop/item_costs.html', context)


@login_required
def item_costs_export(request):
    """Барааны өртөг - идэвхтэй шүүлтийн үр дүнг Excel-ээр: "Өртөг" (бараа тус бүр) ба
    "Өөрчлөлтийн түүх" (өртөг өөрчлөгдсөн өдрүүд, татан авалтууд) гэсэн 2 хуудас."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    import openpyxl
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    from io import BytesIO
    from django.http import HttpResponse
    from shop.services import item_costs as service

    filters = service.parse_filters(request.GET)
    rows = service.build_costs(filters)['rows']
    period = f"{filters['date_from']:%Y-%m-%d} ~ {filters['date_to']:%Y-%m-%d}"
    bold = Font(bold=True)
    money, pct = '#,##0.00', '0.0%'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Өртөг'
    ws.cell(row=1, column=1, value=f'Барааны өртөг: {period}').font = Font(bold=True, size=12)
    headers = [
        ('Код', 10, None), ('Бараа', 48, None), ('Нэгж', 10, None),
        ('Татан авалтын өртөг (НӨАТ-тэй)', 16, money), ('Татан авалтын өртөг (НӨАТ-гүй)', 16, money),
        ('Сүүлийн татан авалт', 14, None), ('Баримт', 14, None),
        ('Дундаж өртөг', 16, money), ('Дундаж − татан авалт (НӨАТ-гүй)', 12, pct),
        ('Эхэн дэх дундаж өртөг', 16, money), ('Дундаж өртгийн өөрчлөлт', 12, pct),
        ('Эхэн дэх татан авалтын өртөг (НӨАТ-тэй)', 16, money), ('Эхэн дэх татан авалтын өртөг (НӨАТ-гүй)', 16, money),
        ('Татан авалтын өртгийн өөрчлөлт', 12, pct),
        ('Хугацаанд татан авалт', 10, None), ('Одоогийн үлдэгдэл', 12, '#,##0.##'), ('Үлдэгдлийн өртөг', 16, '#,##0'),
    ]
    for col, (title, width, _) in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=title)
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[get_column_letter(col)].width = width
    for row_idx, r in enumerate(rows, start=4):
        values = [
            r['id'], r['name'], r['measure'], r['landed_end'], r['landed_end_novat'],
            r['landed_end_date'], r['landed_end_numbers'],
            r['avg_end'], r['diff'], r['avg_start'], r['avg_change'],
            r['landed_start'], r['landed_start_novat'], r['landed_change'],
            r['landed_count'], r['stock_qty'], r['stock_value'],
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            if headers[col - 1][2]:
                cell.number_format = headers[col - 1][2]
    ws.freeze_panes = 'C4'

    hist = wb.create_sheet('Өөрчлөлтийн түүх')
    hist_headers = [('Код', 10), ('Бараа', 48), ('Огноо', 12), ('Дундаж өртөг', 16),
                    ('Татан авалтын өртөг (НӨАТ-тэй)', 16), ('Татан авалтын өртөг (НӨАТ-гүй)', 16), ('Татан авалт', 16)]
    for col, (title, width) in enumerate(hist_headers, start=1):
        cell = hist.cell(row=1, column=col, value=title)
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        hist.column_dimensions[get_column_letter(col)].width = width
    row_idx = 2
    for r in rows:
        for day, avg, landed, landed_novat, note in r['history']:
            for col, value in enumerate([r['id'], r['name'], day, avg, landed, landed_novat, note or None], start=1):
                cell = hist.cell(row=row_idx, column=col, value=value)
                if col in (4, 5, 6):
                    cell.number_format = money
            row_idx += 1
    hist.freeze_panes = 'C2'

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        f'attachment; filename="item_costs_{filters["date_from"]:%Y%m%d}_{filters["date_to"]:%Y%m%d}.xlsx"'
    )
    return response


@login_required
def sales_report_ht(request):
    """Борлуулалтын тайлан ХТ-р - мөр: бараа/харилцагч (аль нь ч дээд түвшинд), багана: он-сар."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from shop.services import sales_ht

    filters, scope = _sales_ht_filters(request)
    context = _sales_ht_filter_context(request, filters, scope)
    from shop.services.report_cache import cached
    context['report'] = cached('ht_report', ['OpenDataSale', 'OpenDataCustomer', 'OpenDataDistributionChannel', 'OpenDataCustomerGroup'], filters, lambda: sales_ht.build_report(filters))
    return render(request, 'shop/sales_report_ht.html', context)


@login_required
def sales_report_ht_dashboard(request):
    """Борлуулалтын тайлан ХТ-р - Dashboard. Хүснэгттэй ижил шүүлттэй; үзүүлэлт, хэсгүүдийг хөтөч дээр сонгоно."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from shop.services import sales_ht
    from shop.services.sales_ht_dashboard import build_dashboard

    filters, scope = _sales_ht_filters(request)
    context = _sales_ht_filter_context(request, filters, scope)
    from shop.services.report_cache import cached
    context['dashboard'] = cached('ht_dashboard', ['OpenDataSale', 'OpenDataCustomer', 'OpenDataDistributionChannel', 'OpenDataCustomerGroup'], filters, lambda: build_dashboard(filters))
    return render(request, 'shop/sales_report_ht_dashboard.html', context)


@login_required
def sales_report_ht_breakdown(request):
    """Борлуулалтын тайлан ХТ-р - Задаргаа: мөр нь ХТ/харилцагч/бараа г.м. задардаг, багана нь сонгох боломжтой
    үзүүлэлтүүд. Огноо заагаагүй бол энэ сараар."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from shop.services import sales_ht
    from shop.services.sales_ht_breakdown import build_breakdown

    params = request.GET.copy()
    if not params.get('date_from') and not params.get('date_to'):
        today = timezone.localdate()
        params['date_from'] = today.replace(day=1).isoformat()
        params['date_to'] = today.isoformat()
    filters, scope = _sales_ht_filters(request, params)
    context = _sales_ht_filter_context(request, filters, scope)
    from shop.services.report_cache import cached
    context['breakdown'] = cached('ht_breakdown', ['OpenDataSale', 'OpenDataCustomer', 'OpenDataDistributionChannel', 'OpenDataCustomerGroup'], filters, lambda: build_breakdown(filters))
    return render(request, 'shop/sales_report_ht_breakdown.html', context)


def _sales_ht_filters(request, params=None):
    """ХТ тайлангийн шүүлт + нэвтэрсэн ХТ-ийн сувгийн хязгаарлалт. Бүх ХТ хуудас/экспорт үүгээр л шүүлтээ авна."""
    from shop.services import sales_ht
    from shop.services.sales_ht_access import apply_channel_scope, get_channel_scope

    filters = sales_ht.parse_filters(params if params is not None else request.GET)
    scope = get_channel_scope(request.user)
    return apply_channel_scope(filters, scope), scope


def _sales_ht_filter_context(request, filters, scope=None):
    """ХТ тайлангийн хүснэгт/dashboard/задаргааны нийтлэг шүүлтийн context (_sales_ht_filters.html).
    ХТ-д (scope) суваг, харилцагчийн сонголтыг зөвхөн түүний сувгаар хязгаарлана."""
    from shop.services import sales_ht

    from shop.services.report_cache import cached
    channel_ids = scope.channel_ids if scope else None
    options = cached('ht_options', ['OpenDataSale', 'OpenDataCustomer', 'OpenDataCustomerGroup'],
                     {'channels': sorted(channel_ids) if channel_ids is not None else None},
                     lambda: sales_ht.get_filter_options(channel_ids=channel_ids))
    channel_options = cached('ht_channels', ['OpenDataDistributionChannel'], {}, sales_ht.get_channels)
    if scope:
        channel_options = [c for c in channel_options if c[0] in scope.channel_ids]
    return {
        'filters': filters,
        'date_from': filters['date_from'].isoformat(),
        'date_to': filters['date_to'].isoformat(),
        'item_options': options['items'],
        'customer_options': options['customers'],
        # Модны түвшнээр догол мөр гаргана (Choices.js эхний хоосон зайг хасдаг тул "— " ашиглав)
        'customer_group_options': [(name, '— ' * depth + name) for name, depth in options['customer_groups']],
        'sales_mode_labels': sales_ht.SALES_MODE_LABELS,
        'item_mode_labels': sales_ht.ITEM_MODE_LABELS,
        'customer_status_labels': sales_ht.CUSTOMER_STATUS_LABELS,
        'channel_options': channel_options,
        'ht_scope': scope,
        'group_labels': sales_ht.GROUP_LABELS,
        'metric_labels': sales_ht.METRIC_LABELS,
        'compare_labels': sales_ht.COMPARE_LABELS,
        'channel_scope_labels': sales_ht.CHANNEL_SCOPE_LABELS,
        'querystring': request.GET.urlencode(),
    }


@login_required
def sales_report_ht_export(request):
    """Борлуулалтын тайлан ХТ-р - идэвхтэй шүүлтийн үр дүнг Excel (.xlsx) файлаар татна.
    Доод түвшний мөрүүдийг Excel-ийн outline бүлэглэлтээр (+/-) нээж/хааж болно."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from io import BytesIO
    from django.http import HttpResponse
    from shop.services import sales_ht

    filters, _scope = _sales_ht_filters(request)
    report = sales_ht.build_report(filters)
    rows = sales_ht.pivot_rows(
        report, filters['group_by'], filters['metric'],
        sort_key=request.GET.get('sort', 'name'), sort_desc=request.GET.get('dir') == 'desc',
    )
    top_is_item = filters['group_by'] == sales_ht.GROUP_ITEM

    stock_headers = []
    if report['show_jijig']:
        stock_headers.append(('Үлдэгдэл жижиг', 0))
    if report['show_tolgoit']:
        stock_headers.append(('Үлдэгдэл толгойт', 1))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Борлуулалт ХТ'
    ws.sheet_properties.outlinePr.summaryBelow = False

    bold = Font(bold=True)
    group_fill = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
    number_format = '#,##0' if filters['metric'] == sales_ht.METRIC_AMOUNT else '#,##0.##'

    ws.cell(row=1, column=1, value=(
        f"Борлуулалтын тайлан ХТ-р ({sales_ht.METRIC_LABELS[filters['metric']]}): "
        f"{filters['date_from']:%Y-%m-%d} ~ {filters['date_to']:%Y-%m-%d}"
    )).font = Font(bold=True, size=12)

    header_row = 3
    top_label = sales_ht.GROUP_LABELS[filters['group_by']]
    sub_label = sales_ht.GROUP_LABELS[sales_ht.GROUP_CUSTOMER if top_is_item else sales_ht.GROUP_ITEM]
    # Мөр дугаарлах үед эхний багана нь "№", нэрийн багана нэгээр баруун тийш шилжинэ
    name_col = 2 if filters['numbering'] else 1
    group_col = name_col + 1  # харилцагчийн мөрөнд OpenDataCustomer-ын бүлэг
    value_col = group_col + 1
    headers = (['№'] if filters['numbering'] else []) + [f'{top_label} / {sub_label}', 'Харилцагчийн бүлэг'] \
        + report['months'] + ['Нийт'] + [label for label, _ in stock_headers]
    for col_idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=text)
        cell.font = bold
        cell.alignment = Alignment(horizontal='center' if col_idx >= value_col else 'left', wrap_text=True)

    def write_row(row_idx, number, row, is_group):
        if filters['numbering']:
            ws.cell(row=row_idx, column=1, value=number)
        ws.cell(row=row_idx, column=name_col, value=row['name']).alignment = Alignment(indent=0 if is_group else 2)
        ws.cell(row=row_idx, column=group_col, value=row['group'] or None)
        cells = row['values'] + [row['total']]
        if stock_headers:
            cells += [(row['stock'][pos] if row['stock'] else None) for _, pos in stock_headers]
        for col_idx, value in enumerate(cells, start=value_col):
            cell = ws.cell(row=row_idx, column=col_idx, value=value or None)
            cell.number_format = number_format
        if is_group:
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).font = bold
                ws.cell(row=row_idx, column=col_idx).fill = group_fill

    row_idx = header_row + 1
    month_count = len(report['months'])
    grand = [0.0] * month_count
    for group_no, group in enumerate(rows, start=1):
        write_row(row_idx, group_no, group, True)
        row_idx += 1
        for child_no, child in enumerate(group['children'], start=1):
            write_row(row_idx, f'{group_no}.{child_no}', child, False)
            ws.row_dimensions[row_idx].outlineLevel = 1
            ws.row_dimensions[row_idx].hidden = True
            row_idx += 1
        for i, v in enumerate(group['values']):
            grand[i] += v

    ws.cell(row=row_idx, column=name_col, value='НИЙТ').font = bold
    for col_idx, value in enumerate(grand + [sum(grand)], start=value_col):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = bold
        cell.number_format = number_format

    if filters['numbering']:
        ws.column_dimensions['A'].width = 8
    ws.column_dimensions[get_column_letter(name_col)].width = 48
    ws.column_dimensions[get_column_letter(group_col)].width = 26
    for col_idx in range(value_col, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14
    ws.freeze_panes = ws.cell(row=header_row + 1, column=value_col)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        f'attachment; filename="sales_ht_{filters["date_from"]:%Y%m%d}_{filters["date_to"]:%Y%m%d}.xlsx"'
    )
    return response


def purchase_report(request):
    """Татан авалтын тайлан - OpenDataLandedCost өгөгдлөөс (Pivot: Он-Сар-Харилцагч-Бараа-Өдөр)"""
    # Зөвхөн ажилтанд зөвшөөрөх
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')
    
    from datamigration.models import OpenDataLandedCost
    
    # Tab сонголт (summary, by_vendor, by_item, by_date)
    active_tab = request.GET.get('tab', 'summary')
    
    # Шүүлтийн параметрүүд
    date_from_str = request.GET.get('date_from')
    date_to_str = request.GET.get('date_to')
    vendor_filter = request.GET.get('vendor')
    item_filter = request.GET.get('item')
    brand_filter = request.GET.get('brand')
    year_filter = request.GET.get('year')
    
    # "None" string болон хоосон утгуудыг цэвэрлэх
    def clean_param(value):
        if not value or value.strip() == '' or value.strip().lower() == 'none':
            return None
        return value.strip()
    
    vendor_filter = clean_param(vendor_filter)
    item_filter = clean_param(item_filter)
    brand_filter = clean_param(brand_filter)
    year_filter = clean_param(year_filter)
    
    # Огнооны шүүлт
    date_from = None
    date_to = None
    
    if date_from_str and date_from_str.strip():
        try:
            date_from = datetime.strptime(date_from_str, '%Y-%m-%d').date()
        except ValueError:
            pass
    
    if date_to_str and date_to_str.strip():
        try:
            date_to = datetime.strptime(date_to_str, '%Y-%m-%d').date()
        except ValueError:
            pass
    
    # OpenDataLandedCost QuerySet - БҮХ өгөгдлөөс эхэлнэ
    purchase_qs = OpenDataLandedCost.objects.all()
    
    # Огноогоор шүүх
    if date_from:
        purchase_qs = purchase_qs.filter(documentdate__gte=date_from)
    
    if date_to:
        purchase_qs = purchase_qs.filter(documentdate__lte=date_to)
    
    context = {
        'active_tab': active_tab,
        'date_from': date_from_str or '',
        'date_to': date_to_str or '',
        'selected_vendor': vendor_filter,
        'selected_item': item_filter,
        'selected_brand': brand_filter,
        'selected_year': year_filter,
    }
    
    if active_tab == 'summary':
        # НЭГТГЭЛ - Pivot хүснэгт (Харилцагч->Бараа->Он->Сар->Өдөр)
        summary_qs = purchase_qs
        pivot_result = _get_purchase_pivot_summary(summary_qs, year_filter)
        context.update({
            'pivot_data': pivot_result['data'],
            'grand_total': pivot_result['grand_total'],
        })
        
        # Статистик - TEXT талбаруудыг CAST хийж тооцоолох
        from django.db.models.functions import Cast
        from django.db.models import FloatField
        
        stats = summary_qs.annotate(
            qty_float=Cast('qty', FloatField()),
            price_float=Cast('price', FloatField()),
            amount_float=Cast('itemamountfc', FloatField())
        ).aggregate(
            total_qty=Sum('qty_float'),
            total_price=Avg('price_float'),
            total_amount=Sum('amount_float'),
            total_orders=Count('documentpkid', distinct=True),
            total_vendors=Count('vendorname', distinct=True),
        )
        context['stats'] = stats
        context['total_records'] = summary_qs.count()
        
    elif active_tab == 'by_vendor':
        # ХАРИЛЦАГЧААР (Vendor)
        vendor_qs = purchase_qs
        
        if vendor_filter:
            vendor_qs = vendor_qs.filter(vendorname__icontains=vendor_filter)
        
        if item_filter:
            vendor_qs = vendor_qs.filter(itemname__icontains=item_filter)
        
        vendor_analysis = _get_vendor_analysis(vendor_qs)
        
        # Dropdown lists
        vendors_list = purchase_qs.values_list(
            'vendorname', flat=True
        ).distinct().exclude(vendorname__isnull=True).exclude(vendorname='').order_by('vendorname')
        
        items_list = purchase_qs.values_list(
            'itemname', flat=True
        ).distinct().exclude(itemname__isnull=True).exclude(itemname='').order_by('itemname')
        
        context.update({
            'vendor_analysis': vendor_analysis,
            'vendors': list(vendors_list),
            'items': [{'name': i} for i in items_list],
        })
        
        # Статистик
        stats = vendor_qs.aggregate(
            total_qty=Sum('qty'),
            total_price=Avg('price'),
            total_amount=Sum('itemamountfc'),
            total_orders=Count('documentpkid', distinct=True),
        )
        context['stats'] = stats
        context['total_records'] = vendor_qs.count()
        
    elif active_tab == 'by_item':
        # БАРААГААР
        item_qs = purchase_qs
        
        if item_filter:
            item_qs = item_qs.filter(itemname__icontains=item_filter)
        
        if vendor_filter:
            item_qs = item_qs.filter(vendorname__icontains=vendor_filter)
        
        item_analysis = _get_item_analysis(item_qs)
        
        # Dropdown lists
        items_list = purchase_qs.values_list(
            'itemname', flat=True
        ).distinct().exclude(itemname__isnull=True).exclude(itemname='').order_by('itemname')
        
        vendors_list = purchase_qs.values_list(
            'vendorname', flat=True
        ).distinct().exclude(vendorname__isnull=True).exclude(vendorname='').order_by('vendorname')
        
        context.update({
            'item_analysis': item_analysis,
            'items': [{'name': i} for i in items_list],
            'vendors': list(vendors_list),
        })
        
        # Статистик
        stats = item_qs.aggregate(
            total_qty=Sum('qty'),
            total_price=Avg('price'),
            total_amount=Sum('itemamountfc'),
            total_orders=Count('documentpkid', distinct=True),
        )
        context['stats'] = stats
        context['total_records'] = item_qs.count()
        
    elif active_tab == 'by_date':
        # ӨДРӨӨР
        date_qs = purchase_qs
        
        date_analysis = _get_purchase_date_analysis(date_qs)
        
        # Статистик
        stats = date_qs.aggregate(
            total_qty=Sum('qty'),
            total_price=Avg('price'),
            total_amount=Sum('itemamountfc'),
            total_orders=Count('documentpkid', distinct=True),
        )
        context['stats'] = stats
        context['total_records'] = date_qs.count()
        
        context['date_analysis'] = date_analysis
    
    return render(request, 'shop/purchase_report.html', context)


# ========== HELPER FUNCTIONS - Purchase Report ==========

def _get_purchase_pivot_summary(purchase_qs, selected_year=None):
    """Харилцагч->Бараа->Он->Сар->Өдөр бүтэцтэй pivot хүснэгт"""
    from collections import defaultdict
    from datetime import datetime
    
    # БҮХ өгөгдлийг авах (зөвхөн шаардлагатай талбаруудыг)
    all_data = purchase_qs.values(
        'vendorname', 'itemname', 'documentdate', 
        'qty', 'price', 'itemamountfc', 'documentpkid'
    ).order_by('vendorname', 'itemname', 'documentdate')
    
    # Nested dictionary бүтэц: vendor -> item -> year -> month -> day
    pivot = defaultdict(lambda: {
        'items': defaultdict(lambda: {
            'years': defaultdict(lambda: {
                'months': defaultdict(lambda: {
                    'days': defaultdict(lambda: {
                        'qty': 0,
                        'price': 0,
                        'amount': 0,
                        'orders': set()
                    }),
                    'total_qty': 0,
                    'total_price': 0,
                    'total_amount': 0,
                    'total_orders': 0
                }),
                'total_qty': 0,
                'total_price': 0,
                'total_amount': 0,
                'total_orders': 0
            }),
            'total_qty': 0,
            'total_price': 0,
            'total_amount': 0,
            'total_orders': 0
        }),
        'total_qty': 0,
        'total_price': 0,
        'total_amount': 0,
        'total_orders': 0
    })
    
    # Өгөгдлийг бүлэглэх
    for row in all_data:
        vendor = row['vendorname'] or 'Тодорхойгүй'
        item = row['itemname'] or 'Тодорхойгүй'
        date = row['documentdate']
        
        # Type conversion - датабазаас string ирж магадгүй
        try:
            qty = float(row['qty']) if row['qty'] else 0
        except (ValueError, TypeError):
            qty = 0
            
        try:
            price = float(row['price']) if row['price'] else 0
        except (ValueError, TypeError):
            price = 0
            
        try:
            amount = float(row['itemamountfc']) if row['itemamountfc'] else 0
        except (ValueError, TypeError):
            amount = 0
            
        order_id = row['documentpkid']
        
        if not date:
            continue
        
        # Хэрэв date string бол date object болгох
        if isinstance(date, str):
            try:
                date = datetime.strptime(date, '%Y-%m-%d').date()
            except:
                continue
        
        year = date.year
        month = date.month
        day = date.day
        
        # Харилцагчийн items доторх item
        item_data = pivot[vendor]['items'][item]
        year_data = item_data['years'][year]
        month_data = year_data['months'][month]
        day_data = month_data['days'][day]
        
        # Өдрийн өгөгдөл
        day_data['qty'] += qty
        day_data['price'] = price  # Хамгийн сүүлийн үнэ
        day_data['amount'] += amount
        day_data['orders'].add(order_id)
        
        # Сарын нийлбэр
        month_data['total_qty'] += qty
        month_data['total_price'] = price  # Хамгийн сүүлийн үнэ
        month_data['total_amount'] += amount
        if order_id:
            month_data['total_orders'] = len([d['orders'] for d in month_data['days'].values()])
        
        # Оны нийлбэр
        year_data['total_qty'] += qty
        year_data['total_price'] = price
        year_data['total_amount'] += amount
        
        # Барааны нийлбэр
        item_data['total_qty'] += qty
        item_data['total_price'] = price
        item_data['total_amount'] += amount
        
        # Харилцагчийн нийлбэр
        pivot[vendor]['total_qty'] += qty
        pivot[vendor]['total_price'] = price
        pivot[vendor]['total_amount'] += amount
    
    # Dict-рүү хөрвүүлэх
    result = {}
    for vendor, vendor_data in sorted(pivot.items()):
        result[vendor] = {
            'total_qty': vendor_data['total_qty'],
            'total_price': vendor_data['total_price'],
            'total_amount': vendor_data['total_amount'],
            'items': {}
        }
        
        for item, item_data in sorted(vendor_data['items'].items()):
            result[vendor]['items'][item] = {
                'total_qty': item_data['total_qty'],
                'total_price': item_data['total_price'],
                'total_amount': item_data['total_amount'],
                'years': {}
            }
            
            for year, year_data in sorted(item_data['years'].items(), reverse=True):
                result[vendor]['items'][item]['years'][year] = {
                    'total_qty': year_data['total_qty'],
                    'total_price': year_data['total_price'],
                    'total_amount': year_data['total_amount'],
                    'months': {}
                }
                
                for month, month_data in sorted(year_data['months'].items()):
                    result[vendor]['items'][item]['years'][year]['months'][month] = {
                        'total_qty': month_data['total_qty'],
                        'total_price': month_data['total_price'],
                        'total_amount': month_data['total_amount'],
                        'days': dict(sorted(month_data['days'].items()))
                    }
                    
                    # Өдрийн orders set-ийг тоо болгох
                    for day, day_data in result[vendor]['items'][item]['years'][year]['months'][month]['days'].items():
                        day_data['orders'] = len(day_data['orders']) if day_data['orders'] else 0
    
    # Grand total
    grand_total = {
        'total_qty': sum(v['total_qty'] for v in result.values()),
        'total_price': sum(v['total_price'] for v in result.values()) / len(result) if result else 0,  # Дундаж үнэ
        'total_amount': sum(v['total_amount'] for v in result.values()),
    }
    
    return {'data': result, 'grand_total': grand_total}


def _get_purchase_available_years(purchase_qs):
    """Боломжит жилүүдийг авах"""
    years = purchase_qs.extra(
        select={'year': "EXTRACT(year FROM \"DocumentDate\")"}
    ).values_list('year', flat=True).distinct().order_by('-year')
    
    return [int(y) for y in years if y]


def _get_vendor_analysis(purchase_qs):
    """Харилцагчийн (vendor) дэлгэрэнгүй шинжилгээ"""
    from django.db.models.functions import Cast
    from django.db.models import FloatField
    
    vendor_stats = purchase_qs.annotate(
        qty_float=Cast('qty', FloatField()),
        price_float=Cast('price', FloatField()),
        amount_float=Cast('itemamountfc', FloatField())
    ).values('vendorname').annotate(
        total_qty=Sum('qty_float'),
        total_price=Avg('price_float'),
        total_amount=Sum('amount_float'),
        orders=Count('documentpkid', distinct=True),
        items=Count('itemname', distinct=True)
    ).order_by('-total_amount')
    
    return vendor_stats


def _get_item_analysis(purchase_qs):
    """Барааны дэлгэрэнгүй шинжилгээ"""
    from django.db.models.functions import Cast
    from django.db.models import FloatField
    
    item_stats = purchase_qs.annotate(
        qty_float=Cast('qty', FloatField()),
        price_float=Cast('price', FloatField()),
        amount_float=Cast('itemamountfc', FloatField())
    ).values('itemname').annotate(
        total_qty=Sum('qty_float'),
        total_price=Avg('price_float'),
        total_amount=Sum('amount_float'),
        orders=Count('documentpkid', distinct=True),
        vendors=Count('vendorname', distinct=True)
    ).order_by('-total_amount')
    
    return item_stats


def _get_purchase_date_analysis(purchase_qs):
    """Өдрөөр дэлгэрэнгүй шинжилгээ"""
    from django.db.models.functions import Cast
    from django.db.models import FloatField
    
    date_stats = purchase_qs.annotate(
        qty_float=Cast('qty', FloatField()),
        price_float=Cast('price', FloatField()),
        amount_float=Cast('itemamountfc', FloatField())
    ).values('documentdate').annotate(
        total_qty=Sum('qty_float'),
        total_price=Avg('price_float'),
        total_amount=Sum('amount_float'),
        orders=Count('documentpkid', distinct=True),
        vendors=Count('vendorname', distinct=True)
    ).order_by('-documentdate')
    
    return date_stats


"""
Нэме: delivery_report view
"""

@login_required
def delivery_report(request):
    """Түгээлтийн тайлан - OpenDataSaleHeader өгөгдлөөс"""
    from collections import defaultdict
    
    # Эрхийн удирдлага - нэвтэрсэн хэрэглэгчийн email шалгах
    logged_in_email = request.user.email
    is_admin = False
    user_distributor_id = None
    
    # Employee мэдээлэл олох
    try:
        employee = OpenDataEmployee.objects.get(email=logged_in_email)
        # Захирал эсвэл Ерөнхий нягтлан бол бүгдийг харуулна
        if employee.positionname in ["Захирал", "Ерөнхий нягтлан"]:
            is_admin = True
        else:
            # Жирийн ажилтны хувьд UserProfile-с DistributorId олох
            try:
                profile = request.user.profile
                user_distributor_id = profile.distributor_id
                
                # Хэрэв UserProfile-д distributor_id байхгүй бол Employee-ийн pkid-г ашиглах
                # (Түгээгч ажилтан өөрийн түгээлтийг харах)
                if not user_distributor_id:
                    user_distributor_id = employee.id
            except:
                # UserProfile байхгүй бол Employee-ийн pkid-г ашиглах
                user_distributor_id = employee.id
    except OpenDataEmployee.DoesNotExist:
        # Employee бүртгэлгүй бол UserProfile-с distributor_id-г шалгах
        try:
            profile = request.user.profile
            user_distributor_id = profile.distributor_id
        except:
            pass
    
    # Огноо фильтр
    date_from_str = request.GET.get('date_from', '')
    date_to_str = request.GET.get('date_to', '')
    seller_filter = request.GET.get('seller', '')
    year_filter = request.GET.get('year', '')
    month_filter = request.GET.get('month', '')
    
    date_from = None
    date_to = None
    
    if date_from_str and date_from_str.strip():
        try:
            date_from = datetime.strptime(date_from_str, '%Y-%m-%d').date()
        except ValueError:
            pass
    
    if date_to_str and date_to_str.strip():
        try:
            date_to = datetime.strptime(date_to_str, '%Y-%m-%d').date()
        except ValueError:
            pass
    
    # OpenDataSaleHeader QuerySet
    delivery_qs = OpenDataSaleHeader.objects.all()
    
    # Эрхийн шүүлт - админ биш бол зөвхөн өөрийн түгээгчийн өгөгдлийг харуулна
    if not is_admin and user_distributor_id:
        delivery_qs = delivery_qs.filter(distributorid=user_distributor_id)
    elif not is_admin and not user_distributor_id:
        # Эрхгүй хэрэглэгч - хоосон өгөгдол
        delivery_qs = delivery_qs.none()
    
    # DocumentDate нь TEXT учраас эхлээд Cast хийнэ
    delivery_qs = delivery_qs.annotate(
        date_field_filter=Cast('documentdate', DateField())
    )
    
    # Огноогийн автомат шүүлт - сүүлийн 1 жил
    one_year_ago = timezone.now() - timedelta(days=365)
    
    # Сүүлийн 1 жилийн өгөгдөл
    delivery_qs = delivery_qs.filter(date_field_filter__gte=one_year_ago.date())
    
    # Огноогоор шүүх (хэрэглэгчийн оруулсан)
    if date_from:
        delivery_qs = delivery_qs.filter(date_field_filter__gte=date_from)
    
    if date_to:
        delivery_qs = delivery_qs.filter(date_field_filter__lte=date_to)
    
    # Жолоочоор шүүх
    if seller_filter:
        delivery_qs = delivery_qs.filter(distributorname__icontains=seller_filter)
    
    # Хүснэгтийн баганы толгойн шүүлт
    filter_year = request.GET.get('filter_year', '').strip()
    filter_month = request.GET.get('filter_month', '').strip()
    filter_day = request.GET.get('filter_day', '').strip()
    filter_seller = request.GET.get('filter_seller', '').strip()
    
    # Өдрөөр бүлэглэх - Он-Сар-Өдөр-Жолооч
    daily_data = delivery_qs.annotate(
        date_field=Cast('documentdate', DateField()),
        year=ExtractYear('date_field'),
        month=ExtractMonth('date_field'),
        day=ExtractDay('date_field'),
        payamount_float=Cast('payamount', FloatField()),
        deliveryqty_float=Cast('deliveryqty', FloatField())
    )
    
    # Хүснэгтийн баганы толгойн шүүлт (annotate хийсний дараа)
    if filter_year:
        daily_data = daily_data.filter(year=int(filter_year))
    if filter_month:
        daily_data = daily_data.filter(month=int(filter_month))
    if filter_day:
        daily_data = daily_data.filter(day=int(filter_day))
    if filter_seller:
        daily_data = daily_data.filter(distributorname__icontains=filter_seller)
    
    # Бүлэглэх
    daily_data = daily_data.values('year', 'month', 'day', 'distributorname').annotate(
        total_amount=Sum('payamount_float'),
        total_qty=Sum('deliveryqty_float'),
        total_orders=Count('documentpkid', distinct=True)
    ).order_by('-year', '-month', '-day', 'distributorname')
    
    # Статистик - шүүлтийн дараах нийлбэр (Python дээр тооцоолох)
    daily_data_list = list(daily_data)
    stats = {
        'total_amount': sum(row['total_amount'] or 0 for row in daily_data_list),
        'total_qty': sum(row['total_qty'] or 0 for row in daily_data_list),
        'total_orders': sum(row['total_orders'] or 0 for row in daily_data_list),
        'total_sellers': len(set(row['distributorname'] for row in daily_data_list if row['distributorname']))
    }
    
    # Жолоочдын жагсаалт
    sellers_list = OpenDataSaleHeader.objects.values_list(
        'distributorname', flat=True
    ).distinct().exclude(distributorname__isnull=True).exclude(distributorname='').order_by('distributorname')
    
    # Боломжит жилүүд
    years = delivery_qs.annotate(
        year=ExtractYear('date_field_filter')
    ).values_list('year', flat=True).distinct().order_by('-year')
    
    context = {
        'daily_data': daily_data_list,
        'sellers': list(sellers_list),
        'years': [int(y) for y in years if y],
        'stats': stats,
        'date_from': date_from_str or '',
        'date_to': date_to_str or '',
        'selected_seller': seller_filter,
        'selected_year': year_filter,
        'selected_month': month_filter,
        'total_records': len(daily_data_list),
    }
    
    return render(request, 'shop/delivery_report.html', context)


@login_required
def expense_report(request):
    """Зардлын тооцоолол - Expense Calculation Report"""
    # Зөвхөн ажилтанд зөвшөөрөх
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')
    
    from datamigration.models import OpenDataSale
    from django.db import connection
    from collections import defaultdict
    
    # Сонгосон он (default - сүүлийн он)
    selected_year_str = request.GET.get('year', '').replace(',', '')  # Мянгачлалыг арилгах
    
    # OpenDataSale-с жилүүдийг авах
    available_years = _get_available_years(OpenDataSale.objects.all())
    
    # Хэрэв available_years хоосон бол сүүлийн 5 жилийг үүсгэх
    if not available_years:
        current_year = timezone.now().year
        available_years = list(range(current_year, current_year - 5, -1))
    
    # Default он - хамгийн сүүлийн он
    if not selected_year_str and available_years:
        selected_year = available_years[0]
    elif selected_year_str:
        selected_year = int(selected_year_str)
    else:
        selected_year = timezone.now().year
    
    # Сараар өгөгдөл цуглуулах (1-12 сар)
    monthly_data = {i: {} for i in range(1, 13)}
    
    # 1. OpenDataSale - PayAmount (Борлуулалтын дүн)
    with connection.cursor() as cursor:
        cursor.execute(
            '''
            SELECT 
                EXTRACT(month FROM CAST("DocumentDate" AS DATE)) as month,
                SUM(CAST("PayAmount" AS FLOAT)) as payamount_sum,
                SUM(CAST("VatAmount" AS FLOAT)) as vatamount_sum,
                SUM(CAST("AmountNonVat" AS FLOAT)) as amountnonvat_sum,
                SUM(CAST("UnitCost" AS FLOAT) * CAST("Qty" AS FLOAT)) as cost_sum
            FROM "OpenDataSale"
            WHERE EXTRACT(year FROM CAST("DocumentDate" AS DATE)) = %s
            GROUP BY EXTRACT(month FROM CAST("DocumentDate" AS DATE))
            ''',
            [selected_year]
        )
        for row in cursor.fetchall():
            month = int(row[0]) if row[0] else 0
            if 1 <= month <= 12:
                monthly_data[month]['payamount'] = row[1] or 0
                monthly_data[month]['vatamount'] = row[2] or 0
                monthly_data[month]['amountnonvat'] = row[3] or 0
                monthly_data[month]['cost'] = row[4] or 0
    
    # 2. OpenDataGeneralLedger - 520101 Борлуулалтын хөнгөлөлт
    with connection.cursor() as cursor:
        cursor.execute(
            '''
            SELECT 
                EXTRACT(month FROM CAST("DocumentDate" AS DATE)) as month,
                SUM(CAST("DebitAmt" AS FLOAT)) as total
            FROM "OpenDataGeneralLedger"
            WHERE EXTRACT(year FROM CAST("DocumentDate" AS DATE)) = %s
              AND "AccountId" = %s
            GROUP BY EXTRACT(month FROM CAST("DocumentDate" AS DATE))
            ''',
            [selected_year, '520101']
        )
        for row in cursor.fetchall():
            month = int(row[0]) if row[0] else 0
            if 1 <= month <= 12:
                monthly_data[month]['discount'] = row[1] or 0
    
    # 3. OpenDataSaleRefund - Буцаалтын дүн
    with connection.cursor() as cursor:
        cursor.execute(
            '''
            SELECT 
                EXTRACT(month FROM CAST("DocumentDate" AS DATE)) as month,
                SUM(CAST("PayAmount" AS FLOAT)) as total
            FROM "OpenDataSaleRefund"
            WHERE EXTRACT(year FROM CAST("DocumentDate" AS DATE)) = %s
            GROUP BY EXTRACT(month FROM CAST("DocumentDate" AS DATE))
            ''',
            [selected_year]
        )
        for row in cursor.fetchall():
            month = int(row[0]) if row[0] else 0
            if 1 <= month <= 12:
                monthly_data[month]['refund'] = row[1] or 0
    
    # 4. Бохир ашиг = AmountNonVat - Өртөг
    # 5. Хөнгөлөлт, буцаалт хассан зардлын өмнөх ашиг = AmountNonVat - Хөнгөлөлт - Буцаалт - Өртөг
    for month in range(1, 13):
        amountnonvat = monthly_data[month].get('amountnonvat', 0)
        discount = monthly_data[month].get('discount', 0)
        refund = monthly_data[month].get('refund', 0)
        cost = monthly_data[month].get('cost', 0)
        
        # Бохир ашиг (зөвхөн өртөг хассан)
        monthly_data[month]['gross_profit'] = amountnonvat - cost
        
        # Хөнгөлөлт, буцаалт хассан зардлын өмнөх ашиг
        monthly_data[month]['gross_profit_after_discount'] = amountnonvat - discount - refund - cost
    
    # 5. OpenDataGeneralLedger - 7-р эхэлсэн зардлууд (дансны дугаараар бүлэглэх)
    with connection.cursor() as cursor:
        cursor.execute(
            '''
            SELECT 
                EXTRACT(month FROM CAST("DocumentDate" AS DATE)) as month,
                "AccountId" as accountid,
                "AccountName" as accountname,
                SUM(CAST("DebitAmt" AS FLOAT)) as total
            FROM "OpenDataGeneralLedger"
            WHERE EXTRACT(year FROM CAST("DocumentDate" AS DATE)) = %s
              AND "AccountId" LIKE '7%%'
            GROUP BY EXTRACT(month FROM CAST("DocumentDate" AS DATE)), "AccountId", "AccountName"
            ORDER BY "AccountId", EXTRACT(month FROM CAST("DocumentDate" AS DATE))
            ''',
            [selected_year]
        )
        expenses_by_month = cursor.fetchall()
    
    # Зардлын дансуудыг цуглуулах
    expense_accounts = {}
    for row in expenses_by_month:
        month = int(row[0]) if row[0] else 0  # month
        account_id = row[1] or 'Unknown'  # accountid
        account_name = row[2] or 'Тодорхойгүй'  # accountname
        total = row[3] or 0  # total
        
        if account_id not in expense_accounts:
            expense_accounts[account_id] = {
                'name': account_name,
                'monthly': {i: 0 for i in range(1, 13)}
            }
        
        if 1 <= month <= 12:
            expense_accounts[account_id]['monthly'][month] = total
    
    # Нийт зардал сараар
    total_expenses_monthly = {i: 0 for i in range(1, 13)}
    for account_data in expense_accounts.values():
        for month in range(1, 13):
            total_expenses_monthly[month] += account_data['monthly'][month]
    
    # 6. Цэвэр ашиг = Хөнгөлөлт, буцаалт хассан зардлын өмнөх ашиг - Нийт зардал
    for month in range(1, 13):
        gross_profit_after_discount = monthly_data[month].get('gross_profit_after_discount', 0)
        total_expense = total_expenses_monthly[month]
        monthly_data[month]['net_profit'] = gross_profit_after_discount - total_expense
    
    # Жилийн нийлбэр тооцоолох
    yearly_totals = {
        'payamount': sum(monthly_data[m].get('payamount', 0) for m in range(1, 13)),
        'vatamount': sum(monthly_data[m].get('vatamount', 0) for m in range(1, 13)),
        'amountnonvat': sum(monthly_data[m].get('amountnonvat', 0) for m in range(1, 13)),
        'discount': sum(monthly_data[m].get('discount', 0) for m in range(1, 13)),
        'refund': sum(monthly_data[m].get('refund', 0) for m in range(1, 13)),
        'cost': sum(monthly_data[m].get('cost', 0) for m in range(1, 13)),
        'gross_profit': sum(monthly_data[m].get('gross_profit', 0) for m in range(1, 13)),
        'gross_profit_after_discount': sum(monthly_data[m].get('gross_profit_after_discount', 0) for m in range(1, 13)),
        'total_expense': sum(total_expenses_monthly.values()),
        'net_profit': sum(monthly_data[m].get('net_profit', 0) for m in range(1, 13)),
    }
    
    # Зардлын дансны жилийн дүн
    for account_id in expense_accounts:
        expense_accounts[account_id]['yearly_total'] = sum(
            expense_accounts[account_id]['monthly'].values()
        )
    
    context = {
        'selected_year': selected_year,
        'available_years': available_years,
        'monthly_data': monthly_data,
        'expense_accounts': dict(sorted(expense_accounts.items())),
        'total_expenses_monthly': total_expenses_monthly,
        'yearly_totals': yearly_totals,
        'months': range(1, 13),
    }
    
    return render(request, 'shop/expense_report.html', context)


@staff_member_required
def expense_report_by_year(request):
    """Зардлын тооцоолол жилээр - Expense Report by Year (Columns = Years)"""
    from datamigration.models import OpenDataSale
    from django.db import connection
    from collections import defaultdict
    
    # OpenDataSale-с жилүүдийг авах
    available_years = _get_available_years(OpenDataSale.objects.all())
    
    # Хэрэв available_years хоосон бол сүүлийн 5 жилийг үүсгэх
    if not available_years:
        current_year = timezone.now().year
        available_years = list(range(current_year, current_year - 5, -1))
    
    # Жил бүрийн өгөгдөл
    yearly_data = {year: {} for year in available_years}
    
    # 1. OpenDataSale - Борлуулалт, НӨАТ, өртөг
    for year in available_years:
        with connection.cursor() as cursor:
            cursor.execute(
                '''
                SELECT 
                    SUM(CAST("PayAmount" AS FLOAT)) as payamount_sum,
                    SUM(CAST("VatAmount" AS FLOAT)) as vatamount_sum,
                    SUM(CAST("AmountNonVat" AS FLOAT)) as amountnonvat_sum,
                    SUM(CAST("UnitCost" AS FLOAT) * CAST("Qty" AS FLOAT)) as cost_sum
                FROM "OpenDataSale"
                WHERE EXTRACT(year FROM CAST("DocumentDate" AS DATE)) = %s
                ''',
                [year]
            )
            row = cursor.fetchone()
            if row:
                yearly_data[year]['payamount'] = row[0] or 0
                yearly_data[year]['vatamount'] = row[1] or 0
                yearly_data[year]['amountnonvat'] = row[2] or 0
                yearly_data[year]['cost'] = row[3] or 0
    
    # 2. Борлуулалтын хөнгөлөлт (520101)
    for year in available_years:
        with connection.cursor() as cursor:
            cursor.execute(
                '''
                SELECT SUM(CAST("DebitAmt" AS FLOAT)) as total
                FROM "OpenDataGeneralLedger"
                WHERE EXTRACT(year FROM CAST("DocumentDate" AS DATE)) = %s
                  AND "AccountId" = %s
                ''',
                [year, '520101']
            )
            row = cursor.fetchone()
            yearly_data[year]['discount'] = (row[0] or 0) if row else 0
    
    # 3. Буцаалт
    for year in available_years:
        with connection.cursor() as cursor:
            cursor.execute(
                '''
                SELECT SUM(CAST("PayAmount" AS FLOAT)) as total
                FROM "OpenDataSaleRefund"
                WHERE EXTRACT(year FROM CAST("DocumentDate" AS DATE)) = %s
                ''',
                [year]
            )
            row = cursor.fetchone()
            yearly_data[year]['refund'] = (row[0] or 0) if row else 0
    
    # 4. Бохир ашиг болон хөнгөлөлт/буцаалт хассан ашиг
    for year in available_years:
        amountnonvat = yearly_data[year].get('amountnonvat', 0)
        cost = yearly_data[year].get('cost', 0)
        discount = yearly_data[year].get('discount', 0)
        refund = yearly_data[year].get('refund', 0)
        
        yearly_data[year]['gross_profit'] = amountnonvat - cost
        yearly_data[year]['gross_profit_after_discount'] = amountnonvat - cost - discount - refund
    
    # 5. Зардлууд (7-р данс) - дансаар бүлэглэх
    expense_accounts = {}
    for year in available_years:
        with connection.cursor() as cursor:
            cursor.execute(
                '''
                SELECT 
                    "AccountId" as accountid,
                    "AccountName" as accountname,
                    SUM(CAST("DebitAmt" AS FLOAT)) as total
                FROM "OpenDataGeneralLedger"
                WHERE EXTRACT(year FROM CAST("DocumentDate" AS DATE)) = %s
                  AND "AccountId" LIKE '7%%'
                GROUP BY "AccountId", "AccountName"
                ORDER BY "AccountId"
                ''',
                [year]
            )
            for row in cursor.fetchall():
                account_id = row[0] or 'Unknown'
                account_name = row[1] or 'Тодорхойгүй'
                total = row[2] or 0
                
                if account_id not in expense_accounts:
                    expense_accounts[account_id] = {
                        'name': account_name,
                        'yearly': {y: 0 for y in available_years}
                    }
                
                expense_accounts[account_id]['yearly'][year] = total
    
    # Нийт зардал жил бүрээр
    total_expenses_yearly = {year: 0 for year in available_years}
    for account_data in expense_accounts.values():
        for year in available_years:
            total_expenses_yearly[year] += account_data['yearly'][year]
    
    # 6. Цэвэр ашиг
    for year in available_years:
        gross_profit_after_discount = yearly_data[year].get('gross_profit_after_discount', 0)
        total_expense = total_expenses_yearly[year]
        yearly_data[year]['total_expense'] = total_expense
        yearly_data[year]['net_profit'] = gross_profit_after_discount - total_expense
    
    context = {
        'available_years': available_years,
        'yearly_data': yearly_data,
        'expense_accounts': dict(sorted(expense_accounts.items())),
        'total_expenses_yearly': total_expenses_yearly,
    }
    
    return render(request, 'shop/expense_report_by_year.html', context)


@staff_member_required
def expense_transaction_detail(request):
    """Зардлын гүйлгээний дэлгэрэнгүй"""
    from django.db import connection
    import logging
    logger = logging.getLogger(__name__)
    
    year = request.GET.get('year')
    month = request.GET.get('month')
    account_id = request.GET.get('account_id')
    transaction_type = request.GET.get('type')  # 'discount', 'sale', 'refund', 'expense'
    
    logger.info(f"Transaction detail requested: year={year}, month={month}, type={transaction_type}, account={account_id}")
    
    if not year or not month:
        # Параметр дутуу бол error template үзүүлэх
        context = {
            'title': 'Алдаа - Параметр дутуу',
            'year': year or 'N/A',
            'month': month or 'N/A',
            'transactions': [],
            'total_amount': 0,
            'transaction_count': 0,
            'error_message': f'Он эсвэл сар дутуу байна. Year={year}, Month={month}',
        }
        return render(request, 'shop/expense_transaction_detail.html', context)
    
    transactions = []
    title = ""
    
    try:
        with connection.cursor() as cursor:
            if transaction_type == 'discount' or account_id == '520101':
                # Борлуулалтын хөнгөлөлт (520101)
                title = f"{year} оны {month}-р сар - Борлуулалтын хөнгөлөлт (520101)"
                cursor.execute('''
                    WITH main_line AS (
                        SELECT 
                            "DocumentPkId",
                            CAST("DocumentDate" AS DATE) as date,
                            COALESCE("DocumentDesc", "DocumentNumber", '-') as description,
                            "AccountId" || ' - ' || "AccountName" as main_account,
                            CAST("DebitAmt" AS FLOAT) as main_debit,
                            CAST("CreditAmt" AS FLOAT) as main_credit
                        FROM "OpenDataGeneralLedger"
                        WHERE EXTRACT(year FROM CAST("DocumentDate" AS DATE)) = %s
                          AND EXTRACT(month FROM CAST("DocumentDate" AS DATE)) = %s
                          AND "AccountId" = %s
                    ),
                    paired_lines AS (
                        SELECT 
                            gl."DocumentPkId",
                            STRING_AGG(gl."AccountId" || ' - ' || gl."AccountName", ', ' ORDER BY gl."AccountId") as other_accounts
                        FROM "OpenDataGeneralLedger" gl
                        INNER JOIN main_line ml ON gl."DocumentPkId" = ml."DocumentPkId"
                        WHERE gl."AccountId" != %s
                        GROUP BY gl."DocumentPkId"
                    )
                    SELECT 
                        ml.date,
                        ml.description,
                        ml.main_account as debit_account,
                        COALESCE(pl.other_accounts, '-') as credit_account,
                        CASE WHEN ml.main_debit > 0 THEN ml.main_debit ELSE ml.main_credit END as amount
                    FROM main_line ml
                    LEFT JOIN paired_lines pl ON ml."DocumentPkId" = pl."DocumentPkId"
                    ORDER BY ml.date
                ''', [year, month, '520101', '520101'])
                logger.info(f"Discount query executed")
                
            elif transaction_type == 'sale':
                # Борлуулалт (OpenDataSale)
                title = f"{year} оны {month}-р сар - Борлуулалт"
                cursor.execute('''
                    SELECT 
                        CAST("DocumentDate" AS DATE) as date,
                        "ItemName" || ' (' || "CustomerName" || ')' as description,
                        'Барааны борлуулалт' as debit_account,
                        'Орлого' as credit_account,
                        CAST("PayAmount" AS FLOAT) as amount
                    FROM "OpenDataSale"
                    WHERE EXTRACT(year FROM CAST("DocumentDate" AS DATE)) = %s
                      AND EXTRACT(month FROM CAST("DocumentDate" AS DATE)) = %s
                    ORDER BY CAST("DocumentDate" AS DATE), "DocumentPkId"
                ''', [year, month])
                logger.info(f"Sale query executed")
                
            elif transaction_type == 'refund':
                # Буцаалт
                title = f"{year} оны {month}-р сар - Буцаалт"
                cursor.execute('''
                    SELECT 
                        "DocumentDate" as date,
                        "ItemName" || ' (' || "CustomerName" || ')' as description,
                        'Буцаалт' as debit_account,
                        'Орлого' as credit_account,
                        CAST("PayAmount" AS FLOAT) as amount
                    FROM "OpenDataSaleRefund"
                    WHERE EXTRACT(year FROM CAST("DocumentDate" AS DATE)) = %s
                      AND EXTRACT(month FROM CAST("DocumentDate" AS DATE)) = %s
                    ORDER BY CAST("DocumentDate" AS DATE)
                ''', [year, month])
                logger.info(f"Refund query executed")
                
            elif account_id and account_id.startswith('7'):
                # 7-р зардлын данс
                cursor.execute('''
                    SELECT 
                        "AccountName"
                    FROM "OpenDataGeneralLedger"
                    WHERE "AccountId" = %s
                    LIMIT 1
                ''', [account_id])
                account_row = cursor.fetchone()
                account_name = account_row[0] if account_row else 'Тодорхойгүй данс'
                
                title = f"{year} оны {month}-р сар - {account_id} ({account_name})"
                cursor.execute('''
                    WITH main_line AS (
                        SELECT 
                            "DocumentPkId",
                            CAST("DocumentDate" AS DATE) as date,
                            COALESCE("DocumentDesc", "DocumentNumber", '-') as description,
                            "AccountId" || ' - ' || "AccountName" as main_account,
                            CAST("DebitAmt" AS FLOAT) as main_debit,
                            CAST("CreditAmt" AS FLOAT) as main_credit
                        FROM "OpenDataGeneralLedger"
                        WHERE EXTRACT(year FROM CAST("DocumentDate" AS DATE)) = %s
                          AND EXTRACT(month FROM CAST("DocumentDate" AS DATE)) = %s
                          AND "AccountId" = %s
                    ),
                    paired_lines AS (
                        SELECT 
                            gl."DocumentPkId",
                            STRING_AGG(gl."AccountId" || ' - ' || gl."AccountName", ', ' ORDER BY gl."AccountId") as other_accounts
                        FROM "OpenDataGeneralLedger" gl
                        INNER JOIN main_line ml ON gl."DocumentPkId" = ml."DocumentPkId"
                        WHERE gl."AccountId" != %s
                        GROUP BY gl."DocumentPkId"
                    )
                    SELECT 
                        ml.date,
                        ml.description,
                        ml.main_account as debit_account,
                        COALESCE(pl.other_accounts, '-') as credit_account,
                        CASE WHEN ml.main_debit > 0 THEN ml.main_debit ELSE ml.main_credit END as amount
                    FROM main_line ml
                    LEFT JOIN paired_lines pl ON ml."DocumentPkId" = pl."DocumentPkId"
                    ORDER BY ml.date
                ''', [year, month, account_id, account_id])
                logger.info(f"Expense account query executed for {account_id}")
            
            else:
                # Гүйлгээний төрөл тодорхойгүй
                logger.warning(f"Invalid transaction type: {transaction_type}")
                context = {
                    'title': 'Алдаа - Гүйлгээний төрөл тодорхойгүй',
                    'year': year,
                    'month': month,
                    'transactions': [],
                    'total_amount': 0,
                    'transaction_count': 0,
                    'error_message': f'Гүйлгээний төрөл тодорхойгүй: type={transaction_type}, account_id={account_id}',
                }
                return render(request, 'shop/expense_transaction_detail.html', context)
            
            # Гүйлгээнүүдийг цуглуулах
            for row in cursor.fetchall():
                transactions.append({
                    'date': row[0],
                    'description': row[1] or '-',
                    'debit_account': row[2] or '-',
                    'credit_account': row[3] or '-',
                    'amount': row[4] or 0
                })
            
            logger.info(f"Found {len(transactions)} transactions")
    
    except Exception as e:
        logger.error(f"Error in expense_transaction_detail: {str(e)}", exc_info=True)
        context = {
            'title': 'Алдаа гарлаа',
            'year': year,
            'month': month,
            'transactions': [],
            'total_amount': 0,
            'transaction_count': 0,
            'error_message': f'Өгөгдлийн сангаас мэдээлэл авахад алдаа гарлаа: {str(e)}',
        }
        return render(request, 'shop/expense_transaction_detail.html', context)
    
    # Нийт дүн
    total_amount = sum(t['amount'] for t in transactions)
    
    logger.info(f"Rendering template with {len(transactions)} transactions, total: {total_amount}")
    
    context = {
        'title': title,
        'year': year,
        'month': month,
        'transactions': transactions,
        'total_amount': total_amount,
        'transaction_count': len(transactions),
    }
    
    return render(request, 'shop/expense_transaction_detail.html', context)


@login_required
def profile(request):
    """Хувийн мэдээлэл - OpenDataEmployee + UserProfile"""
    from django.contrib import messages
    from shop.models import UserProfile

    employee = None
    if request.user.email:
        employee = OpenDataEmployee.objects.filter(email__iexact=request.user.email).first()

    user_profile, _ = UserProfile.objects.get_or_create(user=request.user)

    if request.method == 'POST':
        user_profile.phone = request.POST.get('phone', '').strip()
        user_profile.save()
        messages.success(request, 'Мэдээлэл амжилттай хадгалагдлаа.')
        return redirect('shop:profile')

    context = {
        'employee': employee,
        'user_profile': user_profile,
    }
    return render(request, 'shop/profile.html', context)


MONTH_NAMES_MN = ['1-р сар', '2-р сар', '3-р сар', '4-р сар', '5-р сар', '6-р сар',
                  '7-р сар', '8-р сар', '9-р сар', '10-р сар', '11-р сар', '12-р сар']


@login_required
def my_attendance(request):
    """Хувийн мэдээлэл - Цагийн бүртгэл: нэвтэрсэн ажилтны зөвхөн өөрийнх нь ирц/хоцролтыг сараар харуулна.

    Ажилтныг Хувийн мэдээлэлтэй адил имэйлээр (OpenDataEmployee) тодорхойлж, мөрүүдийг Цаг бүртгэлийн
    тооцоотой ижил дүрмээр (build_attendance_rows) тооцоолно. Анхдагчаар тухайн сар, ?year=&month=-ээр бусад сар.
    """
    import calendar
    from shop.models import AttendanceRecord, AttendanceRawLog, EmployeeFingerprint, MerchandiserVisitLog
    from shop.services.attendance import build_attendance_rows

    employee = None
    if request.user.email:
        employee = OpenDataEmployee.objects.filter(email__iexact=request.user.email).first()

    today = timezone.localdate()
    year, month = _month_from_request(request.GET)

    date_from = datetime(year, month, 1).date()
    date_to = date_from.replace(day=calendar.monthrange(year, month)[1])

    prev_month = (date_from - timedelta(days=1)).replace(day=1)
    next_month = date_to + timedelta(days=1)

    rows = []
    has_fingerprint = False
    years = [today.year]
    if employee and employee.id:
        rows = build_attendance_rows(date_from, date_to, employee_code=employee.id)
        rows.sort(key=lambda r: r['date'])

        device_user_ids = list(
            EmployeeFingerprint.objects.filter(employee_code=employee.id)
            .exclude(device_user_id='').values_list('device_user_id', flat=True)
        )
        has_fingerprint = bool(device_user_ids)

        # Он сонгох жагсаалт: ажилтны хамгийн эртний бүртгэлийн оноос одоог хүртэл
        earliest = [
            AttendanceRawLog.objects.filter(device_user_id__in=device_user_ids).aggregate(m=Min('timestamp'))['m'],
            MerchandiserVisitLog.objects.filter(employee_code=employee.id).aggregate(m=Min('timestamp'))['m'],
            AttendanceRecord.objects.filter(employee_code=employee.id).aggregate(m=Min('date'))['m'],
        ]
        earliest_years = [d.year for d in earliest if d]
        if earliest_years:
            years = list(range(min(earliest_years), today.year + 1))
    if year not in years:
        years = sorted(set(years) | {year})

    summary = {
        'days': len(rows),
        'late_days': sum(1 for r in rows if r['total_late_minutes']),
        'late_minutes': sum(r['total_late_minutes'] for r in rows),
        'missing_left': sum(1 for r in rows if not r['left_time']),
    }

    context = {
        'employee': employee,
        'rows': rows,
        'summary': summary,
        'has_fingerprint': has_fingerprint,
        'has_customers': any(r['customer_name'] for r in rows),
        'year': year,
        'month': month,
        'years': list(reversed(years)),
        'months': list(enumerate(MONTH_NAMES_MN, start=1)),
        'month_label': f'{year} оны {MONTH_NAMES_MN[month - 1]}',
        'prev_month': prev_month,
        'next_month': next_month,
        'is_current_month': (year, month) == (today.year, today.month),
    }
    return render(request, 'shop/my_attendance.html', context)


def _personal_period(request, earliest_year):
    """Хувийн мэдээллийн цалин / нэмэгдлийн хуудасны шүүлтүүр: харах хэлбэр (сар / жил), он, сар. Ирээдүйн сарыг
    одоогийн сар руу шахна. Он сонгох жагсаалт earliest_year-ээс одоог хүртэл (буурахаар)."""
    today = timezone.localdate()
    mode = 'year' if request.GET.get('view') == 'year' else 'month'
    year, month = _month_from_request(request.GET)
    if (year, month) > (today.year, today.month):
        year, month = today.year, today.month
    first_year = min(earliest_year or today.year, year)
    date_from = datetime(year, month, 1).date()
    prev_month = (date_from - timedelta(days=1)).replace(day=1)
    next_month = (date_from + timedelta(days=32)).replace(day=1)
    return {
        'mode': mode,
        'year': year,
        'month': month,
        'years': list(range(today.year, first_year - 1, -1)),
        'months': list(enumerate(MONTH_NAMES_MN, start=1)),
        'month_label': f'{year} оны {MONTH_NAMES_MN[month - 1]}',
        'prev_month': prev_month,
        'next_month': next_month if (next_month.year, next_month.month) <= (today.year, today.month) else None,
        'is_current_month': (year, month) == (today.year, today.month),
        'current_year': today.year,
        # Жилээр харахад тухайн жилийн харуулах сарууд (энэ жил бол одоогийн сар хүртэл)
        'year_months': list(range(1, (today.month if year == today.year else 12) + 1)),
    }


PERSONAL_PAYROLL_SUM_KEYS = (
    'gross', 'sales_bonus', 'total_tax', 'deductions_total', 'benefits_total', 'net_pay', 'deposit', 'payout',
)


@login_required
def my_payroll(request):
    """Хувийн мэдээлэл - Цалингийн мэдээлэл: нэвтэрсэн ажилтны зөвхөн өөрийнх нь цалингийн задаргаа.

    Цалин бодолтын (build_payroll) тооцоог тухайн ажилтанд л бодож, карт / бэлэн хуудсаар нь харуулна.
    Шүүлтүүр: ?view=month|year (сараар задаргаа / жилийн хураангуй), ?year=&month=, ?sheet=all|card|cash
    (хосолсон олгох хэлбэртэй ажилтанд)."""
    from shop.models import EmployeePayProfile, PayrollEntry
    from shop.services.payroll import BENEFIT_FIELDS, DEDUCTION_COLUMNS
    from shop.services.payroll_snapshot import get_payroll

    employee = None
    if request.user.email:
        employee = OpenDataEmployee.objects.filter(email__iexact=request.user.email).first()
    code = employee.id if employee and employee.id else None

    period = _personal_period(request, PayrollEntry.objects.aggregate(m=Min('year'))['m'])
    year, month = period['year'], period['month']

    profile = EmployeePayProfile.objects.filter(employee_code=code).first() if code else None
    pay_type = profile.pay_type if profile else ''
    sheet_labels = {PayrollEntry.SHEET_CARD: 'Картын цалин', PayrollEntry.SHEET_CASH: 'Бэлэн цалин'}
    available = [
        s for s, types in ((PayrollEntry.SHEET_CARD, ('card', 'mixed')), (PayrollEntry.SHEET_CASH, ('cash', 'mixed')))
        if pay_type in types
    ]
    sheet = request.GET.get('sheet')
    if sheet in available:
        sheets = [sheet]
    else:
        sheet, sheets = 'all', available

    def payroll_row(y, m, s):
        rows, _, _, close = get_payroll(y, m, s, code)
        row = rows[0] if rows else None
        if row:
            row['closed'] = close is not None
            row['deductions_total'] = sum((row[f] for f, _, _ in DEDUCTION_COLUMNS), Decimal(0))
            row['benefits_total'] = sum((row[f] for f, _ in BENEFIT_FIELDS), Decimal(0))
        return row

    context = {
        **period,
        'employee': employee,
        'has_profile': bool(available),
        'pay_type_label': dict(EmployeePayProfile.PAY_CHOICES).get(pay_type, ''),
        'sheet': sheet,
        # Хосолсон олгох хэлбэртэй ажилтанд л карт / бэлэн сонгох шүүлтүүр
        'sheet_choices': [(s, sheet_labels[s]) for s in available] if len(available) > 1 else [],
    }

    if code and sheets and period['mode'] == 'month':
        details = []
        for s in sheets:
            row = payroll_row(year, month, s)
            if not row:
                continue
            is_card = s == PayrollEntry.SHEET_CARD
            calc = [
                ('Бодогдсон цалин', row['earned_salary'],
                 '' if row['is_shift'] or not row['required_hours']
                 else f"{row['nominal_salary']:,.0f} ÷ {row['required_hours']:g} цаг × {row['worked_hours']:g} цаг"),
                ('Бямбад ажилласан нэмэгдэл', row['saturday_bonus'],
                 f"{row['saturday_hours']:g} цаг, {row['saturday_bonus_percent']:g}%" if row['saturday_hours'] else ''),
                ('Амралтын мөнгө', row['holiday_pay'], ''),
                ('Борлуулалтын нэмэгдэл', row['sales_bonus'], ''),
                ('Унааны мөнгө', row['transport'], ''),
            ]
            details.append({
                'sheet': s,
                'label': sheet_labels[s],
                'is_card': is_card,
                'row': row,
                'base': [
                    ('Нэг гарын мөнгө × ажилласан хоног' if row['is_shift'] else 'Үндсэн цалин', row['base_salary']),
                    ('Удаан жилийн нэмэгдэл', row['seniority_bonus']),
                ],
                # Тэг мөрийг (Бодогдсон цалингаас бусад) харуулахгүй
                'calc': [c for i, c in enumerate(calc) if i == 0 or c[1]],
                'taxes': [
                    (f"НДШ ({row['ndsh_employee_rate']:g}%)", row['ndsh_employee'], ''),
                    ('ХХОАТ', row['pit'],
                     f"({row['taxable']:,.0f} × хувь) − хөнгөлөлт {row['credit']:,.0f}" if row['taxable'] else ''),
                ] if is_card else [],
                'deductions': [(label, row[f]) for f, label, _ in DEDUCTION_COLUMNS if row[f]],
                'benefits': [(label, row[f]) for f, label in BENEFIT_FIELDS if row[f]],
            })
        context['details'] = details
        # Хаагдаагүй (нягтлан баталгаажуулаагүй) хуудас байвал урьдчилсан тооцоо
        context['is_provisional'] = any(not d['row']['closed'] for d in details)
        context['month_total'] = {
            key: sum((d['row'][key] for d in details), Decimal(0)) for key in PERSONAL_PAYROLL_SUM_KEYS
        }
        # Урьдчилгаа суутгалд орсон тул сарын нийт авах дүн = олгох дүн + урьдчилгаа
        context['month_total']['advance'] = sum((d['row']['advance'] for d in details), Decimal(0))
        context['month_total']['payout_with_advance'] = context['month_total']['payout'] + context['month_total']['advance']

    if code and sheets and period['mode'] == 'year':
        year_rows = []
        for m in period['year_months']:
            by_sheet = {s: payroll_row(year, m, s) for s in sheets}
            found = [r for r in by_sheet.values() if r]
            if not found:
                continue
            item = {
                'month': m,
                'label': MONTH_NAMES_MN[m - 1],
                'worked_days': max(r['worked_days'] for r in found),
                'provisional': any(not r['closed'] for r in found),
                'by_sheet': [by_sheet[s]['payout'] if by_sheet[s] else None for s in sheets],
                **{key: sum((r[key] for r in found), Decimal(0)) for key in PERSONAL_PAYROLL_SUM_KEYS},
            }
            # Ажилласан хоноггүй, бодогдсон цалингүй сарыг хоосон гэж тэмдэглэж нийт дүнд оруулахгүй
            item['empty'] = not item['worked_days'] and not item['gross']
            year_rows.append(item)
        filled = [r for r in year_rows if not r['empty']]
        year_total = {key: sum((r[key] for r in filled), Decimal(0)) for key in PERSONAL_PAYROLL_SUM_KEYS}
        year_total['worked_days'] = sum((Decimal(r['worked_days']) for r in filled), Decimal(0))
        year_total['by_sheet'] = [
            sum((r['by_sheet'][i] or 0 for r in filled), Decimal(0)) for i in range(len(sheets))
        ]
        context.update({
            'year_rows': year_rows,
            'year_total': year_total,
            'year_months_count': len(filled),
            'year_avg_payout': year_total['payout'] / len(filled) if filled else Decimal(0),
            'sheet_columns': [sheet_labels[s] for s in sheets] if len(sheets) > 1 else [],
        })

    return render(request, 'shop/my_payroll.html', context)


@login_required
def my_sales_bonus(request):
    """Хувийн мэдээлэл - Борлуулалтын нэмэгдэл: нэвтэрсэн ажилтны хадгалагдсан (цалин бодолтод орсон) нэмэгдэл.

    Сараар харахад хүснэгт бүрийн (Нярав, ХТ, Борлуулагч, Түгээгч) задаргааг хадгалсан утгуудаар нь
    (sales_bonus.compute) дахин бодож харуулна; олгох дүн нь хадгалсан SalesBonusEntry.total.
    Шүүлтүүр: ?view=month|year, ?year=&month=, ?scheme= (ажилтан хэд хэдэн хүснэгтэд орсон бол)."""
    from types import SimpleNamespace
    from shop.models import PayrollSetting, SalesBonusEntry, SalesBonusMember
    from shop.services import sales_bonus as sb

    employee = None
    if request.user.email:
        employee = OpenDataEmployee.objects.filter(email__iexact=request.user.email).first()
    code = employee.id if employee and employee.id else None

    my_entries = SalesBonusEntry.objects.filter(employee_code=code) if code else SalesBonusEntry.objects.none()
    period = _personal_period(request, my_entries.aggregate(m=Min('year'))['m'])
    year, month = period['year'], period['month']

    scheme_order = [s for s, _ in SalesBonusMember.SCHEME_CHOICES]
    scheme_labels = dict(SalesBonusMember.SCHEME_CHOICES)
    my_schemes = [s for s in scheme_order if s in set(my_entries.values_list('scheme', flat=True))]
    scheme = request.GET.get('scheme')
    if scheme in my_schemes:
        my_entries = my_entries.filter(scheme=scheme)
    else:
        scheme = 'all'

    def amounts(entries):
        return {
            'total': sum((e.total for e in entries), Decimal(0)),
            'card': sum((e.card_amount for e in entries), Decimal(0)),
            'cash': sum((e.cash_amount for e in entries), Decimal(0)),
        }

    year_entries = sorted(my_entries.filter(year=year), key=lambda e: (e.month, scheme_order.index(e.scheme)))
    context = {
        **period,
        'employee': employee,
        'has_any': bool(my_schemes),
        'scheme': scheme,
        # Хэд хэдэн хүснэгтэд (жиш: ХТ нь няравыг орлосон) нэмэгдэл бодогдсон бол л хүснэгт сонгох шүүлтүүр
        'scheme_choices': [(s, scheme_labels[s]) for s in my_schemes] if len(my_schemes) > 1 else [],
        'year_total': amounts(year_entries),
    }

    if period['mode'] == 'year':
        by_month = {}
        for e in year_entries:
            by_month.setdefault(e.month, []).append(e)
        context.update({
            'year_rows': [
                {
                    'month': m,
                    'label': MONTH_NAMES_MN[m - 1],
                    'entries': [{'scheme_label': scheme_labels.get(e.scheme, e.scheme), 'total': e.total,
                                 'card': e.card_amount, 'cash': e.cash_amount} for e in by_month[m]],
                    **amounts(by_month[m]),
                }
                for m in sorted(by_month)
            ],
            'months_with_bonus': len(by_month),
            'year_avg': context['year_total']['total'] / len(by_month) if by_month else Decimal(0),
        })
        return render(request, 'shop/my_sales_bonus.html', context)

    entries = [e for e in year_entries if e.month == month]
    settings = sb.get_settings(PayrollSetting.get())
    member = SalesBonusMember.objects.filter(employee_code=code).first() if code else None
    data_cache = {}
    details = []
    for e in entries:
        alloc = None
        m_obj = member if member and member.scheme == e.scheme else None
        if e.scheme == 'storekeeper':
            people = {m.employee_code: (m, a) for m, a in sb.storekeeper_people(year, month, settings)}
            m_obj, alloc = people.get(code, (m_obj, {}))
        if m_obj is None:
            m_obj = SimpleNamespace(employee_code=code, scheme=e.scheme, channels=[], source=None, sort_order=999)
        rng = sb.period_range(year, month, e.scheme, settings)
        if rng not in data_cache:
            data_cache[rng] = sb.month_data(year, month, *rng)
        transfers = sb.transfers_map(year, month, e.scheme) if e.scheme in ('sales_rep', 'cash_seller') else None
        result = sb.compute(m_obj, e.inputs or {}, data_cache[rng], settings, alloc, transfers)
        # Бусад ажилтанд шилжүүлсэн дүнгийн хүлээн авагчийн нэр
        out_codes = {t['code'] for line in result['lines'] for t in line.get('transfers_out') or []}
        names = dict(OpenDataEmployee.objects.filter(id__in=out_codes).values_list('id', 'name')) if out_codes else {}
        for line in result['lines']:
            line['transfers_out'] = [{**t, 'name': names.get(t['code'], t['code'])} for t in line.get('transfers_out') or []]
        details.append({
            'scheme': e.scheme,
            'label': scheme_labels.get(e.scheme, e.scheme),
            'period': sb.period_label(*rng),
            'entry': e,
            'note': (e.inputs or {}).get('note') or '',
            'result': result,
            # Хадгалсны дараа эх өгөгдөл (борлуулалт, буцаалт) өөрчлөгдсөн бол задаргаа олгох дүнгээс зөрж болно
            'differs': result['total'] != e.total,
        })
    context.update({
        'details': details,
        'month_total': amounts(entries),
        'distributor_settings': settings['distributor'],
    })
    return render(request, 'shop/my_sales_bonus.html', context)


@login_required
def menu_permission_config(request):
    """Sidebar цэсийг албан тушаал тус бүрээр харуулах эсэх, мөн бүлэглэлтийг тохируулах хуудас (зөвхөн superuser).

    Албан тушаал нэг нэгээр сонгож эрхийг нь тохируулдаг тул цэс хэдий чинээ
    олон болсон ч зөвхөн мөр нэмэгдэнэ, баганаар өсдөг матриц болохгүй.
    """
    from django.contrib import messages
    from shop.models import MenuPermission, MenuGroupAssignment
    from shop.menuconfig import ALWAYS_ALL_MENU_POSITIONS, get_configurable_menu_items, get_group_choices, get_effective_menu_groups

    if not request.user.is_superuser:
        messages.warning(request, 'Энэ хуудас зөвхөн админд зориулагдсан.')
        return redirect('shop:dashboard')

    positions = list(
        OpenDataEmployee.objects.exclude(positionname__isnull=True)
        .exclude(positionname='')
        .values_list('positionname', flat=True)
        .distinct()
        .order_by('positionname')
    )
    menu_items = get_configurable_menu_items()
    group_choices = get_group_choices()

    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'groups':
            for item in menu_items:
                group_key = request.POST.get(f"group__{item['key']}")
                if group_key:
                    MenuGroupAssignment.objects.update_or_create(
                        menu_key=item['key'],
                        defaults={'group_key': group_key},
                    )
            messages.success(request, 'Цэсний бүлэглэлт амжилттай хадгалагдлаа.')
            return redirect('shop:menu_permission_config')

        if action == 'permissions':
            position = request.POST.get('position', '')
            if position in ALWAYS_ALL_MENU_POSITIONS:
                messages.info(request, f'"{position}" албан тушаал бүх цэсийг автоматаар харах тул тохиргоо хийх шаардлагагүй.')
            elif position in positions:
                for item in menu_items:
                    field_name = f"visible__{item['key']}"
                    MenuPermission.objects.update_or_create(
                        position_name=position,
                        menu_key=item['key'],
                        defaults={'is_visible': field_name in request.POST},
                    )
                messages.success(request, f'"{position}" албан тушаалын цэсний эрх хадгалагдлаа.')
            return redirect(f"{request.path}?position={position}")

    # Идэвхтэй албан тушаал: query param эсвэл жагсаалтын эхнийх
    selected_position = request.GET.get('position') or (positions[0] if positions else '')
    is_always_all_position = selected_position in ALWAYS_ALL_MENU_POSITIONS

    existing_for_position = dict(
        MenuPermission.objects.filter(position_name=selected_position).values_list('menu_key', 'is_visible')
    )

    # Групээр бүлэглэсэн checkbox жагсаалт - шинэ цэс нэмэгдэхэд зөвхөн мөр нэмэгдэнэ, багана биш
    permission_groups = []
    for group in get_effective_menu_groups():
        rows = []
        for item in group['items']:
            if item.get('always_visible'):
                continue
            rows.append({
                'key': item['key'],
                'label': item['label'],
                'icon': item['icon'],
                'field_name': f"visible__{item['key']}",
                # Ерөнхий нягтлан шинэ цэсийг ч оролцуулаад бүгдийг үргэлж харна;
                # бусад албан тушаалд тохиргоогүй цэс default-аар нуугдана (opt-in)
                'is_visible': True if is_always_all_position else existing_for_position.get(item['key'], False),
            })
        if rows:
            permission_groups.append({'label': group['label'] or group['key'], 'rows': rows})

    group_overrides = dict(MenuGroupAssignment.objects.values_list('menu_key', 'group_key'))
    group_rows = [
        {
            'key': item['key'],
            'label': item['label'],
            'icon': item['icon'],
            'field_name': f"group__{item['key']}",
            'current_group_key': group_overrides.get(item['key'], item['group_key']),
        }
        for item in menu_items
    ]

    context = {
        'positions': positions,
        'selected_position': selected_position,
        'is_always_all_position': is_always_all_position,
        'always_all_positions': ALWAYS_ALL_MENU_POSITIONS,
        'permission_groups': permission_groups,
        'group_choices': group_choices,
        'group_rows': group_rows,
    }
    return render(request, 'shop/menu_permission_config.html', context)




@login_required
def meal_list(request):
    """Хоолны бүртгэлийн жагсаалт"""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from shop.models import Meal, MealIngredient, MealNumberAssignment, MealCycleSettings, get_menu_number_for_date
    meals = (
        Meal.objects.prefetch_related('ingredients', 'numbers')
        .annotate(min_number=Min('numbers__number'))
        .order_by('-is_active', 'min_number', 'meal_type', 'name')
    )

    # Хайлт / шүүлт - join-оор min_number annotate гажихаас сэргийлж pk__in ашиглана
    q = request.GET.get('q', '').strip()
    meal_type = request.GET.get('meal_type', '')
    number = request.GET.get('number', '')
    status = request.GET.get('status', '')
    if q:
        # Орцын нэр нь OpenDataItem view-д байгаа тул хоолонд ашиглагдсан кодуудаас нэрээр нь хайна
        used_item_ids = MealIngredient.objects.values_list('item_id', flat=True).distinct()
        matched_item_ids = list(
            OpenDataItem.objects.filter(id__in=list(used_item_ids), name__icontains=q)
            .values_list('id', flat=True)
        )
        meals = meals.filter(
            Q(name__icontains=q)
            | Q(description__icontains=q)
            | Q(pk__in=MealIngredient.objects.filter(
                Q(item_id__icontains=q) | Q(item_id__in=matched_item_ids)
            ).values('meal_id'))
        )
    if meal_type in dict(Meal.MEAL_TYPE_CHOICES):
        meals = meals.filter(meal_type=meal_type)
    if number.isdigit():
        meals = meals.filter(pk__in=MealNumberAssignment.objects.filter(number=int(number)).values('meal_id'))
    if status == 'active':
        meals = meals.filter(is_active=True)
    elif status == 'inactive':
        meals = meals.filter(is_active=False)

    today = timezone.localdate()
    cycle_effective_from = MealCycleSettings.get_effective_from()
    context = {
        'meals': meals,
        'q': q,
        'filter_meal_type': meal_type,
        'filter_number': number,
        'filter_status': status,
        'meal_type_choices': Meal.MEAL_TYPE_CHOICES,
        'is_filtered': bool(q or meal_type or number or status),
        'cycle_effective_from': cycle_effective_from,
        'cycle_effective_number': get_menu_number_for_date(cycle_effective_from) if cycle_effective_from else None,
        'today': today,
        'today_menu_number': get_menu_number_for_date(today),
        'number_choices': range(1, 11),
    }
    return render(request, 'shop/meal_list.html', context)


@login_required
def meal_cycle_settings_update(request):
    """Сонгосон огноог өгсөн хоолны дугаартай тохируулж, циклийг шинэчилнэ.

    Тохируулсан огнооноос хойшхи бүх ажлын өдрүүдэд шинэ дараалал (тухайн
    дугаараас цааш 1-10 хүртэл эргэлдэх) хүчинтэй болно.
    """
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from django.contrib import messages
    from shop.models import MealCycleSettings

    if request.method == 'POST':
        date_str = request.POST.get('date')
        number_str = request.POST.get('number')
        try:
            target_date = datetime.strptime(date_str, '%Y-%m-%d').date()
            number = int(number_str)
            if target_date.weekday() >= 5:
                raise ValueError('weekend')
            if not (1 <= number <= 10):
                raise ValueError('number out of range')
        except (TypeError, ValueError):
            messages.error(request, 'Огноо, дугаарыг зөв сонгоно уу (амралтын өдөр байж болохгүй, дугаар 1-10 хооронд байна).')
            return redirect('shop:meal_list')

        MealCycleSettings.set_number_for_date(target_date, number)
        messages.success(
            request,
            f'{target_date} өдрийг {number} дугаартай болгож циклийг шинэчиллээ. '
            f'Энэ өдрөөс хойшхи бүх ажлын өдрүүдэд дараалал үргэлжилнэ.'
        )

    return redirect('shop:meal_list')


@login_required
def meal_material_calc(request):
    """Сонгосон хугацаанд ажилтнуудад хэрэгтэй орцны материалын нийт хэмжээг тооцоолно.

    Тухайн хугацаанд ажлын өдөр бүрийн хоолны дугаар хэд удаа давтагдахыг тоолж,
    орц тус бүрийн (нэг хүнд хэмжээ × давтагдсан тоо × хүний тоо)-г нэгтгэнэ.
    """
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from collections import Counter
    from shop.models import Meal, get_menu_number_for_date

    today = timezone.localdate()
    date_from_str = request.GET.get('date_from')
    date_to_str = request.GET.get('date_to')
    headcount_str = request.GET.get('headcount')

    try:
        date_from = datetime.strptime(date_from_str, '%Y-%m-%d').date() if date_from_str else today
    except ValueError:
        date_from = today
    try:
        date_to = datetime.strptime(date_to_str, '%Y-%m-%d').date() if date_to_str else today + timedelta(days=29)
    except ValueError:
        date_to = today + timedelta(days=29)
    if date_to < date_from:
        date_from, date_to = date_to, date_from

    try:
        headcount = int(headcount_str)
        if headcount < 0:
            headcount = 0
    except (TypeError, ValueError):
        headcount = 0

    # Тухайн хугацаанд ажлын өдөр бүрийн хоолны дугаарыг тоолно (амралтын өдөр/
    # тохируулаагүй өдөр None буцаах тул тоологдохгүй)
    number_counts = Counter()
    dates_by_number = {}
    d = date_from
    while d <= date_to:
        n = get_menu_number_for_date(d)
        if n:
            number_counts[n] += 1
            dates_by_number.setdefault(n, []).append(d)
        d += timedelta(days=1)

    # Тухайн хугацаанд давтагдсан хоолнуудыг төрлөөр нь (1-р/2-р хоол, хачир)
    # тусдаа хүснэгт болгож нэрээр нь харуулна. Нэг хоол 2 өөр дугаартай байсан ч
    # нэр нь ижил тул давхардуулахгүй, тоог нь нэгтгэж нэг мөр болгоно.
    meals = list(
        Meal.objects.filter(Meal.visible_on_q(date_from), numbers__number__in=number_counts.keys())
        .prefetch_related('ingredients', 'numbers')
        .distinct()
        .order_by('meal_type', 'name')
    ) if number_counts else []

    # Идэвхгүй болсон хоолыг зөвхөн идэвхгүй болохоос өмнөх өдрүүдэд тоолно
    occurrence_by_meal = {
        meal.pk: sum(
            1 for n in meal.number_list for day in dates_by_number.get(n, [])
            if meal.is_visible_on(day)
        )
        for meal in meals
    }

    repetition_by_type = {'1': [], '2': [], '3': []}
    for meal in meals:
        occurrence_count = occurrence_by_meal[meal.pk]
        if occurrence_count == 0:
            continue
        repetition_by_type.setdefault(meal.meal_type, []).append({
            'name': meal.name, 'count': occurrence_count,
        })
    for rows in repetition_by_type.values():
        rows.sort(key=lambda r: (-r['count'], r['name']))

    # Ажлын өдөр бүрийн хоолны хуваарь (огноо, 2-р хоол, 1-р хоол)
    meals_by_number = {}
    for meal in meals:
        for n in meal.number_list:
            meals_by_number.setdefault(n, []).append(meal)
    daily_rows = []
    for n, days in dates_by_number.items():
        for day in days:
            day_meals = [m for m in meals_by_number.get(n, []) if m.is_visible_on(day)]
            daily_rows.append({
                'date': day,
                'number': n,
                'meal_2': ', '.join(m.name for m in day_meals if m.meal_type == '2'),
                'meal_1': ', '.join(m.name for m in day_meals if m.meal_type == '1'),
            })
    daily_rows.sort(key=lambda r: r['date'])

    detail_rows = []
    summary = {}

    if headcount > 0 and number_counts:
        for meal in meals:
            occurrence_count = occurrence_by_meal[meal.pk]
            if occurrence_count == 0:
                continue
            for ing in meal.ingredients.all():
                per_person = ing.quantity or 0
                total = float(per_person) * occurrence_count * headcount
                ingredient_name = ing.item.name if ing.item else ing.item_id
                detail_rows.append({
                    'meal_type': meal.get_meal_type_display(),
                    'meal_name': meal.name,
                    'ingredient': ingredient_name,
                    'per_person': per_person,
                    'unit': ing.unit,
                    'occurrence_count': occurrence_count,
                    'total': total,
                })
                key = ing.item_id
                if key not in summary:
                    summary[key] = {'label': ingredient_name, 'unit': ing.unit, 'total': 0.0}
                summary[key]['total'] += total

    summary_rows = []
    for key, data in sorted(summary.items(), key=lambda kv: kv[1]['label']):
        unit = (data['unit'] or '').strip().lower()
        converted = None
        converted_unit = None
        if unit in ('гр', 'г'):
            converted = data['total'] / 1000
            converted_unit = 'кг'
        elif unit == 'мл':
            converted = data['total'] / 1000
            converted_unit = 'л'
        summary_rows.append({
            'label': data['label'],
            'unit': data['unit'],
            'total': data['total'],
            'converted': converted,
            'converted_unit': converted_unit,
        })

    context = {
        'date_from': date_from,
        'date_to': date_to,
        'headcount': headcount,
        'repetition_1': repetition_by_type.get('1', []),
        'repetition_2': repetition_by_type.get('2', []),
        'repetition_3': repetition_by_type.get('3', []),
        'has_repetitions': any(repetition_by_type.values()),
        'daily_rows': daily_rows,
        'detail_rows': detail_rows,
        'summary_rows': summary_rows,
    }
    return render(request, 'shop/meal_material_calc.html', context)


@login_required
def meal_edit(request, pk=None):
    """Хоол бүртгэх / засах (орцын хамт)"""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from django.contrib import messages
    from django.shortcuts import get_object_or_404
    from shop.models import Meal, MealNumberAssignment
    from shop.forms import MealForm, MealIngredientFormSet, MealNumberFormSet

    meal = get_object_or_404(Meal, pk=pk) if pk else Meal()

    if request.method == 'POST':
        form = MealForm(request.POST, instance=meal)
        formset = MealIngredientFormSet(request.POST, instance=meal)
        number_formset = MealNumberFormSet(request.POST, instance=meal)
        if form.is_valid() and formset.is_valid() and number_formset.is_valid():
            meal_type = form.cleaned_data['meal_type']
            conflict = False
            for number_form in number_formset.forms:
                if not number_form.cleaned_data or number_form.cleaned_data.get('DELETE'):
                    continue
                number = number_form.cleaned_data.get('number')
                if number is None:
                    continue
                clash = MealNumberAssignment.objects.filter(
                    meal__meal_type=meal_type, meal__is_active=True, number=number,
                )
                if meal.pk:
                    clash = clash.exclude(meal=meal)
                other_meal = clash.select_related('meal').first()
                if other_meal:
                    number_form.add_error('number', f'{number} дугаарыг "{other_meal.meal.name}" хоол аль хэдийн эзэлсэн байна.')
                    conflict = True

            if not conflict:
                meal = form.save()
                formset.instance = meal
                formset.save()
                number_formset.instance = meal
                number_formset.save()

                messages.success(request, 'Хоол амжилттай хадгалагдлаа.')
                return redirect('shop:meal_list')
    else:
        form = MealForm(instance=meal)
        formset = MealIngredientFormSet(instance=meal)
        number_formset = MealNumberFormSet(instance=meal)

    context = {
        'form': form,
        'formset': formset,
        'number_formset': number_formset,
        'meal': meal if meal.pk else None,
    }
    return render(request, 'shop/meal_form.html', context)


@login_required
def meal_delete(request, pk):
    """Хоол устгах"""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from django.contrib import messages
    from django.shortcuts import get_object_or_404
    from shop.models import Meal

    meal = get_object_or_404(Meal, pk=pk)
    if request.method == 'POST':
        meal.delete()
        messages.success(request, 'Хоол устгагдлаа.')
        return redirect('shop:meal_list')

    context = {'meal': meal}
    return render(request, 'shop/meal_confirm_delete.html', context)


@login_required
def meal_toggle_active(request, pk):
    """Хоолыг идэвхгүй болгох / дахин идэвхжүүлэх"""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from django.contrib import messages
    from django.shortcuts import get_object_or_404
    from shop.models import Meal, MealNumberAssignment

    meal = get_object_or_404(Meal, pk=pk)
    if request.method == 'POST':
        if meal.is_active:
            meal.is_active = False
            meal.deactivated_at = timezone.localdate()
            meal.save(update_fields=['is_active', 'deactivated_at', 'updated_at'])
            messages.success(request, f'"{meal.name}" хоолыг идэвхгүй болголоо.')
        else:
            # Идэвхгүй байх хугацаанд дугаарыг нь өөр хоол эзэлсэн бол идэвхжүүлэхгүй
            clash = (
                MealNumberAssignment.objects
                .filter(meal__meal_type=meal.meal_type, meal__is_active=True, number__in=meal.number_list)
                .exclude(meal=meal)
                .select_related('meal')
                .first()
            )
            if clash:
                messages.error(
                    request,
                    f'{clash.number} дугаарыг "{clash.meal.name}" хоол эзэлсэн тул идэвхжүүлэх боломжгүй. '
                    f'Эхлээд дугаарыг нь засна уу.'
                )
            else:
                meal.is_active = True
                meal.deactivated_at = None
                meal.save(update_fields=['is_active', 'deactivated_at', 'updated_at'])
                messages.success(request, f'"{meal.name}" хоолыг идэвхжүүллээ.')
    return redirect('shop:meal_list')


@login_required
def meal_attendance_list(request):
    """Хоол идэх нэрсийн бүртгэлийн жагсаалт"""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    import calendar
    from shop.models import MealAttendance

    # Анхдагчаар тухайн сарын бүртгэлийг харуулна
    today = timezone.localdate()
    month_start = today.replace(day=1)
    month_end = today.replace(day=calendar.monthrange(today.year, today.month)[1])
    try:
        date_from = datetime.strptime(request.GET.get('date_from', ''), '%Y-%m-%d').date()
    except ValueError:
        date_from = month_start
    try:
        date_to = datetime.strptime(request.GET.get('date_to', ''), '%Y-%m-%d').date()
    except ValueError:
        date_to = month_end
    if date_to < date_from:
        date_from, date_to = date_to, date_from

    status = request.GET.get('status', '')
    if status not in ('', 'changed', 'unchanged'):
        status = ''
    q = request.GET.get('q', '').strip()

    attendances = MealAttendance.objects.filter(date__gte=date_from, date__lte=date_to)
    if status == 'changed':
        attendances = attendances.filter(ingredients_changed=True)
    elif status == 'unchanged':
        attendances = attendances.filter(ingredients_changed=False)
    if q:
        attendances = attendances.filter(
            Q(names__name__icontains=q) | Q(change_comment__icontains=q)
        ).distinct()
    attendances = list(
        attendances.annotate(headcount=Count('names', distinct=True)).order_by('-date')
    )

    context = {
        'attendances': attendances,
        'date_from': date_from,
        'date_to': date_to,
        'status': status,
        'q': q,
        'total_headcount': sum(a.headcount for a in attendances),
        'is_filtered': bool(status or q) or (date_from, date_to) != (month_start, month_end),
    }
    return render(request, 'shop/meal_attendance_list.html', context)


@login_required
def meal_attendance_summary(request):
    """Сонгосон огнооны хооронд хоол идсэн хүмүүс/орцын нэгтгэл (огноогоор багана болгосон хүснэгт)"""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from shop.models import MealAttendance

    view_mode = request.GET.get('view', 'ingredients')
    if view_mode not in ('ingredients', 'people'):
        view_mode = 'ingredients'

    today = timezone.localdate()
    date_from_str = request.GET.get('date_from')
    date_to_str = request.GET.get('date_to')
    try:
        date_from = datetime.strptime(date_from_str, '%Y-%m-%d').date() if date_from_str else today - timedelta(days=6)
    except ValueError:
        date_from = today - timedelta(days=6)
    try:
        date_to = datetime.strptime(date_to_str, '%Y-%m-%d').date() if date_to_str else today
    except ValueError:
        date_to = today

    attendances = list(
        MealAttendance.objects.filter(date__gte=date_from, date__lte=date_to)
        .prefetch_related('names')
        .order_by('date')
    )
    dates = [a.date for a in attendances]

    row_order = []
    row_labels = {}
    row_units = {}
    cell_values = {}

    for attendance in attendances:
        headcount = len(attendance.names.all())
        if view_mode == 'people':
            for person in attendance.names.all():
                name = person.name.strip()
                if not name:
                    continue
                if name not in row_labels:
                    row_labels[name] = name
                    row_units[name] = ''
                    row_order.append(name)
                    cell_values[name] = {}
                cell_values[name][attendance.date] = cell_values[name].get(attendance.date, 0) + 1
        else:
            for meal in attendance.meals:
                for ingredient in meal.ingredients.all():
                    key = ingredient.item_id
                    if key not in row_labels:
                        item = ingredient.item
                        row_labels[key] = item.name if item else ingredient.item_id
                        row_units[key] = ingredient.unit
                        row_order.append(key)
                        cell_values[key] = {}
                    qty = float(ingredient.quantity or 0) * headcount
                    cell_values[key][attendance.date] = cell_values[key].get(attendance.date, 0) + qty

    row_order.sort(key=lambda k: row_labels[k])

    column_totals = {d: 0 for d in dates}
    rows = []
    for key in row_order:
        cells = []
        row_total = 0
        for d in dates:
            value = cell_values[key].get(d)
            cells.append(value)
            if value:
                row_total += value
                column_totals[d] += value
        rows.append({
            'label': row_labels[key],
            'unit': row_units.get(key, ''),
            'cells': cells,
            'total': row_total,
        })

    context = {
        'view_mode': view_mode,
        'date_from': date_from,
        'date_to': date_to,
        'dates': dates,
        'rows': rows,
        'column_totals': [column_totals[d] for d in dates],
        'grand_total': sum(column_totals.values()),
    }
    return render(request, 'shop/meal_attendance_summary.html', context)


@login_required
def meal_attendance_form(request, pk=None):
    """Хоол идэх нэрсийн бүртгэл - огноо сонгоход тухайн өдрийн хоол/орц автоматаар харагдаж, хүмүүсийн нэрсийг бүртгэнэ"""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from django.contrib import messages
    from django.shortcuts import get_object_or_404
    from django.forms import inlineformset_factory
    from shop.models import MealAttendance, MealAttendanceName, MealAttendanceIngredient
    from shop.forms import (
        MealAttendanceForm, MealAttendanceNameFormSet, MealAttendanceNameForm,
        MealAttendanceIngredientForm,
    )

    IngredientFormSet = inlineformset_factory(
        MealAttendance, MealAttendanceIngredient,
        form=MealAttendanceIngredientForm, extra=0, can_delete=True,
    )

    attendance = get_object_or_404(MealAttendance, pk=pk) if pk else None

    if request.method == 'POST':
        if not attendance:
            posted_date = request.POST.get('date')
            try:
                posted_date = datetime.strptime(posted_date, '%Y-%m-%d').date()
            except (TypeError, ValueError):
                messages.error(request, 'Огноо буруу байна.')
                return redirect('shop:meal_attendance_list')
            attendance, _ = MealAttendance.objects.get_or_create(date=posted_date)

        form = MealAttendanceForm(request.POST, instance=attendance)
        formset = MealAttendanceNameFormSet(request.POST, instance=attendance)
        ingredient_formset = IngredientFormSet(request.POST, instance=attendance)

        if form.is_valid() and formset.is_valid() and ingredient_formset.is_valid():
            meals_before_save = list(attendance.meals)

            attendance = form.save()
            formset.instance = attendance
            formset.save()
            ingredient_formset.instance = attendance
            ingredient_formset.save()

            after_snapshot = {
                (o.meal_id, o.item_id, o.quantity, o.unit)
                for o in attendance.ingredient_overrides.all()
            }
            template_snapshot = {
                (m.pk, ing.item_id, ing.quantity, ing.unit)
                for m in meals_before_save for ing in m.ingredients.all()
            }
            changed = after_snapshot != template_snapshot
            if changed != attendance.ingredients_changed:
                attendance.ingredients_changed = changed
                attendance.save(update_fields=['ingredients_changed'])

            messages.success(request, 'Хоол идэх нэрсийн бүртгэл хадгалагдлаа.')
            return redirect('shop:meal_attendance_edit', pk=attendance.pk)
    else:
        if not attendance:
            date_param = request.GET.get('date')
            try:
                selected_date = datetime.strptime(date_param, '%Y-%m-%d').date() if date_param else timezone.localdate()
            except ValueError:
                selected_date = timezone.localdate()

            existing = MealAttendance.objects.filter(date=selected_date).first()
            if existing:
                return redirect('shop:meal_attendance_edit', pk=existing.pk)
            attendance = MealAttendance(date=selected_date)

            # Шинэ бүртгэл нээхэд өмнөх хамгийн сүүлийн өдрийн нэрсийг урьдчилан бөглөнө -
            # тэндээсээ засах/нэмэх/хасах хийхэд хялбар байх зорилготой
            previous = (
                MealAttendance.objects.filter(date__lt=selected_date)
                .prefetch_related('names')
                .order_by('-date')
                .first()
            )
            initial_names = [{'name': n.name} for n in previous.names.all()] if previous else []

            NameFormSet = inlineformset_factory(
                MealAttendance, MealAttendanceName,
                form=MealAttendanceNameForm, extra=len(initial_names) + 1, can_delete=True,
            )
            formset = NameFormSet(instance=attendance, initial=initial_names)
        else:
            formset = MealAttendanceNameFormSet(instance=attendance)

        form = MealAttendanceForm(instance=attendance)

        # Орцны formset - тухайн өдөрт хадгалагдсан хувилбар байвал түүнийг, үгүй бол
        # стандарт жороос (Meal.ingredients) урьдчилан бөглөж харуулна
        meals = list(attendance.meals) if attendance.date else []
        existing_override_meal_ids = set(
            MealAttendanceIngredient.objects.filter(attendance=attendance).values_list('meal_id', flat=True)
        ) if attendance.pk else set()

        initial_ingredients = []
        for meal in meals:
            if meal.pk in existing_override_meal_ids:
                continue
            for ing in meal.ingredients.all():
                initial_ingredients.append({
                    'meal': meal.pk, 'item_id': ing.item_id, 'quantity': ing.quantity, 'unit': ing.unit,
                })

        ExtraIngredientFormSet = inlineformset_factory(
            MealAttendance, MealAttendanceIngredient,
            form=MealAttendanceIngredientForm, extra=len(initial_ingredients), can_delete=True,
        )
        ingredient_formset = ExtraIngredientFormSet(
            instance=attendance if attendance.pk else None, initial=initial_ingredients,
        )

        # Байгаа орцонд зөвхөн нэрийг нь текстээр харуулна (сонголтыг зөвхөн шинэ
        # орц нэмэхэд ашиглана) - нэрийг энд урьдчилан тооцоолж form дээр наана
        from datamigration.models import OpenDataItem
        item_ids = {
            (f.instance.item_id if f.instance.pk else f.initial.get('item_id'))
            for f in ingredient_formset.forms
        }
        item_ids.discard(None)
        item_names = dict(OpenDataItem.objects.filter(pk__in=item_ids).values_list('id', 'name'))
        for ing_form in ingredient_formset.forms:
            current_item_id = ing_form.instance.item_id if ing_form.instance.pk else ing_form.initial.get('item_id')
            ing_form.display_name = item_names.get(current_item_id, current_item_id)

    # Орцны form-уудыг хоол тус бүрээр нь бүлэглэж, загварт харуулахад бэлтгэнэ
    meals_for_context = list(attendance.meals) if attendance.date else []
    forms_by_meal = {}
    for ing_form in ingredient_formset.forms:
        meal_id = ing_form.instance.meal_id or ing_form.initial.get('meal')
        forms_by_meal.setdefault(meal_id, []).append(ing_form)
    meal_ingredient_groups = [(meal, forms_by_meal.get(meal.pk, [])) for meal in meals_for_context]

    context = {
        'form': form,
        'formset': formset,
        'ingredient_formset': ingredient_formset,
        'meal_ingredient_groups': meal_ingredient_groups,
        'attendance': attendance if attendance.pk else None,
        'selected_date': attendance.date,
        'menu_number': attendance.menu_number,
        'meals': meals_for_context,
        'is_weekend': attendance.date.weekday() >= 5,
    }
    return render(request, 'shop/meal_attendance_form.html', context)


@login_required
def meal_attendance_delete(request, pk):
    """Хоол идэх нэрсийн бүртгэл устгах"""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from django.contrib import messages
    from django.shortcuts import get_object_or_404
    from shop.models import MealAttendance

    attendance = get_object_or_404(MealAttendance, pk=pk)
    if request.method == 'POST':
        attendance.delete()
        messages.success(request, 'Бүртгэл устгагдлаа.')
        return redirect('shop:meal_attendance_list')

    context = {'attendance': attendance}
    return render(request, 'shop/meal_attendance_confirm_delete.html', context)


def _lookup_product_stock(query):
    """Чөлөөт текстээс утга бүхий үгсийг гаргаж аваад InventorySnapshot-с тохирох
    эхний 5 барааг үлдэгдлийн хамт буцаана. AI чат болон public JSON endpoint хоёулаа
    ашиглана.
    """
    import re
    from shop.models_inventory import InventorySnapshot
    from django.db.models import Q

    query = (query or "").strip()
    if not query:
        return []

    # 3-с дээш тэмдэгттэй үгсийг л ашиглана (богино холбогч үгсийг үл тооцно)
    words = [w for w in re.split(r"\s+", query) if len(w) >= 3][:8]
    if not words:
        return []

    q_filter = Q()
    for word in words:
        q_filter |= Q(itemname__icontains=word)

    snapshots = InventorySnapshot.objects.filter(q_filter).order_by("-total_qty")[:5]

    return [
        {
            "name": s.itemname,
            "brand": s.brandname or "",
            "qty": float(s.total_qty),
            "in_stock": s.total_qty > 0,
        }
        for s in snapshots
    ]


AI_CHAT_SYSTEM_PROMPT = (
    "Чи бол Төгс Амин Эрдэнэ ХХК компанийн вэбсайт дээрх AI туслах. Энэ компани нь Европын "
    "(Итали, Франц, Герман, Испани, Грек гэх мэт) шилдэг брэндүүдийн хүнсний бүтээгдэхүүнийг "
    "Монгол улс руу импортолж, хувь хүн (B2C) болон ресторан/дэлгүүр зэрэг бизнес түншүүдэд "
    "(B2B) борлуулдаг. Чиний цорын ганц үүрэг бол ЗӨВХӨН энэ компани болон түүний "
    "бүтээгдэхүүний талаарх асуултад хариулах. Хэрэв хэрэглэгч компани эсвэл бүтээгдэхүүнтэй "
    "огт холбоогүй өөр сэдвээр асуувал, өөр ямар ч мэдээлэл өгөлгүй, зөвхөн дараах хариултыг "
    "эелдэгээр өг: 'Уучлаарай, би зөвхөн манай байгууллага болон бүтээгдэхүүний талаарх таны "
    "сонирхсон асуултад хариулах боломжтой.' Хэрэв хэрэглэгч тодорхой барааны нөөц/үлдэгдлийн "
    "талаар асуувал, доор өгөгдсөн Барааны нөөцийн систем-ийн мэдээллийг ҮНДЭСЛЭН тухайн бараа "
    "НӨӨЦТЭЙ эсвэл НӨӨЦГҮЙ гэдгийг л тодорхой хэлээрэй, тоо ширхэгийг хэзээ ч бүү дурьд, өөрөө "
    "таамаглаж бүү хариул. Хэрэв тохирох бараа олдоогүй бол уучлалт хүсээд тодорхой барааны "
    "нэрийг лавлаарай. Хэрэглэгч заримдаа Монгол үгээ Латин үсгээр галиглаж бичдэг (жишээ нь: "
    "'sain bnu', 'unetei bhuu', 'ta yamar unaa hudaldj bgaa' гэх мэт) - ийм тохиолдолд бичсэн "
    "үгийн Монгол хэл дээрх утгыг ойлгож, хариултаа энгийн монгол (кирилл) хэлээр өгнө үү. "
    "Хариултаа товч, эелдэг, монгол хэлээр өгнө үү."
)


def ai_chat_message(request):
    """AI чат - Anthropic Claude API-г Django-с шууд дуудна (n8n дундын шат ашигладаггүй)."""
    if request.method != "POST":
        return JsonResponse({"error": "POST хүсэлт байх ёстой"}, status=405)

    import json
    import logging
    import requests
    from django.conf import settings

    logger = logging.getLogger(__name__)

    try:
        payload = json.loads(request.body or "{}")
    except json.JSONDecodeError:
        return JsonResponse({"error": "Буруу бус JSON"}, status=400)

    message = (payload.get("message") or "").strip()
    if not message:
        return JsonResponse({"error": "Зурвас хоосон байна"}, status=400)

    stock_results = _lookup_product_stock(message)
    if stock_results:
        lines = [
            f"- {r['name']}: {'НӨӨЦТЭЙ' if r['in_stock'] else 'НӨӨЦГҮЙ'}"
            for r in stock_results
        ]
        stock_context = (
            "Барааны нөөцийн систем (эх сурвалж, найдвартай) дараах тохирох бараануудыг олсон:\n"
            + "\n".join(lines)
        )
    else:
        stock_context = "Барааны нөөцийн системээс энэ асуултад тохирох бараа олдсонгүй."

    body = {
        "model": "claude-sonnet-5",
        "max_tokens": 1024,
        "system": AI_CHAT_SYSTEM_PROMPT,
        "messages": [
            {"role": "user", "content": f"{stock_context}\n\nХэрэглэгчийн асуулт: {message}"}
        ],
    }
    headers = {
        "x-api-key": settings.ANTHROPIC_API_KEY,
        "anthropic-version": "2023-06-01",
        "content-type": "application/json",
    }

    # Anthropic талд заримдаа тохиолддог түр зуурын сүлжээний алдааг даван туулахын тулд нэг удаа дахин оролдоно
    data = None
    last_error = None
    for attempt in range(2):
        try:
            response = requests.post(
                "https://api.anthropic.com/v1/messages",
                json=body,
                headers=headers,
                timeout=30,
            )
            response.raise_for_status()
            response.encoding = "utf-8"
            data = response.json()
            break
        except (requests.RequestException, ValueError) as exc:
            last_error = exc

    if data is None:
        logger.exception("Anthropic API руу хүсэлт 2 оролдлогоор ч амжилтгүй боллоо", exc_info=last_error)
        return JsonResponse({"error": "AI чат үйлчилгээ түр боломжгүй байна"}, status=502)

    reply = next(
        (c.get("text", "") for c in data.get("content", []) if c.get("type") == "text"),
        "",
    )
    return JsonResponse({"reply": reply})


def product_stock_lookup(request):
    """AI чат (n8n)-с барааны үлдэгдэл асуухад ашиглах public JSON endpoint.

    ?q= параметрт ирсэн чөлөөт текстээс утга бүхий үгсийг гаргаж аваад,
    InventorySnapshot-с тохирох эхний 5 барааг үлдэгдлийн хамт буцаана.
    """
    query = request.GET.get("q") or ""
    return JsonResponse({"results": _lookup_product_stock(query)})


def _attendance_filter(request):
    """Цаг бүртгэлийн шүүлтийг буцаана: (date_from, date_to, name_query).
    Анхдагч: тухайн сарын 1-нээс өнөөдөр хүртэл, нэрийн шүүлтгүй (name_query='').

    Шүүлтийг session-д түлхүүр тус бүрээр хадгална: URL-д ирсэн параметрээр шинэчилж, ирээгүйг нь
    (refresh, засварын дараах redirect зөвхөн огноотой ирэх, цэснээс орох, Excel татах) сүүлд шүүсэн
    утгаар нөхнө. Шинээр нэвтрэхэд session цэвэрлэгдэнэ."""
    today = timezone.localdate()
    saved_filter = dict(request.session.get('attendance_filter') or {})
    # missing: '1' бол ажиллаагүй өдрүүдийг харуулна (хуудсан дээр JS-ээр асааж/унтраана);
    # age_min/age_max: насны шүүлт (_age_range)
    for key in ('date_from', 'date_to', 'q', 'missing', 'age_min', 'age_max'):
        if key in request.GET:
            saved_filter[key] = request.GET.get(key)
    request.session['attendance_filter'] = saved_filter
    date_from_str = saved_filter.get('date_from')
    date_to_str = saved_filter.get('date_to')
    name_query = (saved_filter.get('q') or '').strip()
    try:
        date_from = datetime.strptime(date_from_str, '%Y-%m-%d').date() if date_from_str else today.replace(day=1)
    except ValueError:
        date_from = today.replace(day=1)
    try:
        date_to = datetime.strptime(date_to_str, '%Y-%m-%d').date() if date_to_str else today
    except ValueError:
        date_to = today
    return date_from, date_to, name_query


def _matches_name_query(row, name_query):
    """Цаг бүртгэлийн мөр нэрийн шүүлтэд (нэр эсвэл ажилтны кодын хэсэг, том жижиг үсэг ялгахгүй) тохирох эсэх -
    хуудасны JS шүүлттэй ижил дүрэм."""
    if not name_query:
        return True
    needle = name_query.lower()
    return needle in (row.get('name') or '').lower() or needle in (row.get('employee_code') or '').lower()


@login_required
def attendance_filter_save(request):
    """Цаг бүртгэлийн нэрийн шүүлт, 'ажиллаагүй өдрийг харуулах' сонголтыг (хуудсан дээр өөрчлөх бүрт JS-ээс
    дуудна) session-д хадгална - ингэснээр refresh, засварын дараа ч хэвээр үлдэнэ."""
    from django.http import HttpResponse
    if request.user.is_staff and any(k in request.GET for k in ('q', 'missing', 'age_min', 'age_max')):
        _attendance_filter(request)
    return HttpResponse(status=204)


def _age_range(request, session_key):
    """Session-д хадгалсан насны шүүлтийг (age_min, age_max) бүхэл тоогоор буцаана - хоосон/буруу бол None.
    Цаг бүртгэлд _attendance_filter, Хоног бүртгэлд _timesheet_month шүүлтийг session-д хадгалсны дараа дуудна."""
    saved = request.session.get(session_key) or {}

    def as_int(value):
        try:
            return int(str(value).strip()) if str(value or '').strip() else None
        except ValueError:
            return None

    return as_int(saved.get('age_min')), as_int(saved.get('age_max'))


def _timesheet_month(request):
    """Хоног бүртгэлийн он/сарыг буцаана. Шүүлтийг (он/сар, насны шүүлт) session-д түлхүүр тус бүрээр
    хадгална - гараар сонгосон үед шинэчилж, параметргүй (refresh, цэснээс орох, Excel татах) үед сүүлд
    сонгосон утгаа ашиглана."""
    saved_filter = dict(request.session.get('timesheet_filter') or {})
    source = request.POST if request.method == 'POST' else request.GET
    if 'year' not in source and 'month' not in source:
        source = saved_filter
    year, month = _month_from_request(source)
    if request.method == 'GET':
        saved_filter.update({'year': str(year), 'month': str(month)})
        for key in ('age_min', 'age_max'):
            if key in request.GET:
                saved_filter[key] = request.GET.get(key)
        request.session['timesheet_filter'] = saved_filter
    return year, month


def _age_range_label(age_min, age_max):
    """Excel-ийн гарчигт насны шүүлтийг тэмдэглэх (жиш: ' · нас: 60-аас дээш')."""
    if age_min is not None and age_max is not None:
        return f' · нас: {age_min}-{age_max}'
    if age_min is not None:
        return f' · нас: {age_min}-аас дээш'
    if age_max is not None:
        return f' · нас: {age_max} хүртэл'
    return ''


def _filter_timesheet_by_age(request, rows):
    """Хоног бүртгэлийн мөрүүдийг session-д хадгалсан насны шүүлтээр шүүнэ: (rows, age_min, age_max, ages)."""
    from shop.services.attendance import age_in_range, get_employee_ages

    age_min, age_max = _age_range(request, 'timesheet_filter')
    ages = get_employee_ages()
    rows = [r for r in rows if age_in_range(ages.get(r['employee_code']), age_min, age_max)]
    return rows, age_min, age_max, ages


def _xlsx_response(sheet_title, title, headers, rows, filename, widths, totals=None, number_cols=(), row_fills=None):
    """Энгийн хүснэгтийг Excel (.xlsx) файлаар татуулна: 1-р мөрөнд гарчиг, 3-р мөрөнд толгой, дараа нь мөрүүд,
    totals өгвөл төгсгөлд нь тодруулсан нийт мөр. row_fills - мөр бүрийн дэвсгэр өнгө (hex эсвэл None)."""
    import openpyxl
    from io import BytesIO
    from django.http import HttpResponse
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=12)

    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_row = 3
    for col_idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=text)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='F3F4F6', end_color='F3F4F6', fill_type='solid')
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    def write(row_idx, values, fill=None, bold=False):
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = border
            if col_idx in number_cols:
                cell.number_format = '#,##0' if isinstance(value, int) else '#,##0.00'
            if fill:
                cell.fill = PatternFill(start_color=fill, end_color=fill, fill_type='solid')
            if bold:
                cell.font = Font(bold=True)

    row_idx = header_row
    for i, values in enumerate(rows):
        row_idx += 1
        write(row_idx, values, fill=row_fills[i] if row_fills else None)
    if totals:
        row_idx += 1
        write(row_idx, totals, fill='F3F4F6', bold=True)

    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _excel_number(value):
    """Decimal-ыг Excel-д тоогоор бичнэ (бүхэл бол int)."""
    value = Decimal(value or 0)
    return int(value) if value == value.to_integral_value() else float(value)


@login_required
def attendance_export(request):
    """Цаг бүртгэлийн хүснэгтийг (хуудсан дээрх огнооны мужаар) Excel файлаар татна."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from shop.services.attendance import age_in_range, build_attendance_rows, get_employee_ages

    date_from, date_to, name_query = _attendance_filter(request)
    age_min, age_max = _age_range(request, 'attendance_filter')
    ages = get_employee_ages()
    rows = [
        r for r in build_attendance_rows(date_from, date_to)
        if _matches_name_query(r, name_query) and age_in_range(ages.get(r['employee_code']), age_min, age_max)
    ]
    headers = [
        'Хурууны код', 'Ажилтны код', 'Нэр', 'Албан тушаал', 'Огноо', 'Гариг', 'Ирсэн цаг', 'Тарсан цаг',
        '1-10 минут', '11-20 минут', '20-с дээш минут', 'Нийт хоцорсон минут', 'Харилцагчийн нэр',
    ]
    data = [[
        r['device_user_id'] or None, r['employee_code'] or None, r['name'],
        (f"{r['position_name']} (хавсарсан)" if r['is_additional'] else r['position_name']) or None,
        r['date'].strftime('%Y-%m-%d'), r['weekday'], r['arrived_time'] or None, r['left_time'] or None,
        r['late_1_10'], r['late_11_20'], r['late_20_plus'], r['total_late_minutes'], r['customer_name'] or None,
    ] for r in rows]
    # Хуудастай ижил: гараар засварласан мөр шар, амралтын өдөр саарал
    fills = [
        'FEFCE8' if r['is_saved'] else ('F3F4F6' if r['weekday'] in ('Бямба', 'Ням') else None)
        for r in rows
    ]
    totals = ['Нийт'] + [None] * 7 + [
        sum(r[k] for r in rows) for k in ('late_1_10', 'late_11_20', 'late_20_plus', 'total_late_minutes')
    ] + [None]
    return _xlsx_response(
        'Цаг бүртгэл',
        f'Цаг бүртгэл: {date_from:%Y-%m-%d} ~ {date_to:%Y-%m-%d}' + (f' · нэрээр шүүсэн: "{name_query}"' if name_query else '')
        + _age_range_label(age_min, age_max),
        headers, data,
        f'attendance_{date_from:%Y%m%d}_{date_to:%Y%m%d}.xlsx',
        widths=[10, 10, 22, 20, 11, 9, 9, 9, 9, 9, 10, 11, 40],
        totals=totals, number_cols=(9, 10, 11, 12), row_fills=fills,
    )


@login_required
def timesheet_export(request):
    """Хоног бүртгэлийн хүснэгтийг (хуудсан дээрх сараар) Excel файлаар татна."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from shop.services.timesheet import build_timesheet

    from shop.services.timesheet import timesheet_totals

    year, month = _timesheet_month(request)
    _, _, rows, totals, _ = build_timesheet(year, month)
    rows, age_min, age_max, _ = _filter_timesheet_by_age(request, rows)
    totals = timesheet_totals(rows)
    headers = [
        '№', 'Нэрс', 'Албан тушаал', 'Ажиллавал зохих хоног', 'Ажиллавал зохих цаг', 'Ажилласан хоног',
        'Ажилласан цаг', 'Бямбад ажилласан', 'Ээлжийн амралт /хоног/', 'Хоцролт /минут/ 1-10',
        'Хоцролт /минут/ 11-20', 'Хоцролт /минут/ 21-30', 'Чөлөөтэй хоног', 'Цагийн чөлөө',
        'Хоцролтын суутгал мөнгөн дүнгээр', 'Зөвшөөрсөн',
    ]
    value_keys = (
        'required_days', 'required_hours', 'worked_days', 'worked_hours', 'saturday_days', 'leave_days',
        'late_1_10', 'late_11_20', 'late_21_30', 'excused_days', 'excused_hours', 'deduction',
    )
    registry = _employee_registry(r['employee_code'] for r in rows)
    # Эхний 2 багана: ажилтны код, регистрийн дугаар (хэвлэх хувилбарт ороогүй)
    headers = ['Ажилтны код', 'Регистрийн дугаар'] + headers
    data = [
        [r['employee_code'], registry.get(r['employee_code']),
         i, r['name'], f"{r['positionname']} (шилжин)" if not r['is_main'] else (r['positionname'] or None)]
        + [_excel_number(r[k]) for k in value_keys] + [None]
        for i, r in enumerate(rows, start=1)
    ]
    # Хуудастай ижил: шилжин ажилласан мөр улбар шар, ажиллавал зохих хоногоос зөрүүтэй мөр шар
    fills = ['FFF7ED' if not r['is_main'] else ('FEF3C7' if r['day_gap'] else None) for r in rows]
    total_row = [None, None, 'Нийт', None, None] + [_excel_number(totals[k]) for k in value_keys] + [None]
    return _xlsx_response(
        'Хоног бүртгэл',
        f'Хоног бүртгэл: {year} оны {MONTH_NAMES_MN[month - 1]}' + _age_range_label(age_min, age_max),
        headers, data,
        f'timesheet_{year}_{month:02d}.xlsx',
        widths=[10, 13, 5, 22, 24, 11, 11, 11, 11, 10, 11, 10, 10, 10, 10, 10, 14, 12],
        totals=total_row, number_cols=tuple(range(6, 18)), row_fills=fills,
    )


def _employee_registry(codes):
    """{ажилтны код: регистрийн дугаар} - Excel экспортод."""
    return dict(OpenDataEmployee.objects.filter(id__in=set(codes)).values_list('id', 'registrynumber'))


@login_required
def attendance_calc(request):
    """Цаг бүртгэлийн тооцоо - төхөөрөмжийн ирц болон Ажилтны хурууны мэдээллийг (EmployeeFingerprint)
    нэгтгэж, сонгосон огнооны мужид харуулна. Хадгалсан засвар байвал тооцооллыг тэрээр орлуулна."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from shop.models import EmployeeFingerprint
    from shop.services.attendance import build_attendance_rows, get_position_choices, get_position_start_times

    today = timezone.localdate()
    # Нэрийн шүүлтийг (name_query) хуудсан дээр JS шүүнэ - бичих бүрт сервер рүү хандахгүй, бүх мөрийг дамжуулна
    date_from, date_to, name_query = _attendance_filter(request)

    from shop.models import AttendanceRecord

    rows = build_attendance_rows(date_from, date_to)
    # Ажиллавал зохих боловч мөргүй (ажиллаагүй) өдрүүд - "Ажиллаагүй өдрийг харуулах" чекээр хүснэгтэд
    # тодруулж харуулна (хоног дутуу үед аль өдөр ажиллаагүйг олоход). Хүснэгтэд нэр асц, огноо буурахаар холино
    from shop.services.timesheet import build_missing_days
    missing_days = build_missing_days(date_from, date_to, rows)
    table_rows = rows + missing_days
    table_rows.sort(key=lambda r: r.get('is_additional', False))
    table_rows.sort(key=lambda r: r['date'], reverse=True)
    table_rows.sort(key=lambda r: r['name'] or '')
    show_missing = (request.session.get('attendance_filter') or {}).get('missing') == '1'
    # Насны шүүлтийг ч нэрийн шүүлтийн адил хуудсан дээр JS шүүнэ - мөр бүрт ажилтны насыг (регистрээс) дамжуулна
    from shop.services.attendance import get_employee_ages
    age_min, age_max = _age_range(request, 'attendance_filter')
    ages = get_employee_ages()
    for r in table_rows:
        r['age'] = ages.get(r['employee_code'])
    error = None
    if not EmployeeFingerprint.objects.exclude(device_user_id='').exists():
        error = 'Ажилтны хурууны мэдээлэл бүртгэгдээгүй байна. "Ажилтны хурууны мэдээлэл" хуудаснаас ажилтан бүрийн төхөөрөмжийн кодыг тохируулна уу.'

    # Мөр нэмэх маягтын ажилтны сонголт
    employees = (
        OpenDataEmployee.objects.filter(isreclusion='N')
        .exclude(id__isnull=True).exclude(id='')
        .order_by('name')
        .values('id', 'name', 'positionname')
    )
    # Тухайн хугацаанд устгасан мөрүүд - эндээс сэргээнэ
    deleted_records = list(AttendanceRecord.objects.filter(
        date__gte=date_from, date__lte=date_to, is_deleted=True
    ).order_by('name', '-date'))
    for rec in deleted_records:
        rec.age = ages.get(rec.employee_code)

    context = {
        'date_from': date_from,
        'date_to': date_to,
        'name_query': name_query,
        'age_min': age_min,
        'age_max': age_max,
        'table_rows': table_rows,
        'missing_count': len(missing_days),
        'show_missing': show_missing,
        'rows': rows,
        'error': error,
        'employees': employees,
        'deleted_records': deleted_records,
        'today': today,
        # Мөрийн албан тушаалыг (шилжин ажилласан өдөр) сонгох жагсаалт, JS-д хоцролтыг дахин тооцох эхлэх цагууд
        'positions': get_position_choices(),
        'position_start_times': {
            name: start.strftime('%H:%M') if start else ''
            for name, start in get_position_start_times().items()
        },
    }
    return render(request, 'shop/attendance_calc.html', context)


@login_required
def attendance_sync(request):
    """Хурууны хээ таних төхөөрөмжөөс шинэ ирцийн бүртгэлийг татаж AttendanceRawLog-д хадгална"""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ үйлдэл зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    if request.method != 'POST':
        return redirect('shop:attendance_calc')

    from django.contrib import messages
    from shop.models import AttendanceRecord
    from shop.services.attendance import sync_from_device

    date_from = request.POST.get('date_from')
    date_to = request.POST.get('date_to')

    try:
        fetched, created = sync_from_device()
    except Exception as e:
        messages.error(request, f'Төхөөрөмжтэй холбогдоход алдаа гарлаа: {e}')
    else:
        # Хуудсан дээр харж буй хугацаанд ✕ товчоор устгасан мөрүүдийг сэргээнэ - устгасан тэмдэглэгээг
        # арилгахад мөр төхөөрөмжийн (түүхий) бүртгэлээс дахин тооцоологдож гарч ирнэ
        today = timezone.localdate()
        try:
            restore_from = datetime.strptime(date_from, '%Y-%m-%d').date() if date_from else today.replace(day=1)
            restore_to = datetime.strptime(date_to, '%Y-%m-%d').date() if date_to else today
        except ValueError:
            restore_from, restore_to = today.replace(day=1), today
        restored, _ = AttendanceRecord.objects.filter(
            is_deleted=True, date__gte=restore_from, date__lte=restore_to
        ).delete()

        message = f'Төхөөрөмжөөс {fetched} бүртгэл татагдлаа, шинээр {created} мөр хадгалагдлаа.'
        if restored:
            message += f' Устгасан {restored} мөр сэргээгдлээ.'
        messages.success(request, message)

    url = reverse('shop:attendance_calc')
    qs = []
    if date_from:
        qs.append(f'date_from={date_from}')
    if date_to:
        qs.append(f'date_to={date_to}')
    if qs:
        url += '?' + '&'.join(qs)
    return redirect(url)


@login_required
def attendance_sync_merchandiser(request):
    """Түгээлтийн программаас экспортолсон Excel файлыг (сонгож upload хийсэн) уншиж, худалдааны
    төлөөлөгч/мерчандайзерийн зочилсон (зураг илгээсэн) бүртгэлийг MerchandiserVisitLog-д импортолно"""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ үйлдэл зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    if request.method != 'POST':
        return redirect('shop:attendance_calc')

    from django.contrib import messages
    from shop.services.attendance import import_merchandiser_visits

    uploaded_file = request.FILES.get('document_file')
    if not uploaded_file:
        messages.error(request, 'Эхлээд Excel файлаа сонгоно уу.')
    else:
        try:
            fetched, created = import_merchandiser_visits(file_obj=uploaded_file)
            messages.success(request, f'{uploaded_file.name}-с {fetched} бүртгэл уншигдлаа, шинээр {created} мөр хадгалагдлаа.')
        except ValueError as e:
            messages.error(request, f'Файлын бүтэц таарахгүй байна: {e}')
        except Exception as e:
            messages.error(request, f'Файл уншихад алдаа гарлаа: {e}')

    url = reverse('shop:attendance_calc')
    date_from = request.POST.get('date_from')
    date_to = request.POST.get('date_to')
    qs = []
    if date_from:
        qs.append(f'date_from={date_from}')
    if date_to:
        qs.append(f'date_to={date_to}')
    if qs:
        url += '?' + '&'.join(qs)
    return redirect(url)


def _parse_decimal(value, default=Decimal(0)):
    """Маягтын тоон утгыг Decimal болгоно (хоосон/буруу бол default, сөрөг утгыг 0 болгоно)."""
    text = (value or '').strip().replace(',', '.')
    if not text:
        return default
    try:
        number = Decimal(text)
    except InvalidOperation:
        return default
    if not number.is_finite():
        return default
    return max(number, Decimal(0))


def _parse_money(value, default=Decimal(0)):
    """Мянгачилж харуулсан мөнгөн дүнг (1,234,567.89) Decimal болгоно - таслал, зайг мянгатын тусгаарлагч гэж хасна."""
    return _parse_decimal((value or '').replace(',', '').replace(' ', '').replace('\u00a0', ''), default=default)


@login_required
def timesheet(request):
    """Хоног бүртгэл - идэвхтэй ажилчдын сарын ажиллавал зохих/ажилласан хоног, хоцролт, суутгалыг
    Цаг бүртгэлийн мэдээллээс нэгтгэж харуулна (shop.services.timesheet.build_timesheet).

    Хуудасны эхэнд тухайн сарын өдрүүдээс байгууллагын амралтын өдрийг чеклэж, ажилтан бүрийн
    ээлжийн амралт/чөлөө болон хоцролтын минут тутмын
    суутгалын үнэлгээг гараар оруулж хадгална."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    import calendar
    from django.contrib import messages
    from shop.models import OrganizationHoliday, TimesheetEntry, TimesheetSetting
    from shop.services.timesheet import build_timesheet

    today = timezone.localdate()
    year, month = _timesheet_month(request)

    date_from = datetime(year, month, 1).date()
    date_to = date_from.replace(day=calendar.monthrange(year, month)[1])
    redirect_url = f"{reverse('shop:timesheet')}?year={year}&month={month}"

    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'holidays':
            selected = set()
            for value in request.POST.getlist('holiday'):
                try:
                    d = datetime.strptime(value, '%Y-%m-%d').date()
                except ValueError:
                    continue
                if date_from <= d <= date_to:
                    selected.add(d)
            OrganizationHoliday.objects.filter(date__gte=date_from, date__lte=date_to).exclude(date__in=selected).delete()
            for d in selected:
                OrganizationHoliday.objects.update_or_create(date=d, defaults={'updated_by': request.user})
            messages.success(request, f'Амралтын өдөр хадгалагдлаа ({len(selected)} өдөр).')

        elif action == 'entries':
            codes = request.POST.getlist('employee_code')
            # Хоосон бол үндсэн мөр, үгүй бол шилжин ажилласан албан тушаалын тусдаа мөр
            positions = request.POST.getlist('entry_position')
            leave = request.POST.getlist('leave_days')
            excused_days = request.POST.getlist('excused_days')
            excused_hours = request.POST.getlist('excused_hours')
            # Гараар засах хоног ('худалдагч' албан тушаал): тооцоолсон (auto) утгаас өөр бол л хадгална -
            # ингэснээр засаагүй мөр цаг бүртгэлийн өөрчлөлтийг дагаж шинэчлэгдсээр байна
            required_days = request.POST.getlist('required_days')
            required_auto = request.POST.getlist('required_days_auto')
            worked_days = request.POST.getlist('worked_days')
            worked_auto = request.POST.getlist('worked_days_auto')
            saturday_days = request.POST.getlist('saturday_days')
            saturday_auto = request.POST.getlist('saturday_days_auto')

            def override(value, auto):
                number = _parse_decimal(value, default=None)
                if number is None or number == _parse_decimal(auto, default=None):
                    return None
                return number

            saved = 0
            for code, position, lv, exd, exh, req, req_auto, wrk, wrk_auto, sat, sat_auto in zip(
                codes, positions, leave, excused_days, excused_hours, required_days, required_auto, worked_days,
                worked_auto, saturday_days, saturday_auto,
            ):
                values = {
                    'leave_days': _parse_decimal(lv),
                    'excused_days': _parse_decimal(exd),
                    'excused_hours': _parse_decimal(exh),
                    'required_days': override(req, req_auto),
                    'worked_days': override(wrk, wrk_auto),
                    'saturday_days': override(sat, sat_auto),
                }
                if not any(v for v in values.values() if v is not None):
                    TimesheetEntry.objects.filter(
                        employee_code=code, year=year, month=month, position_name=position,
                    ).delete()
                    continue
                TimesheetEntry.objects.update_or_create(
                    employee_code=code, year=year, month=month, position_name=position,
                    defaults={**values, 'updated_by': request.user},
                )
                saved += 1
            messages.success(request, f'Хоног бүртгэл хадгалагдлаа ({saved} ажилтан).')

        elif action == 'reset_entries':
            # Тухайн сарын бүх гар утгыг устгана - хоног бүртгэл цаг бүртгэлээс дахин тооцоологдож,
            # дутуу хоногийг чөлөөтэй хоногт дахин санал болгоно
            deleted, _ = TimesheetEntry.objects.filter(year=year, month=month).delete()
            messages.success(request, f'Хадгалсан хоног бүртгэл цэвэрлэгдлээ ({deleted} ажилтан). Утгууд цаг бүртгэлээс дахин тооцоологдлоо.')

        elif action == 'rates':
            setting = TimesheetSetting.get()
            setting.late_rate_1_10 = _parse_decimal(request.POST.get('late_rate_1_10'))
            setting.late_rate_11_20 = _parse_decimal(request.POST.get('late_rate_11_20'))
            setting.late_rate_21_30 = _parse_decimal(request.POST.get('late_rate_21_30'))
            setting.save()
            messages.success(request, 'Хоцролтын суутгалын үнэлгээ хадгалагдлаа.')

        return redirect(redirect_url)

    days, calendar_summary, rows, totals, setting = build_timesheet(year, month)
    # Насны шүүлт (регистрээс тооцсон өнөөдрийн нас) - шүүсэн мөрүүдээр нийт дүнг дахин тооцно
    from shop.services.timesheet import timesheet_totals
    rows, age_min, age_max, ages = _filter_timesheet_by_age(request, rows)
    totals = timesheet_totals(rows)
    for r in rows:
        r['age'] = ages.get(r['employee_code'])

    # Хуанлийг Даваагаас эхлэх 7 баганатай долоо хоногуудад хуваана (сарын эхний өдрийн өмнө хоосон нүд)
    weeks = []
    week = [None] * days[0]['date'].weekday()
    for d in days:
        week.append(d)
        if len(week) == 7:
            weeks.append(week)
            week = []
    if week:
        weeks.append(week + [None] * (7 - len(week)))

    prev_month = (date_from - timedelta(days=1)).replace(day=1)
    next_month = date_to + timedelta(days=1)

    context = {
        'year': year,
        'month': month,
        'years': list(range(today.year, 2019, -1)),
        'months': list(enumerate(MONTH_NAMES_MN, start=1)),
        'month_label': f'{year} оны {MONTH_NAMES_MN[month - 1]}',
        'prev_month': prev_month,
        'next_month': next_month,
        'weeks': weeks,
        'weekday_headers': ['Да', 'Мя', 'Лх', 'Пү', 'Ба', 'Бя', 'Ня'],
        'calendar_summary': calendar_summary,
        'age_min': age_min,
        'age_max': age_max,
        'rows': rows,
        # Шилжин ажилласан албан тушаалын мөрүүдээс болж мөрийн тоо ажилтны тооноос их байж болно
        'employee_count': len({r['employee_code'] for r in rows}),
        'totals': totals,
        'setting': setting,
    }
    return render(request, 'shop/timesheet.html', context)


def _month_from_request(source):
    """?year=&month= (эсвэл POST)-оос сарыг уншина, буруу бол тухайн сар.

    USE_THOUSAND_SEPARATOR асаалттай тул template-д {{ year }} гэж бичвэл '2,026' болдог - тэр тохиолдолд ч
    зөв уншихын тулд мянгатын тусгаарлагч, зайг хасна."""
    today = timezone.localdate()

    def as_int(value, default):
        text = (value or '').replace(',', '').replace(' ', '').replace(' ', '')
        return int(text) if text else default

    try:
        year = as_int(source.get('year'), today.year)
        month = as_int(source.get('month'), today.month)
        if not (1 <= month <= 12 and 2000 <= year <= today.year + 1):
            raise ValueError
    except ValueError:
        year, month = today.year, today.month
    return year, month


def _month_nav_context(year, month):
    """Сар сонгох маягтад хэрэгтэй утгууд (Хоног бүртгэл, Цалин бодолт)."""
    import calendar
    today = timezone.localdate()
    date_from = datetime(year, month, 1).date()
    date_to = date_from.replace(day=calendar.monthrange(year, month)[1])
    return {
        'year': year,
        'month': month,
        'years': list(range(today.year, 2019, -1)),
        'months': list(enumerate(MONTH_NAMES_MN, start=1)),
        'month_label': f'{year} оны {MONTH_NAMES_MN[month - 1]}',
        'prev_month': (date_from - timedelta(days=1)).replace(day=1),
        'next_month': date_to + timedelta(days=1),
    }


@login_required
def payroll(request, sheet):
    """Цалин бодолт - карт (sheet='card', НДШ/ХХОАТ суутгана) эсвэл бэлэн (sheet='cash', татваргүй) хуудас.

    Хоног бүртгэлийн хоног/цагаар цалинг хувьтгаж, сар бүрийн гар утгуудыг (нэмэгдэл, суутгал, ХЧТА,
    хуримтлал) хүснэгтэд бичиж хадгална - shop.services.payroll.build_payroll тайлбарыг үз."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from django.contrib import messages
    from shop.models import PayrollEntry
    from shop.services.payroll import (
        BENEFIT_FIELDS, DEDUCTION_COLUMNS, EARNING_FIELDS, MANUAL_FIELDS, OVERRIDE_FIELDS, RECEIVABLE_ACCOUNTS,
    )

    from shop.services import payroll_snapshot

    is_card = sheet == PayrollEntry.SHEET_CARD
    url_name = 'shop:payroll_card' if is_card else 'shop:payroll_cash'
    year, month = _month_from_request(request.POST if request.method == 'POST' else request.GET)
    close = payroll_snapshot.get_close(year, month, sheet)

    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'close_month':
            count = payroll_snapshot.close_month(year, month, sheet, request.user)
            messages.success(request, f'{year} оны {month}-р сарын {"картын" if is_card else "бэлэн"} цалин хаагдлаа ({count} ажилтан). '
                             'Цалингийн тайлан, хувийн мэдээлэлд энэ бодолтоор харагдана.')
            return redirect(f"{reverse(url_name)}?year={year}&month={month}")
        if action == 'reopen_month':
            payroll_snapshot.reopen_month(year, month, sheet)
            messages.warning(request, 'Сарыг дахин нээлээ - цалин одоогийн мэдээллээр дахин бодогдоно. Засварласны дараа дахин хаана уу.')
            return redirect(f"{reverse(url_name)}?year={year}&month={month}")
        if close:
            messages.error(request, 'Энэ сар хаагдсан тул засварлах боломжгүй. Эхлээд сарыг дахин нээнэ үү.')
            return redirect(f"{reverse(url_name)}?year={year}&month={month}")

    if request.method == 'POST' and request.POST.get('action') == 'pull_receivables':
        from shop.services.payroll import pull_receivables
        count, total, as_of = pull_receivables(year, month, sheet, request.user)
        messages.success(request, f'Ажилчдын авлагыг {as_of:%Y-%m-%d}-ний үлдэгдлээр татлаа: {count} ажилтан, нийт {total:,.0f}₮.')
        return redirect(f"{reverse(url_name)}?year={year}&month={month}")

    if request.method == 'POST':
        codes = request.POST.getlist('employee_code')
        # Нүдэнд мянгачилж харуулдаг тул мянгатын таслалыг хасна (_parse_decimal таслалыг аравтын тэмдэг гэж үздэг)
        values_by_field = {f: [v.replace(',', '').replace(' ', '') for v in request.POST.getlist(f)] for f in MANUAL_FIELDS}
        # Бэлэн хуудсанд Бямбын нэмэгдэл, унааны мөнгийг гараар дарж бичнэ: хоосон бол автомат (NULL), 0 бол 0
        override_fields = [f for _, f in OVERRIDE_FIELDS] if not is_card else []
        overrides_by_field = {f: [v.replace(',', '').replace(' ', '') for v in request.POST.getlist(f)] for f in override_fields}
        saved = 0
        for i, code in enumerate(codes):
            values = {
                f: _parse_decimal(values_by_field[f][i] if i < len(values_by_field[f]) else '')
                for f in MANUAL_FIELDS
            }
            for f in override_fields:
                raw = overrides_by_field[f][i] if i < len(overrides_by_field[f]) else ''
                values[f] = _parse_decimal(raw, default=None)
            if not any(values[f] for f in MANUAL_FIELDS) and all(values[f] is None for f in override_fields):
                # Картын хуудсанд гараар дарж бичсэн утга байхгүй тул мөрийг устгана
                PayrollEntry.objects.filter(employee_code=code, year=year, month=month, sheet=sheet).delete()
                continue
            PayrollEntry.objects.update_or_create(
                employee_code=code, year=year, month=month, sheet=sheet,
                defaults={**values, 'updated_by': request.user},
            )
            saved += 1
        messages.success(request, f'Цалин бодолт хадгалагдлаа ({saved} ажилтан).')
        return redirect(f"{reverse(url_name)}?year={year}&month={month}")

    rows, totals, setting, close = payroll_snapshot.get_payroll(year, month, sheet)
    from shop.services.payroll_print import build_print_pages
    context = {
        'close': close,
        **_month_nav_context(year, month),
        'sheet': sheet,
        'is_card': is_card,
        'active_page': 'payroll_card' if is_card else 'payroll_cash',
        'title': 'Картын цалин бодолт' if is_card else 'Бэлэн цалин бодолт',
        'url_name': url_name,
        'rows': rows,
        'totals': totals,
        'setting': setting,
        'earning_fields': EARNING_FIELDS,
        'deduction_columns': DEDUCTION_COLUMNS,
        'benefit_fields': BENEFIT_FIELDS,
        'receivable_accounts': RECEIVABLE_ACCOUNTS,
        # Хэвлэхэд бүх мөр нь хоосон (нийт 0) нэмэгдэл/суутгалын баганыг гаргахгүй - үлдсэн нь томоор багтана
        # Хэвлэх хуудсууд (баганыг Цалингийн тохиргооноос). Олгох цалингүй (нийт цалин, олгох дүн хоёулаа 0)
        # ажилтны мөрийг хэвлэхгүй - жиш: хосолсон ажилтны бэлэн хэсэг хоосон бол
        'print_pages': build_print_pages(
            setting, sheet, [r for r in rows if r['gross'] or r['payout']], totals,
        ),
    }
    return render(request, 'shop/payroll.html', context)


@login_required
def payroll_export(request, sheet):
    """Цалин бодолтыг (карт/бэлэн) хуудсан дээрх бүх баганаар Excel файлаар татна. Эхний 2 багана: ажилтны код,
    регистрийн дугаар. Хэвлэх хувилбарт нөлөөлөхгүй."""
    if not request.user.is_staff:
        return redirect('shop:home')

    from shop.models import PayrollEntry
    from shop.services.payroll import BENEFIT_FIELDS, DEDUCTION_COLUMNS, EARNING_FIELDS
    from shop.services.payroll_snapshot import get_payroll

    is_card = sheet == PayrollEntry.SHEET_CARD
    year, month = _month_from_request(request.GET)
    rows, totals, _, _ = get_payroll(year, month, sheet)
    registry = _employee_registry(r['employee_code'] for r in rows)

    # (толгой, түлхүүр) - хуудасны хүснэгттэй ижил дараалал
    columns = [
        ('Үндсэн цалин', 'base_salary'), ('Удаан жилийн нэмэгдэл', 'seniority_bonus'), ('Нийт цалин', 'nominal_salary'),
        ('Ажиллавал зохих хоног', 'required_days'), ('Ажиллавал зохих цаг', 'required_hours'),
        ('Ажилласан хоног', 'worked_days'), ('Ажилласан цаг', 'worked_hours'),
        ('Бямбад ажилласан цаг', 'saturday_hours'), ('Бямбад ажилласан нэмэгдэл', 'saturday_bonus'),
    ] + [(label, field) for field, label in EARNING_FIELDS] + [
        ('Унааны мөнгө', 'transport'), ('Бодогдсон нийт цалин', 'gross'),
    ]
    if is_card:
        columns += [
            ('НДШ код', 'ndsh_code'), ('НДШ ажилтан %', 'ndsh_employee_rate'), ('НДШ ажилтан', 'ndsh_employee'),
            ('НДШ байгууллага %', 'ndsh_employer_rate'), ('НДШ байгууллага', 'ndsh_employer'),
            ('ХХОАТ тооцох дүн', 'taxable'), ('Хөнгөлөлт шатлалаар', 'credit'), ('ХХОАТ', 'pit'),
            ('Нийт татварын суутгал', 'total_tax'),
        ]
    columns += [(label, field) for field, label, _ in DEDUCTION_COLUMNS]
    columns += [(label, field) for field, label in BENEFIT_FIELDS]
    columns += [('Гарт олгох', 'net_pay'), ('Хадгаламж', 'deposit'),
                ('Картанд орох' if is_card else 'Бэлнээр олгох', 'payout')]
    if not is_card:
        columns.append(('Бэлэн цалингийн данс', 'cash_bank_account'))
    text_keys = {'ndsh_code', 'cash_bank_account'}
    no_total = text_keys | {'ndsh_employee_rate', 'ndsh_employer_rate'}

    lead = ['Ажилтны код', 'Регистрийн дугаар', '№', 'Ажилчдын нэрс', 'Албан тушаал']
    headers = lead + [label for label, _ in columns]
    data = [
        [r['employee_code'], registry.get(r['employee_code']), i, r['name'], r['positionname']]
        + [(r[key] or None) if key in text_keys else _excel_number(r[key]) for _, key in columns]
        for i, r in enumerate(rows, start=1)
    ]
    total_row = [None, None, 'Нийт', None, None] + [
        None if key in no_total else _excel_number(totals.get(key, 0)) for _, key in columns
    ]
    number_cols = tuple(len(lead) + i for i, (_, key) in enumerate(columns, start=1) if key not in text_keys)
    title = 'Картын цалин бодолт' if is_card else 'Бэлэн цалин бодолт'
    return _xlsx_response(
        'Цалин', f'{title}: {year} оны {MONTH_NAMES_MN[month - 1]}', headers, data,
        f'payroll_{sheet}_{year}_{month:02d}.xlsx',
        widths=[10, 13, 5, 22, 22] + [12] * len(columns),
        totals=total_row, number_cols=number_cols,
    )


PAYROLL_REPORT_FILTER_DIMS = ['department', 'position', 'gender', 'pay_type', 'age_band', 'tenure_band']
PAYROLL_REPORT_MAX_MONTHS = 36


def _parse_year_month(value, default):
    try:
        y, m = (int(x) for x in (value or '').split('-')[:2])
        if 2000 <= y <= 2100 and 1 <= m <= 12:
            return y, m
    except ValueError:
        pass
    return default


@login_required
def cash_bank_accounts(request):
    """Бэлэн цалингийн данс: бэлэн / хосолсон цалинтай ажилчдын дансыг (Цалингийн тохиргооноос) цэгцэлж харуулна -
    "MN"-ийн араас тоонуудыг зайгүй (shop.services.bank_accounts). Сонгосон сарын бэлэн цалингийн олгох дүн хамт.
    Шүүлтүүр: хайлт, банк, албан тушаал, хэлтэс, олгох хэлбэр, дансны төлөв, идэвхтэй эсэх. Excel."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from urllib.parse import urlencode
    from shop.models import EmployeePayProfile, PayrollEntry
    from shop.services.bank_accounts import account_info
    from shop.services.payroll_snapshot import get_payroll

    year, month = _month_from_request(request.GET)
    profiles = list(EmployeePayProfile.objects.filter(
        pay_type__in=[EmployeePayProfile.PAY_CASH, EmployeePayProfile.PAY_MIXED]))
    employees = {e.id: e for e in OpenDataEmployee.objects.filter(id__in=[p.employee_code for p in profiles])}
    cash_rows, _, _, close = get_payroll(year, month, PayrollEntry.SHEET_CASH)
    payouts = {r['employee_code']: r['payout'] for r in cash_rows}
    pay_labels = dict(EmployeePayProfile.PAY_CHOICES)

    all_rows = []
    for p in profiles:
        e = employees.get(p.employee_code)
        info = account_info(p.cash_bank_account)
        all_rows.append({
            'code': p.employee_code,
            'name': e.name if e else p.employee_code,
            'position': (e.positionname if e else '') or '',
            'department': (e.departmentname if e else '') or '',
            'active': bool(e and e.isreclusion == 'N'),
            'pay_type': p.pay_type,
            'pay_label': pay_labels.get(p.pay_type, ''),
            'payout': payouts.get(p.employee_code),
            **info,
        })

    f = {
        'q': (request.GET.get('q') or '').strip(),
        'bank': request.GET.get('bank') or '',
        'position': request.GET.get('position') or '',
        'department': request.GET.get('department') or '',
        'pay_type': request.GET.get('pay_type') or '',
        'status': request.GET.get('status') or '',
        # Анхдагчаар зөвхөн идэвхтэй ажилчид
        'active': request.GET.get('active') or '1',
    }

    def keep(r):
        q = f['q'].lower()
        return ((not q or q in r['name'].lower() or q in r['code'] or q.replace(' ', '') in r['account'].lower())
                and (not f['bank'] or r['bank'] == f['bank'])
                and (not f['position'] or r['position'] == f['position'])
                and (not f['department'] or r['department'] == f['department'])
                and (not f['pay_type'] or r['pay_type'] == f['pay_type'])
                and (not f['status'] or r['status'] == f['status'])
                and (f['active'] != '1' or r['active']))

    from shop.services.timesheet import position_sort_key

    # Эрэмбэ: ?sort=<түлхүүр> (өсөхөөр) эсвэл -<түлхүүр> (буурахаар), анхдагч албан тушаалаар - цалин бодолттой ижил
    # дарааллаар (Ерөнхий захирал эхэнд г.м.). Хоосон утга үргэлж төгсгөлд, тэнцүү бол нэрээр
    position_key = position_sort_key()
    sort_keys = {
        'name': lambda r: r['name'].lower(),
        'position': lambda r: position_key(r['position']),
        'pay_type': lambda r: r['pay_label'],
        'bank': lambda r: r['bank'],
        'account': lambda r: r['account'],
        'payout': lambda r: r['payout'] if r['payout'] is not None else Decimal(0),
    }
    sort = request.GET.get('sort') or 'position'
    sort_field = sort.lstrip('-') if sort.lstrip('-') in sort_keys else 'position'
    descending = sort.startswith('-') and sort_field == sort[1:]
    sort = ('-' if descending else '') + sort_field
    rows = [r for r in all_rows if keep(r)]
    rows.sort(key=lambda r: r['name'] or '')
    rows.sort(key=sort_keys[sort_field], reverse=descending)
    empty_last = {'position': 'position', 'bank': 'bank', 'account': 'account', 'payout': 'payout'}.get(sort_field)
    if empty_last:
        rows.sort(key=lambda r: r[empty_last] in ('', None))
    for i, r in enumerate(rows, start=1):
        r['no'] = i

    if request.GET.get('export') == 'xlsx':
        headers = ['№', 'Ажилтны код', 'Ажилтан', 'Албан тушаал', 'Хэлтэс', 'Олгох хэлбэр', 'Банк', 'Данс',
                   f'Бэлэн олгох дүн ({year}.{month:02d})']
        data = [[r['no'], r['code'], r['name'], r['position'], r['department'], r['pay_label'], r['bank'], r['account'],
                 _excel_number(r['payout']) if r['payout'] is not None else None] for r in rows]
        total = sum((r['payout'] or 0 for r in rows), Decimal(0))
        return _xlsx_response(
            'Бэлэн цалингийн данс', f'Бэлэн цалингийн данс - {year} оны {MONTH_NAMES_MN[month - 1]}', headers, data,
            f'cash_bank_accounts_{year}_{month:02d}.xlsx', widths=[5, 10, 24, 22, 18, 12, 22, 26, 16],
            totals=['Нийт', None, f'{len(rows)} ажилтан', None, None, None, None, None, _excel_number(total)],
            number_cols=(9,),
        )

    def options(key):
        return sorted({r[key] for r in all_rows if r[key]})

    params = {k: v for k, v in request.GET.items() if k != 'export'}
    context = {
        **_month_nav_context(year, month),
        'active_page': 'cash_bank_accounts',
        'rows': rows,
        'f': f,
        'sort': sort,
        # Баганын толгой: (түлхүүр, нэр, баруун тийш, эрэмбэлэх холбоос, идэвхтэй чиглэл)
        'columns': [
            (key, label, right, '?' + urlencode({**params, 'sort': key if sort != key else f'-{key}'}),
             ('asc' if sort == key else 'desc' if sort == f'-{key}' else ''))
            for key, label, right in [
                ('name', 'Ажилтан', False), ('position', 'Албан тушаал', False), ('pay_type', 'Олгох хэлбэр', False),
                ('bank', 'Банк', False), ('account', 'Данс', False), ('payout', 'Бэлэн олгох дүн', True),
            ]
        ],
        'close': close,
        'bank_options': options('bank'),
        'position_options': options('position'),
        'department_options': options('department'),
        'pay_type_options': [(EmployeePayProfile.PAY_CASH, 'Бэлэн'), (EmployeePayProfile.PAY_MIXED, 'Хосолсон')],
        'status_options': [('ok', 'Зөв'), ('invalid', 'Шалгах (урт буруу)'), ('missing', 'Данс оруулаагүй')],
        'has_filters': any(f[k] for k in ('q', 'bank', 'position', 'department', 'pay_type', 'status')),
        'summary': {
            'count': len(rows),
            'with_account': sum(1 for r in rows if r['status'] != 'missing'),
            'invalid': sum(1 for r in rows if r['status'] == 'invalid'),
            'missing': sum(1 for r in rows if r['status'] == 'missing'),
            'payout': sum((r['payout'] or 0 for r in rows), Decimal(0)),
            'by_bank': [
                (bank, sum(1 for r in rows if r['bank'] == bank), sum((r['payout'] or 0 for r in rows if r['bank'] == bank), Decimal(0)))
                for bank in sorted({r['bank'] for r in rows if r['bank']})
            ],
        },
        'export_url': '?' + urlencode({**params, 'export': 'xlsx', 'year': year, 'month': month}),
    }
    return render(request, 'shop/cash_bank_accounts.html', context)


@login_required
def payroll_report(request):
    """Цалингийн тайлан: сонгосон хугацааны (сараар) карт, бэлэн цалинг нийлүүлж эсвэл тусад нь - ажилтан, албан
    тушаал, хэлтэс, хүйс, насны бүлэг, ажилласан жил, олгох хэлбэр, сараар бүлэглэнэ. Нийт дүн эсвэл сараар задалсан
    (pivot) хүснэгт, сарын график, Excel. Хаагдсан сард хаасан үеийн бодолт, хаагдаагүй сард шууд бодолт
    (shop.services.payroll_snapshot)."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from urllib.parse import urlencode
    from shop.services import payroll_snapshot as ps

    today = timezone.localdate()
    current = (today.year, today.month)
    date_from = _parse_year_month(request.GET.get('from'), (today.year, 1))
    date_to = min(_parse_year_month(request.GET.get('to'), current), current)
    if date_from > date_to:
        date_from = date_to
    months = ps.month_range_list(*date_from, *date_to)
    truncated = len(months) > PAYROLL_REPORT_MAX_MONTHS
    if truncated:
        months = months[-PAYROLL_REPORT_MAX_MONTHS:]
        date_from = months[0]

    sheet = request.GET.get('sheet') if request.GET.get('sheet') in ('card', 'cash') else 'all'
    sheets = [sheet] if sheet != 'all' else ['card', 'cash']
    group = request.GET.get('group') if request.GET.get('group') in ps.DIMENSION_LABELS else 'employee'
    layout = 'pivot' if request.GET.get('layout') == 'pivot' and group != 'month' else 'total'
    metric = request.GET.get('metric') if request.GET.get('metric') in ps.METRIC_LABELS else 'payout'
    q = (request.GET.get('q') or '').strip()
    # Сонгоогүй шүүлтүүр маягтаас хоосон утгаар (f_department=) ирдэг тул хоосныг шүүлтүүргүй гэж үзнэ
    filters = {d: request.GET[f'f_{d}'] for d in PAYROLL_REPORT_FILTER_DIMS if request.GET.get(f'f_{d}')}

    all_facts, status = ps.payroll_facts(months, sheets)

    # Шүүлтүүрийн сонголтууд - хугацааны бүх баримтаас (шүүхээс өмнө)
    filter_options = []
    for d in PAYROLL_REPORT_FILTER_DIMS:
        values = {}
        for f in all_facts:
            key, label = ps.dim_value(f, d)
            values[key] = label
        if d in ('age_band', 'tenure_band'):
            order = ps.band_order(d)
            options = sorted(values.items(), key=lambda kv: order.get(kv[1], 99))
        else:
            options = sorted(values.items(), key=lambda kv: (kv[0] == ps.UNKNOWN, str(kv[1])))
        filter_options.append({'dim': d, 'label': ps.DIMENSION_LABELS[d], 'options': options, 'value': filters.get(d)})

    facts = [
        f for f in all_facts
        if all(ps.dim_value(f, d)[0] == v for d, v in filters.items())
        and (not q or q.lower() in (f['name'] or '').lower() or q in f['employee_code'])
    ]
    groups = ps.aggregate(facts, group, months, metric if layout == 'pivot' else 'payout')
    total = (ps.aggregate(facts, 'all', months) or [None])[0]
    total_payout = total['payout'] if total else Decimal(0)
    # Сонголтоор бүлэг бүрийг карт, бэлэн дэд мөрөөр задлана (нийлүүлсэн төрөлд, төрлөөр бүлэглээгүй үед)
    split_available = sheet == 'all' and group != 'sheet'
    split = split_available and request.GET.get('split') == '1'
    if split:
        facts_by_group = {}
        for f in facts:
            facts_by_group.setdefault(ps.dim_value(f, group)[0], []).append(f)
        for g in groups:
            g['parts'] = ps.aggregate(facts_by_group.get(g['key'], []), 'sheet', months)
        if total:
            total['parts'] = ps.aggregate(facts, 'sheet', months)

    def decorate(g):
        g['share'] = g['payout'] / total_payout * 100 if total_payout else Decimal(0)
        g['metric_value'] = g[metric]
        g['pivot'] = [v[metric] if v else None for v in g['month_values']]
        g['metric_list'] = [g[k] for k in ps.METRIC_KEYS]
        for part in g.get('parts', []):
            decorate(part)

    for g in groups + ([total] if total else []):
        decorate(g)

    # Сараар задлахад карт, бэлэн нь дэд мөрөөр биш багана болно: сар бүр (карт, бэлэн), төгсгөлд (карт, бэлэн, нийт)
    split_columns = split and layout == 'pivot'
    if split_columns:
        for g in groups + ([total] if total else []):
            parts = {p['key']: p for p in g.pop('parts', [])}
            card, cash = parts.get('card'), parts.get('cash')
            cells = []
            for i in range(len(months)):
                cells += [(card['pivot'][i] if card else None, False), (cash['pivot'][i] if cash else None, False)]
            cells += [(card[metric] if card else None, False), (cash[metric] if cash else None, False), (g[metric], True)]
            g['split_cells'] = cells

    base_params = {k: v for k, v in request.GET.items() if k not in ('export',)}

    def url_with(**changes):
        params = {**base_params, **changes}
        return '?' + urlencode({k: v for k, v in params.items() if v not in (None, '')})

    # Бүлгийн мөрөөс тэр бүлгийн ажилчид руу задлах холбоос
    for g in groups:
        if group in PAYROLL_REPORT_FILTER_DIMS:
            g['drill_url'] = url_with(**{f'f_{group}': g['key'], 'group': 'employee', 'layout': 'total'})
        elif group == 'sheet':
            g['drill_url'] = url_with(sheet=g['key'], group='employee', layout='total')
        elif group == 'month':
            g['drill_url'] = url_with(**{'from': g['key'], 'to': g['key'], 'group': 'employee', 'layout': 'total'})

    if request.GET.get('export') == 'xlsx':
        return _payroll_report_xlsx(groups, total, group, layout, metric, months, date_from, date_to, sheet, split_columns)

    # Сарын график: сонгосон үзүүлэлт карт / бэлэнээр
    chart = []
    for y, m in months:
        values = {s: Decimal(0) for s in sheets}
        for f in facts:
            if f['year'] == y and f['month'] == m:
                values[f['sheet']] += f[metric]
        chart.append({
            'label': f'{y}.{m:02d}', 'card': float(values.get('card', 0)), 'cash': float(values.get('cash', 0)),
            'closed': all(status.get((y, m, s)) for s in sheets),
        })
    closed_months = sum(1 for c in chart if c['closed'])

    sheet_label = {'all': 'Карт + бэлэн', 'card': 'Картын цалин', 'cash': 'Бэлэн цалин'}[sheet]
    active_filters = [
        {'label': fo['label'], 'value': dict(fo['options']).get(filters[fo['dim']], filters[fo['dim']]) or ps.UNKNOWN,
         'remove_url': url_with(**{f"f_{fo['dim']}": None})}
        for fo in filter_options if fo['dim'] in filters
    ]
    if q:
        active_filters.append({'label': 'Хайлт', 'value': q, 'remove_url': url_with(q=None)})
    last_year = today.year - 1
    twelve_ago = ps.month_range_list(today.year - 1, today.month, *current)[1]

    context = {
        'active_page': 'payroll_report',
        'date_from': f'{date_from[0]}-{date_from[1]:02d}',
        'date_to': f'{date_to[0]}-{date_to[1]:02d}',
        'max_month': f'{today.year}-{today.month:02d}',
        'period_label': (f'{date_from[0]}.{date_from[1]:02d} - {date_to[0]}.{date_to[1]:02d}'
                         if date_from != date_to else f'{date_from[0]} оны {MONTH_NAMES_MN[date_from[1] - 1]}'),
        'months_count': len(months),
        'truncated': truncated,
        'max_months': PAYROLL_REPORT_MAX_MONTHS,
        'presets': [
            ('Энэ сар', url_with(**{'from': f'{today.year}-{today.month:02d}', 'to': f'{today.year}-{today.month:02d}'})),
            ('Энэ он', url_with(**{'from': f'{today.year}-01', 'to': f'{today.year}-{today.month:02d}'})),
            ('Сүүлийн 12 сар', url_with(**{'from': f'{twelve_ago[0]}-{twelve_ago[1]:02d}', 'to': f'{today.year}-{today.month:02d}'})),
            (f'{last_year} он', url_with(**{'from': f'{last_year}-01', 'to': f'{last_year}-12'})),
        ],
        'sheet': sheet,
        'sheet_label': sheet_label,
        'group': group,
        'group_label': ps.DIMENSION_LABELS[group],
        'group_by_label': ps.DIMENSION_BY_LABELS[group],
        'split': split,
        'split_columns': split_columns,
        # Хоёр давхар толгойн доод мөр: (нэр, хүснэгтийн баганын индекс, нийт эсэх)
        'split_headers': [
            (label, 2 + i, i >= 2 * len(months) + 2)
            for i, label in enumerate(['Карт', 'Бэлэн'] * (len(months) + 1) + ['Нийт'])
        ] if split_columns else [],
        'split_available': split_available,
        'layout': layout,
        'metric': metric,
        'metric_label': ps.METRIC_LABELS[metric],
        'metrics': ps.METRICS,
        'dimensions': ps.DIMENSIONS,
        'filter_options': filter_options,
        'active_filters': active_filters,
        'clear_filters_url': url_with(q=None, **{f'f_{d}': None for d in PAYROLL_REPORT_FILTER_DIMS}),
        'q': q,
        'groups': groups,
        'total': total,
        'payout_index': ps.METRIC_KEYS.index('payout') + 1,
        'kpi': {
            'tax': total['ndsh_employee'] + total['pit'],
            'payout_with_advance': total['payout'] + total['advance'],
        } if total else {},
        'month_labels': [f'{y}.{m:02d}' for y, m in months],
        'chart_data': {'months': chart, 'sheets': sheets, 'metric': ps.METRIC_LABELS[metric]},
        'closed_months': closed_months,
        'open_months': len(months) - closed_months,
        'export_url': url_with(export='xlsx'),
    }
    return render(request, 'shop/payroll_report.html', context)


def _payroll_report_xlsx(groups, total, group, layout, metric, months, date_from, date_to, sheet, split_columns=False):
    from shop.services import payroll_snapshot as ps

    group_label = ps.DIMENSION_LABELS[group]
    period = f'{date_from[0]}.{date_from[1]:02d}-{date_to[0]}.{date_to[1]:02d}'
    sheet_label = {'all': 'карт + бэлэн', 'card': 'картын цалин', 'cash': 'бэлэн цалин'}[sheet]
    lead = [group_label] + (['Албан тушаал'] if group == 'employee' else ['Ажилтны тоо'])

    def lead_values(g):
        return [g['label'], g['position'] if group == 'employee' else g['employee_count']]

    if split_columns:
        headers = lead + [f'{y}.{m:02d} {s}' for y, m in months for s in ('карт', 'бэлэн')] + ['Нийт карт', 'Нийт бэлэн', 'Нийт']

        def values(g):
            return [None if v is None else _excel_number(v) for v, _ in g['split_cells']]
        title = f'Цалингийн тайлан (карт, бэлэн тусдаа) - {ps.METRIC_LABELS[metric]} сараар, {ps.DIMENSION_BY_LABELS[group].lower()}: {period}'
    elif layout == 'pivot':
        headers = lead + [f'{y}.{m:02d}' for y, m in months] + ['Нийт']

        def values(g):
            return [_excel_number(v[metric]) if v else None for v in g['month_values']] + [_excel_number(g[metric])]
        title = f'Цалингийн тайлан ({sheet_label}) - {ps.METRIC_LABELS[metric]} сараар, {ps.DIMENSION_BY_LABELS[group].lower()}: {period}'
    else:
        headers = lead + [label for _, label, _ in ps.METRICS] + ['Дундаж олголт (ажилтан-сар)']

        def values(g):
            return [_excel_number(g[k]) for k in ps.METRIC_KEYS] + [_excel_number(g['avg_payout'])]
        title = f'Цалингийн тайлан ({sheet_label}) - {ps.DIMENSION_BY_LABELS[group].lower()}: {period}'

    # Карт, бэлэнг тусдаа мөрөөр сонгосон бол бүлэг бүрийн доор дэд мөрүүд (саарал дэвсгэргүй, догол мөртэй)
    data, fills = [], []
    for g in groups:
        data.append(lead_values(g) + values(g))
        fills.append('F9FAFB' if g.get('parts') else None)
        for part in g.get('parts', []):
            data.append([f"    {part['label']}", None if group == 'employee' else part['employee_count']] + values(part))
            fills.append(None)
    if total:
        for part in total.get('parts', []):
            data.append([f"Нийт - {part['label']}", None if group == 'employee' else part['employee_count']] + values(part))
            fills.append(None)
    totals = (['Нийт', total['employee_count'] if group != 'employee' else None] + values(total)) if total else None
    number_cols = tuple(range(3, len(headers) + 1)) + ((2,) if group != 'employee' else ())
    return _xlsx_response(
        'Цалингийн тайлан', title, headers, data, f'payroll_report_{group}_{period}.xlsx',
        widths=[26, 20] + [14] * (len(headers) - 2), totals=totals, number_cols=number_cols, row_fills=fills,
    )


@login_required
def payroll_advance(request, sheet):
    """Урьдчилгаа цалин - карт (sheet='card') эсвэл бэлэн (sheet='cash'). Бичсэн урьдчилгаа нь тухайн сарын
    цалин бодолтын (PayrollEntry.advance) "Урьдчилгаа" баганад шууд орно. Хэвлэх загвар нь Excel-ийн
    "Цалин урьдчилгаа карт/БЭЛЭН" файлын загвартай (бэлэнд Гарын үсэг баганад бэлэн цалингийн данс)."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from django.contrib import messages
    from shop.models import PayrollEntry
    from shop.services.payroll import entry_has_values
    from shop.services.payroll_snapshot import get_close, get_payroll

    is_card = sheet == PayrollEntry.SHEET_CARD
    url_name = 'shop:payroll_advance_card' if is_card else 'shop:payroll_advance_cash'
    year, month = _month_from_request(request.POST if request.method == 'POST' else request.GET)
    close = get_close(year, month, sheet)
    prev_year, prev_month = (year, month - 1) if month > 1 else (year - 1, 12)

    def set_advance(code, value):
        """Урьдчилгааг бичнэ; бүх гар утга нь 0 болсон мөрийг устгана (цалин бодолтын хадгалалттай адил)."""
        entry = PayrollEntry.objects.filter(employee_code=code, year=year, month=month, sheet=sheet).first()
        if entry is None:
            if value:
                PayrollEntry.objects.create(
                    employee_code=code, year=year, month=month, sheet=sheet, advance=value, updated_by=request.user,
                )
            return
        entry.advance = value
        if not entry_has_values(entry):
            entry.delete()
        else:
            entry.updated_by = request.user
            entry.save()

    if request.method == 'POST' and close:
        messages.error(request, 'Энэ сарын цалин хаагдсан тул урьдчилгааг засварлах боломжгүй. Цалин бодолтоос сарыг дахин нээнэ үү.')
        return redirect(f"{reverse(url_name)}?year={year}&month={month}")

    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'copy_prev':
            # Өмнөх сарын урьдчилгааг энэ сард урьдчилгаа бичээгүй ажилчдад хуулна
            codes = set(request.POST.getlist('employee_code'))
            current = dict(PayrollEntry.objects.filter(year=year, month=month, sheet=sheet)
                           .values_list('employee_code', 'advance'))
            copied = 0
            for code, advance in PayrollEntry.objects.filter(
                year=prev_year, month=prev_month, sheet=sheet, employee_code__in=codes,
            ).exclude(advance=0).values_list('employee_code', 'advance'):
                if not current.get(code):
                    set_advance(code, advance)
                    copied += 1
            messages.success(request, f'Өмнөх сарын урьдчилгааг хууллаа ({copied} ажилтан).')
        else:
            codes = request.POST.getlist('employee_code')
            values = request.POST.getlist('advance')
            for i, code in enumerate(codes):
                set_advance(code, _parse_decimal(values[i] if i < len(values) else ''))
            messages.success(request, 'Урьдчилгаа хадгалагдлаа. Цалин бодолтын Урьдчилгаа баганад орсон.')
        return redirect(f"{reverse(url_name)}?year={year}&month={month}")

    rows, totals, setting, close = get_payroll(year, month, sheet)
    prev_advances = dict(
        PayrollEntry.objects.filter(year=prev_year, month=prev_month, sheet=sheet).values_list('employee_code', 'advance')
    )
    for row in rows:
        row['prev_advance'] = prev_advances.get(row['employee_code']) or 0
    print_rows = [r for r in rows if r['advance']]
    title_kind = 'УРЬДЧИЛГАА ЦАЛИН' if is_card else 'БЭЛЭНГИЙН УРЬДЧИЛГАА ЦАЛИН'
    context = {
        **_month_nav_context(year, month),
        'sheet': sheet,
        'is_card': is_card,
        'active_page': 'payroll_advance_card' if is_card else 'payroll_advance_cash',
        'title': 'Картын урьдчилгаа цалин' if is_card else 'Бэлэн урьдчилгаа цалин',
        'close': close,
        'payroll_url_name': 'shop:payroll_card' if is_card else 'shop:payroll_cash',
        'rows': rows,
        'total_advance': totals['advance'],
        'has_prev': any(prev_advances.values()),
        'setting': setting,
        'print_rows': print_rows,
        'print_title': f'{year} ОНЫ {month:02d} САРЫН {setting.advance_day} {title_kind}',
        'approver': setting.card_approver if is_card else setting.cash_approver,
    }
    return render(request, 'shop/payroll_advance.html', context)


@login_required
def payroll_settings(request):
    """Цалин бодолтын тохиргоо: татварын хувь хэмжээ, ХХОАТ-ын хөнгөлөлтийн шатлал, ажилтан бүрийн
    олгох хэлбэр (карт/бэлэн/хосолсон) болон үндсэн цалин, удаан жилийн нэмэгдэл, НДШ-ийн код/хувь."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from django.contrib import messages
    from django.db.models import Count
    from shop.models import AttendancePositionRule, EmployeePayProfile, NdshCodeRate, PayrollSetting, PayrollTaxCredit
    from shop.services.timesheet import ordered_active_employees

    SATURDAY_BONUS_DEFAULT = EmployeePayProfile._meta.get_field('saturday_bonus_percent').default

    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'settings':
            setting = PayrollSetting.get()
            for field in ('ndsh_employee_rate', 'ndsh_employer_rate', 'pit_rate', 'saturday_hours_per_day'):
                setattr(setting, field, _parse_decimal(request.POST.get(field), default=getattr(setting, field)))
            for field in ('ndsh_max_base', 'transport_per_day'):
                setattr(setting, field, _parse_money(request.POST.get(field), default=getattr(setting, field)))
            for field in ('company_name', 'card_approver', 'cash_approver', 'prepared_by'):
                if field in request.POST:
                    setattr(setting, field, request.POST.get(field, '').strip())
            try:
                setting.advance_day = min(max(int(request.POST.get('advance_day', setting.advance_day)), 1), 31)
            except (TypeError, ValueError):
                pass
            setting.save()

            # НДШ-ийн код тус бүрийн хувь (хоёулаа хоосон бол анхдагч хувь хэрэглэгдэнэ)
            for code, employee_rate, employer_rate in zip(
                request.POST.getlist('code'), request.POST.getlist('code_employee_rate'), request.POST.getlist('code_employer_rate'),
            ):
                employee_rate = _parse_decimal(employee_rate, default=None)
                employer_rate = _parse_decimal(employer_rate, default=None)
                if employee_rate is None and employer_rate is None:
                    NdshCodeRate.objects.filter(code=code).delete()
                else:
                    NdshCodeRate.objects.update_or_create(
                        code=code, defaults={'employee_rate': employee_rate, 'employer_rate': employer_rate},
                    )
            messages.success(request, 'Татварын тохиргоо хадгалагдлаа.')

        elif action == 'credits':
            new_credits = []
            for income_to, credit in zip(request.POST.getlist('income_to'), request.POST.getlist('credit_amount')):
                if not (income_to or '').strip() and not (credit or '').strip():
                    continue
                new_credits.append(PayrollTaxCredit(
                    income_to=_parse_money(income_to, default=None),
                    credit_amount=_parse_money(credit),
                ))
            PayrollTaxCredit.objects.all().delete()
            PayrollTaxCredit.objects.bulk_create(new_credits)
            messages.success(request, f'ХХОАТ-ын хөнгөлөлтийн шатлал хадгалагдлаа ({len(new_credits)} шат).')

        elif action == 'profiles':
            valid_pay_types = {key for key, _ in EmployeePayProfile.PAY_CHOICES}
            fields = ['pay_type', 'cash_base_salary', 'card_shift_pay', 'cash_shift_pay',
                      'saturday_bonus_percent', 'ndsh_employee_rate', 'ndsh_employer_rate', 'cash_bank_account']
            lists = {f: request.POST.getlist(f) for f in fields}
            # Checkbox зөвхөн чеклэсэн үед (ажилтны кодоор) илгээгддэг. Албан тушаалын горимтой таарвал хадгалахгүй
            # (None) - ингэснээр горим өөрчлөгдөхөд дагаж өөрчлөгдөнө
            saturday_included = set(request.POST.getlist('saturday_in_worked_days'))
            saturday_positions = set(
                AttendancePositionRule.objects.filter(schedule_type=AttendancePositionRule.SCHEDULE_SATURDAY)
                .values_list('position_name', flat=True)
            )
            positions = dict(OpenDataEmployee.objects.values_list('id', 'positionname'))
            saved = 0
            for i, code in enumerate(request.POST.getlist('employee_code')):
                raw = {f: (lists[f][i] if i < len(lists[f]) else '') for f in fields}
                values = {
                    'pay_type': raw['pay_type'] if raw['pay_type'] in valid_pay_types else '',
                    'cash_base_salary': _parse_money(raw['cash_base_salary']),
                    'card_shift_pay': _parse_money(raw['card_shift_pay']),
                    'cash_shift_pay': _parse_money(raw['cash_shift_pay']),
                    'saturday_bonus_percent': _parse_decimal(raw['saturday_bonus_percent'], default=SATURDAY_BONUS_DEFAULT),
                    'ndsh_employee_rate': _parse_decimal(raw['ndsh_employee_rate'], default=None),
                    'ndsh_employer_rate': _parse_decimal(raw['ndsh_employer_rate'], default=None),
                    'cash_bank_account': ' '.join((raw['cash_bank_account'] or '').split())[:100],
                    'saturday_in_worked_days': (
                        None if (code in saturday_included) == (positions.get(code) in saturday_positions)
                        else code in saturday_included
                    ),
                }
                # Анхдагч утгууд (Бямбын нэмэгдлийн анхдагч хувь, горимоор Бямба) нь "тохиргоо хийсэн" гэж тооцогдохгүй
                if not values['pay_type'] and not any(
                    v for k, v in values.items() if k not in ('pay_type', 'saturday_bonus_percent', 'saturday_in_worked_days')
                ) and values['saturday_bonus_percent'] == SATURDAY_BONUS_DEFAULT and values['saturday_in_worked_days'] is None:
                    EmployeePayProfile.objects.filter(employee_code=code).delete()
                    continue
                EmployeePayProfile.objects.update_or_create(
                    employee_code=code, defaults={**values, 'updated_by': request.user},
                )
                saved += 1
            messages.success(request, f'Ажилчдын цалингийн мэдээлэл хадгалагдлаа ({saved} ажилтан).')

        elif action in ('print_columns', 'print_columns_reset'):
            from shop.services.payroll_print import SLOT_KEYS, clean_columns
            sheet = request.POST.get('sheet')
            if sheet in ('card', 'cash'):
                setting = PayrollSetting.get()
                print_columns = dict(setting.print_columns or {})
                if action == 'print_columns_reset':
                    print_columns.pop(sheet, None)
                    messages.success(request, 'Хэвлэх баганыг анхдагч байдалд буцаалаа.')
                else:
                    print_columns[sheet] = {
                        **{slot: clean_columns(request.POST.get(f'cols_{slot}', '').split(','), sheet)
                           for slot in SLOT_KEYS},
                        'separate_pages': request.POST.get('separate_pages') == '1',
                    }
                    messages.success(request, 'Хэвлэх баганууд хадгалагдлаа.')
                setting.print_columns = print_columns
                setting.save(update_fields=['print_columns', 'updated_at'])

        if action in ('print_columns', 'print_columns_reset'):
            return redirect(f"{reverse('shop:payroll_settings')}#print-columns")
        return redirect('shop:payroll_settings')

    profiles = {p.employee_code: p for p in EmployeePayProfile.objects.all()}
    shift_positions = set(
        AttendancePositionRule.objects.filter(schedule_type=AttendancePositionRule.SCHEDULE_SHIFT)
        .values_list('position_name', flat=True)
    )
    saturday_positions = set(
        AttendancePositionRule.objects.filter(schedule_type=AttendancePositionRule.SCHEDULE_SATURDAY)
        .values_list('position_name', flat=True)
    )
    employees = []
    for e in ordered_active_employees():
        profile = profiles.get(e.id)
        employees.append({
            'employee_code': e.id,
            'name': e.name,
            'positionname': e.positionname,
            'profile': profile,
            # Ээлжийн ажилтанд үндсэн цалингийн оронд нэг гарын мөнгө оруулна
            'is_shift': e.positionname in shift_positions,
            # Бямбыг хоногт оруулах эсэх: тусгайлан тохируулсан, эсвэл албан тушаалын горимоор
            'saturday_default': e.positionname in saturday_positions,
            'saturday_included': (
                profile.saturday_in_worked_days if profile and profile.saturday_in_worked_days is not None
                else e.positionname in saturday_positions
            ),
            'saturday_overridden': bool(profile and profile.saturday_in_worked_days is not None),
            # Ажилтны мэдээллээс (OpenDataEmployee) - засахгүй, харуулах л
            'card_base_salary': e.basesalary or 0,
            'ndsh_code': e.insuredtypeid or '',
            'seniority_bonus': e.d4 or 0,
        })
    # Ажилтны мэдээлэл дэх (картын цалингийн) данс - бэлэн цалингийн данс оруулахад лавлагаа болгож харуулна
    system_accounts = {
        code: ' '.join(filter(None, [bank, account]))
        for code, bank, account in OpenDataEmployee.objects.filter(id__in=[e['employee_code'] for e in employees])
        .extra(select={'bank': '"BankName"', 'account': '"BankAccountId"'}).values_list('id', 'bank', 'account')
        if account
    }
    for e in employees:
        e['system_bank_account'] = system_accounts.get(e['employee_code'], '')

    credits = list(PayrollTaxCredit.objects.all())

    # Идэвхтэй ажилчдын НДШ-ийн кодууд (давхардалгүй) + өмнө нь хувь тохируулсан кодууд
    code_rates = {r.code: r for r in NdshCodeRate.objects.all()}
    ndsh_codes = {}
    for row in (
        OpenDataEmployee.objects.filter(isreclusion='N').exclude(insuredtypeid__isnull=True).exclude(insuredtypeid='')
        .values('insuredtypeid', 'insuredtypename').annotate(count=Count('id')).order_by('insuredtypeid')
    ):
        entry = ndsh_codes.setdefault(row['insuredtypeid'], {'code': row['insuredtypeid'], 'name': row['insuredtypename'], 'count': 0})
        entry['count'] += row['count']
    for code in code_rates:
        ndsh_codes.setdefault(code, {'code': code, 'name': '', 'count': 0})
    for entry in ndsh_codes.values():
        entry['rate'] = code_rates.get(entry['code'])

    context = {
        'ndsh_codes': sorted(ndsh_codes.values(), key=lambda c: c['code']),
        'setting': PayrollSetting.get(),
        'credits': credits + [None, None],  # шинэ шат нэмэх хоосон мөрүүд
        'employees': employees,
        'pay_choices': EmployeePayProfile.PAY_CHOICES,
        'saturday_bonus_default': SATURDAY_BONUS_DEFAULT,
        'print_layouts': _payroll_print_layouts(PayrollSetting.get()),
    }
    return render(request, 'shop/payroll_settings.html', context)


def _payroll_print_layouts(setting):
    """Цалингийн тохиргооны "Хэвлэх баганууд" хэсэгт: хуудас (карт/бэлэн) бүрийн байрлал, сонгосон ба
    сонгож болох баганууд."""
    from shop.services.payroll_print import (
        PRINT_SLOTS, available_columns, column_choice_label, get_print_columns, get_separate_pages, page_counts,
    )
    layouts = []
    for sheet, title in (('card', 'Картын цалин'), ('cash', 'Бэлэн цалин')):
        columns = get_print_columns(setting, sheet)
        grouped_pages, separate_pages = page_counts(columns)
        layouts.append({
            'separate_pages': get_separate_pages(setting, sheet),
            'grouped_pages': grouped_pages,
            'separate_page_count': separate_pages,
            'sheet': sheet,
            'title': title,
            'is_custom': sheet in (setting.print_columns or {}),
            'choices': [{'key': k, 'label': column_choice_label(k, sheet)} for k in available_columns(sheet)],
            'slots': [
                {'key': slot, 'label': label, 'columns': ','.join(columns[slot])}
                for slot, _, label in PRINT_SLOTS
            ],
        })
    return layouts


@login_required
def attendance_position_settings(request):
    """Албан тушаал тус бүрийн ажил эхлэх цагийг тохируулах хуудас (Цаг бүртгэлийн тооцоонд ашиглана).

    Цаг хоосон үлдвэл тухайн албан тушаал 'бусад' дүрмээр (ирсэн цагаас хойш цайны цагтайгаа 9 цаг байх ёстой,
    дутсан минутыг хоцролтоор тооцно) тооцоологдоно."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from django.contrib import messages
    from shop.models import AttendancePositionRule

    if request.method == 'POST':
        position_names = request.POST.getlist('position_name')
        start_times = request.POST.getlist('start_time')
        schedule_types = request.POST.getlist('schedule_type')
        sort_orders = request.POST.getlist('sort_order')
        worked_equals_required = set(request.POST.getlist('worked_equals_required'))
        no_transport = set(request.POST.getlist('no_transport'))
        valid_schedules = {key for key, _ in AttendancePositionRule.SCHEDULE_CHOICES}
        for position_name, start_time_str, schedule_type, sort_order_str in zip(
            position_names, start_times, schedule_types, sort_orders
        ):
            start_time_str = (start_time_str or '').strip()
            if schedule_type not in valid_schedules:
                schedule_type = AttendancePositionRule.SCHEDULE_WEEKDAYS
            sort_order_str = (sort_order_str or '').strip()
            sort_order = int(sort_order_str) if sort_order_str.isdigit() else None
            start_time = None
            if start_time_str:
                try:
                    start_time = datetime.strptime(start_time_str, '%H:%M').time()
                except ValueError:
                    continue
            auto_worked = position_name in worked_equals_required
            skip_transport = position_name in no_transport
            if (start_time is None and schedule_type == AttendancePositionRule.SCHEDULE_WEEKDAYS
                    and sort_order is None and not auto_worked and not skip_transport):
                # Бүх утга анхдагч (8 цагийн дүрэм, Бямбад ажилладаггүй, эрэмбэгүй, унаатай) бол мөр хадгалах шаардлагагүй
                AttendancePositionRule.objects.filter(position_name=position_name).delete()
            else:
                AttendancePositionRule.objects.update_or_create(
                    position_name=position_name,
                    defaults={
                        'start_time': start_time,
                        'schedule_type': schedule_type,
                        'sort_order': sort_order,
                        'worked_equals_required': auto_worked,
                        'no_transport': skip_transport,
                    },
                )

        messages.success(request, 'Албан тушаалын ажлын цагийн тохиргоо хадгалагдлаа.')
        return redirect('shop:attendance_position_settings')

    position_names = set(
        OpenDataEmployee.objects.exclude(positionname__isnull=True)
        .exclude(positionname='')
        .values_list('positionname', flat=True)
        .distinct()
    )
    existing_rules = {r.position_name: r for r in AttendancePositionRule.objects.all()}
    position_names.update(existing_rules.keys())

    def position_sort_key(name):
        rule = existing_rules.get(name)
        order = rule.sort_order if rule and rule.sort_order is not None else None
        return (order is None, order or 0, name)

    rows = []
    for name in sorted(position_names, key=position_sort_key):
        rule = existing_rules.get(name)
        rows.append({
            'position_name': name,
            'start_time': rule.start_time.strftime('%H:%M') if rule and rule.start_time else '',
            'schedule_type': rule.schedule_type if rule else AttendancePositionRule.SCHEDULE_WEEKDAYS,
            'sort_order': rule.sort_order if rule and rule.sort_order is not None else '',
            'worked_equals_required': bool(rule and rule.worked_equals_required),
            'no_transport': bool(rule and rule.no_transport),
        })

    return render(request, 'shop/attendance_position_settings.html', {
        'rows': rows,
        'schedule_choices': AttendancePositionRule.SCHEDULE_CHOICES,
    })


@login_required
def employee_fingerprint(request):
    """Идэвхтэй ажилчдын нэрсийг OpenDataEmployee-ээс харуулж, тус бүрийн хурууны хээ таних
    төхөөрөмж дээрх кодыг (EmployeeFingerprint) засаж хадгалах боломжтой хуудас.

    Цаг бүртгэлийн тооцоо энд хадгалсан кодоор төхөөрөмжийн түүхий бүртгэлийг ажилтантай
    холбодог (data/users.xlsx файлыг орлож байгаа)."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from django.contrib import messages
    from shop.models import EmployeeFingerprint

    if request.method == 'POST':
        employee_codes = request.POST.getlist('employee_code')
        device_user_ids = request.POST.getlist('device_user_id')

        # Нэг хурууны код хоёр өөр ажилтанд зэрэг оноогдож болохгүй (нэг код = нэг бодит хүн
        # төхөөрөмж дээр) тул энэ хүсэлт дотор болон бааз дахь бусад ажилтантай мөргөлдөж
        # байгаа эсэхийг шалгаж, мөргөлдвөл хадгалахгүйгээр анхааруулна
        seen_in_request = {}
        conflicts = {}  # хурууны код -> одоогийн эзэмшигчийн ажилтны код
        saved = 0

        for employee_code, device_user_id in zip(employee_codes, device_user_ids):
            device_user_id = (device_user_id or '').strip()
            if not device_user_id:
                EmployeeFingerprint.objects.filter(employee_code=employee_code).delete()
                continue

            if seen_in_request.get(device_user_id, employee_code) != employee_code:
                conflicts[device_user_id] = seen_in_request[device_user_id]
                continue
            clash_owner = (
                EmployeeFingerprint.objects.filter(device_user_id=device_user_id)
                .exclude(employee_code=employee_code)
                .values_list('employee_code', flat=True)
                .first()
            )
            if clash_owner:
                conflicts[device_user_id] = clash_owner
                continue

            seen_in_request[device_user_id] = employee_code
            try:
                EmployeeFingerprint.objects.update_or_create(
                    employee_code=employee_code,
                    defaults={'device_user_id': device_user_id, 'updated_by': request.user},
                )
                saved += 1
            except IntegrityError:
                # DB-ийн unique constraint (зэрэгцээ хүсэлтийн race condition гэх мэт) сэргийлсэн тохиолдол
                conflicts[device_user_id] = (
                    EmployeeFingerprint.objects.filter(device_user_id=device_user_id)
                    .values_list('employee_code', flat=True)
                    .first()
                )

        if saved:
            messages.success(request, f'{saved} ажилтны хурууны код хадгалагдлаа.')
        if conflicts:
            owner_names = dict(
                OpenDataEmployee.objects.filter(id__in=[c for c in conflicts.values() if c])
                .values_list('id', 'name')
            )
            details = []
            for device_user_id, owner_code in sorted(conflicts.items()):
                if owner_code:
                    owner = f"{owner_names.get(owner_code) or 'Тодорхойгүй ажилтан'} ({owner_code})"
                else:
                    owner = 'тодорхойгүй ажилтан'
                details.append(f'{device_user_id} → {owner}')
            messages.error(
                request,
                'Дараах хурууны код өөр ажилтанд аль хэдийн оноогдсон тул хадгалагдсангүй: '
                + ', '.join(details)
                + '. Эхлээд хуучин эзэмшигчийнх нь кодыг хоослоно уу.'
            )
        # Хадгалсны дараа хэрэглэгчийн тохируулсан шүүлтүүрүүд (query string) хэвээр үлдэнэ
        url = reverse('shop:employee_fingerprint')
        query_string = request.META.get('QUERY_STRING', '')
        return redirect(f'{url}?{query_string}' if query_string else url)

    employees = list(
        OpenDataEmployee.objects.filter(isreclusion='N')
        .exclude(id__isnull=True).exclude(id='')
        .order_by('name')
    )
    # Бүх хадгалсан кодыг (ажлаас гарсан ажилчдынхыг ч) авна - тэд кодоо эзэмшсэн хэвээр байвал
    # шинэ ажилтанд тэр кодыг оноох боломжгүй тул жагсаалтад харуулж, хоослох боломж олгоно
    device_codes = dict(
        EmployeeFingerprint.objects.exclude(device_user_id='').values_list('employee_code', 'device_user_id')
    )
    active_codes = {e.id for e in employees}
    inactive_holders = list(
        OpenDataEmployee.objects.filter(id__in=set(device_codes) - active_codes).order_by('name')
    )
    # OpenDataEmployee-д огт олдохгүй ажилтны кодтой бичлэгүүд
    unknown_codes = sorted(set(device_codes) - active_codes - {e.id for e in inactive_holders})

    # Нэг хурууны код хэд хэдэн ажилтанд зэрэг оноогдсон эсэхийг (өмнөх өгөгдлийн алдаа гэх мэт)
    # илрүүлж, хуудсан дээр улаанаар тодруулж харуулна
    code_counts = {}
    for code in device_codes.values():
        if code:
            code_counts[code] = code_counts.get(code, 0) + 1
    duplicate_codes = {code for code, count in code_counts.items() if count > 1}

    rows = [
        {
            'employee_code': e.id,
            'name': e.name,
            'positionname': e.positionname,
            'device_user_id': device_codes.get(e.id, ''),
            'is_duplicate': device_codes.get(e.id) in duplicate_codes,
            'is_inactive': False,
        }
        for e in employees
    ] + [
        {
            'employee_code': e.id,
            'name': e.name,
            'positionname': e.positionname,
            'device_user_id': device_codes.get(e.id, ''),
            'is_duplicate': device_codes.get(e.id) in duplicate_codes,
            'is_inactive': True,
        }
        for e in inactive_holders
    ] + [
        {
            'employee_code': code,
            'name': 'Тодорхойгүй ажилтан',
            'positionname': '',
            'device_user_id': device_codes[code],
            'is_duplicate': device_codes[code] in duplicate_codes,
            'is_inactive': True,
        }
        for code in unknown_codes
    ]

    positions = sorted({row['positionname'] for row in rows if row['positionname']})

    return render(request, 'shop/employee_fingerprint.html', {
        'rows': rows,
        'positions': positions,
        'inactive_count': len(rows) - len(employees),
        'has_duplicates': bool(duplicate_codes),
    })


@login_required
def attendance_save(request):
    """Цаг бүртгэлийн хүснэгтэд хийсэн засварыг тусдаа жагсаалт (AttendanceRecord) болгож хадгална"""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ үйлдэл зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    if request.method != 'POST':
        return redirect('shop:attendance_calc')

    from django.contrib import messages
    from shop.services.attendance import save_attendance_row

    field_names = [
        'employee_code', 'device_user_id', 'name', 'date', 'weekday',
        'arrived_time', 'left_time', 'late_1_10', 'late_11_20', 'late_20_plus', 'customer_name', 'position_name',
        'record_id',  # хавсарсан албан тушаалын нэмэлт мөрийнх (үндсэн мөрийнх хоосон)
    ]
    lists = [request.POST.getlist(f) for f in field_names]

    saved_count = 0
    errors = []
    for values in zip(*lists):
        row = dict(zip(field_names, values))
        if not row.get('employee_code') or not row.get('date'):
            continue
        try:
            row['date'] = datetime.strptime(row['date'], '%Y-%m-%d').date()
        except ValueError:
            continue
        try:
            save_attendance_row(row, request.user)
            saved_count += 1
        except ValueError as e:
            errors.append(f"{row.get('name') or row['employee_code']} ({row['date']}): {e}")

    if saved_count:
        messages.success(request, f'{saved_count} мөр хадгалагдлаа.')
    if errors:
        # Хэт олон алдаа зэрэг гарвал мессежийг эвдэхгүйн тулд эхний хэдийг л харуулна
        shown = errors[:10]
        more = f' (мөн {len(errors) - 10} мөр)' if len(errors) > 10 else ''
        messages.error(request, 'Дараах мөрүүд буруу утгатай тул хадгалагдсангүй: ' + '; '.join(shown) + more)

    url = reverse('shop:attendance_calc')
    date_from = request.POST.get('date_from')
    date_to = request.POST.get('date_to')
    qs = []
    if date_from:
        qs.append(f'date_from={date_from}')
    if date_to:
        qs.append(f'date_to={date_to}')
    if qs:
        url += '?' + '&'.join(qs)
    return redirect(url)


def _attendance_redirect(request):
    """Цаг бүртгэлийн хуудас руу сонгосон огнооны мужаа хадгалан буцаах URL."""
    url = reverse('shop:attendance_calc')
    date_from = request.POST.get('date_from')
    date_to = request.POST.get('date_to')
    if date_from and date_to:
        url += f'?date_from={date_from}&date_to={date_to}'
    return redirect(url)


@login_required
def attendance_add(request):
    """Хуруу дарахаа мартсан гэх мэт шалтгаанаар бүртгэлгүй өдөрт ажилтны шинэ мөрийг гараар нэмнэ."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ үйлдэл зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')
    if request.method != 'POST':
        return redirect('shop:attendance_calc')

    from django.contrib import messages
    from shop.services.attendance import add_attendance_row

    # Олон мөрийг зэрэг нэмнэ (нэг/олон ажилтан x олон өдөр) - мөр бүр өөрийн албан тушаал, цагтай.
    # Алдаатай мөрийг (тухайн өдөр мөр аль хэдийн байх гэх мэт) алгасаж, бусдыг нь нэмнэ
    fields = ('employee_code', 'date', 'arrived_time', 'left_time', 'customer_name', 'position_name')
    lists = {f: request.POST.getlist(f) for f in fields}
    names = dict(OpenDataEmployee.objects.filter(id__in=lists['employee_code']).values_list('id', 'name'))

    added = 0
    errors = []
    for values in zip(*(lists[f] for f in fields)):
        row = dict(zip(fields, values))
        employee_code = (row['employee_code'] or '').strip()
        label = names.get(employee_code) or employee_code or '?'
        try:
            row_date = datetime.strptime(row['date'] or '', '%Y-%m-%d').date()
        except ValueError:
            errors.append(f'{label}: огноо буруу')
            continue
        if not employee_code:
            errors.append(f'{row_date}: ажилтан сонгоогүй')
            continue
        try:
            add_attendance_row(
                employee_code, row_date, row['arrived_time'], row['left_time'], row['customer_name'],
                request.user, position_name=row['position_name'],
            )
            added += 1
        except ValueError as e:
            errors.append(f'{label} ({row_date}): {e}')

    if added:
        messages.success(request, f'{added} мөр нэмэгдлээ.')
    if errors:
        shown = errors[:10]
        more = f' (мөн {len(errors) - 10} мөр)' if len(errors) > 10 else ''
        messages.error(request, 'Дараах мөрүүд нэмэгдсэнгүй: ' + '; '.join(shown) + more)
    if not added and not errors:
        messages.error(request, 'Нэмэх мөр алга байна.')
    return _attendance_redirect(request)


@login_required
def attendance_hide(request):
    """Мөрийг (төхөөрөмжийн бүртгэлээс автоматаар үүссэн ч) хүснэгтээс устгана - түүхий өгөгдлийг
    устгахгүй, AttendanceRecord.is_deleted тэмдэглэнэ. Хуудаснаас fetch()-ээр дуудагддаг."""
    if not request.user.is_staff:
        return JsonResponse({'error': 'forbidden'}, status=403)
    if request.method != 'POST':
        return JsonResponse({'error': 'method not allowed'}, status=405)

    from shop.services.attendance import mark_attendance_row_deleted

    employee_code = request.POST.get('employee_code')
    try:
        row_date = datetime.strptime(request.POST.get('date') or '', '%Y-%m-%d').date()
    except ValueError:
        return JsonResponse({'error': 'огноо буруу байна'}, status=400)
    if not employee_code:
        return JsonResponse({'error': 'employee_code дутуу байна'}, status=400)

    mark_attendance_row_deleted(
        employee_code, row_date, request.POST.get('name', ''), request.user,
        record_id=request.POST.get('record_id') or None,
    )
    return JsonResponse({'deleted': True})


@login_required
def attendance_delete(request):
    """Хадгалсан (AttendanceRecord) нэг мөрийг устгаж, дахин төхөөрөмж/мерчандайзерийн түүхий
    өгөгдөл дээр суурилсан автомат тооцоолол руу буцаана. Хуудаснаас fetch()-ээр дуудагддаг."""
    if not request.user.is_staff:
        return JsonResponse({'error': 'forbidden'}, status=403)
    if request.method != 'POST':
        return JsonResponse({'error': 'method not allowed'}, status=405)

    from shop.models import AttendanceRecord

    employee_code = request.POST.get('employee_code')
    date_str = request.POST.get('date')
    if not employee_code or not date_str:
        return JsonResponse({'error': 'employee_code/date дутуу байна'}, status=400)
    try:
        record_date = datetime.strptime(date_str, '%Y-%m-%d').date()
    except ValueError:
        return JsonResponse({'error': 'огноо буруу байна'}, status=400)

    # Зөвхөн үндсэн мөрийн засварыг цуцална - хавсарсан албан тушаалын нэмэлт мөрүүд хэвээр
    deleted, _ = AttendanceRecord.objects.filter(
        employee_code=employee_code, date=record_date, is_additional=False,
    ).delete()
    return JsonResponse({'deleted': bool(deleted)})


@login_required
def attendance_bulk_delete(request):
    """Сонгосон олон хадгалсан (AttendanceRecord) мөрийг нэг дор устгаж, автомат тооцоолол руу
    буцаана. Хуудаснаас fetch()-ээр дуудагддаг."""
    if not request.user.is_staff:
        return JsonResponse({'error': 'forbidden'}, status=403)
    if request.method != 'POST':
        return JsonResponse({'error': 'method not allowed'}, status=405)

    from django.db.models import Q
    from shop.models import AttendanceRecord

    employee_codes = request.POST.getlist('employee_code')
    date_strs = request.POST.getlist('date')

    q = Q()
    has_pair = False
    for employee_code, date_str in zip(employee_codes, date_strs):
        if not employee_code:
            continue
        try:
            record_date = datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            continue
        q |= Q(employee_code=employee_code, date=record_date)
        has_pair = True

    if not has_pair:
        return JsonResponse({'deleted': 0})

    deleted, _ = AttendanceRecord.objects.filter(q, is_additional=False).delete()
    return JsonResponse({'deleted': deleted})


# ---------------------------------------------------------------------------
# Авлага / өглөгийн тайлан (OpenDataRecPay) - shop.services.recpay_report-ийг үз
# ---------------------------------------------------------------------------

def _recpay_filters(request, kind):
    """GET параметрээс тайлангийн шүүлт. Огноо заагаагүй бол энэ сарын 1-нээс өнөөдөр хүртэл."""
    from shop.services.recpay_report import AGING_BASES, BALANCE_FILTERS, GROUP_BY

    params = request.GET
    today = timezone.localdate()

    def parse_date(name, default):
        try:
            return datetime.strptime(params.get(name, ''), '%Y-%m-%d').date()
        except ValueError:
            return default

    date_to = parse_date('date_to', today)
    date_from = parse_date('date_from', date_to.replace(day=1))
    if date_from > date_to:
        date_from, date_to = date_to, date_from

    def choice(name, options, default):
        value = params.get(name, default)
        return value if value in options else default

    return {
        'kind': kind,
        'date_from': date_from,
        'date_to': date_to,
        'accounts': [a for a in params.getlist('account') if a],
        'groups': [g for g in params.getlist('group') if g],
        'currency': params.get('currency', '').strip(),
        'q': params.get('q', '').strip(),
        'balance': choice('balance', BALANCE_FILTERS, 'nonzero'),
        'group_by': choice('group_by', GROUP_BY, 'customer'),
        'aging_basis': choice('aging_basis', AGING_BASES, 'doc'),
    }


@login_required
def recpay_report(request, kind):
    """Авлагын (kind='receivable') / өглөгийн (kind='payable') тайлан: үзүүлэлт, насжилт, динамик, харилцагчийн
    жагсаалт (мөр дээр дарахад гүйлгээний хуулга)."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    from shop.services.recpay_report import (
        AGING_BASES, BALANCE_FILTERS, GROUP_BY, KINDS, build_report, filter_options,
    )

    from shop.services.report_cache import cached
    filters = _recpay_filters(request, kind)
    report = cached('recpay', ['OpenDataRecPay'] + ['OpenDataCustomer'], filters, lambda: build_report(kind, filters))
    options = cached('recpay_options', ['OpenDataRecPay'], {'kind': kind}, lambda: filter_options(kind))
    query = request.GET.copy()
    context = {
        'kind': kind,
        'cfg': KINDS[kind],
        'active_page': f'{kind}_report',
        'filters': filters,
        'report': report,
        'options': options,
        'balance_choices': BALANCE_FILTERS,
        'group_by_choices': {k: v['label'] for k, v in GROUP_BY.items()},
        'aging_basis_choices': {k: v['label'] for k, v in AGING_BASES.items()},
        'query_string': query.urlencode(),
        'is_filtered': any([filters['accounts'], filters['groups'], filters['currency'], filters['q']]),
    }
    return render(request, 'shop/recpay_report.html', context)


@login_required
def recpay_ledger(request, kind):
    """Нэг харилцагчийн хугацааны гүйлгээний хуулга (JSON) - тайлангийн мөр дээр дарахад."""
    from django.http import JsonResponse
    if not request.user.is_staff:
        return JsonResponse({'error': 'forbidden'}, status=403)

    from shop.services.recpay_report import ledger

    filters = _recpay_filters(request, kind)
    customer = request.GET.get('customer', '')
    if not customer:
        return JsonResponse({'error': 'customer required'}, status=400)
    return JsonResponse(ledger(kind, filters, customer))


@login_required
def recpay_export(request, kind):
    """Авлага / өглөгийн тайланг идэвхтэй шүүлтээр Excel (.xlsx) файлаар татна."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    import openpyxl
    from io import BytesIO
    from django.http import HttpResponse
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from shop.services.recpay_report import AGING_BASES, GROUP_BY, KINDS, build_report

    filters = _recpay_filters(request, kind)
    cfg = KINDS[kind]
    report = build_report(kind, filters)
    basis = AGING_BASES[filters['aging_basis']]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = cfg['title']
    ws.cell(row=1, column=1, value=f"{cfg['title']} ({GROUP_BY[filters['group_by']]['label'].lower()})").font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value=(
        f"Хугацаа: {filters['date_from']:%Y-%m-%d} ~ {filters['date_to']:%Y-%m-%d} · Насжилт: {basis['label'].lower()}"
    )).font = Font(color='555555')

    keys = ['opening', 'increase', 'decrease', 'closing', 'b0', 'b1', 'b2', 'b3', 'b4', 'risk']
    headers = (['№', 'Нэр', 'Код', 'Бүлэг', 'Данс', 'Эхний үлдэгдэл', cfg['increase_label'], cfg['decrease_label'],
                'Эцсийн үлдэгдэл'] + [label for label, _ in basis['buckets']]
               + [basis['risk_label'], 'Сүүлийн гүйлгээ', 'Сүүлийн төлбөр'])
    head_row = 4
    head_fill = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
    thin = Side(style='thin', color='D1D5DB')
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=head_row, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(bottom=thin)
    for i, r in enumerate(report['rows'], start=1):
        values = [i, r['name'], r['code'], r['group'], r['account']] + [r[k] for k in keys] + [r['last_date'], r['last_payment']]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=head_row + i, column=col, value=value)
            if 6 <= col <= 5 + len(keys):
                cell.number_format = '#,##0;[Red]-#,##0'
    total_row = head_row + len(report['rows']) + 1
    ws.cell(row=total_row, column=2, value='Нийт').font = Font(bold=True)
    for offset in range(len(keys)):
        letter = get_column_letter(6 + offset)
        cell = ws.cell(row=total_row, column=6 + offset, value=f'=SUM({letter}{head_row + 1}:{letter}{total_row - 1})')
        cell.font = Font(bold=True)
        cell.number_format = '#,##0;[Red]-#,##0'
    widths = [5, 36, 14, 22, 26] + [16] * len(keys) + [13, 13]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=head_row + 1, column=3)
    ws.auto_filter.ref = f'A{head_row}:{get_column_letter(len(headers))}{max(total_row - 1, head_row)}'

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    filename = f"{'avlaga' if kind == 'receivable' else 'uglug'}_{filters['date_to']:%Y%m%d}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# ---------------------------------------------------------------------------
# Борлуулалтын нэмэгдэл бодох - shop.services.sales_bonus-ийг үз
# ---------------------------------------------------------------------------

def _sb_json(value):
    """Decimal-ийг JSON-д (JS-ийн тооцоонд) float болгоно."""
    if isinstance(value, dict):
        return {k: _sb_json(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_sb_json(v) for v in value]
    if isinstance(value, Decimal):
        return float(value)
    return value


def _sb_clean_inputs(raw, scheme):
    """Хөтчөөс ирсэн оруулсан утгыг цэвэрлэнэ (зөвхөн мэдэгдэх талбар, тоо эсвэл null)."""
    def num(v):
        if v is None or v == '':
            return None
        try:
            return str(Decimal(str(v)))
        except Exception:
            return None

    inputs = {'adjustment': num(raw.get('adjustment')), 'note': str(raw.get('note') or '')[:300]}
    if scheme == 'altjin_seller':
        inputs.update({'sales': num(raw.get('sales')), 'returns': num(raw.get('returns'))})
    elif scheme == 'distributor':
        inputs.update({
            'delivery_amount': num(raw.get('delivery_amount')),
            'criteria': {str(k): bool(v) for k, v in (raw.get('criteria') or {}).items()},
            'solo_days': num(raw.get('solo_days')),
            'helper_days': num(raw.get('helper_days')),
        })
    else:
        inputs['lines'] = [
            {'channel': str(line.get('channel') or '')[:200], 'rate': num(line.get('rate')) or '0',
             'sales': num(line.get('sales')), 'returns': num(line.get('returns'))}
            for line in (raw.get('lines') or []) if line.get('channel')
        ]
        if scheme == 'storekeeper':
            inputs['inventory'] = num(raw.get('inventory'))
        else:
            inputs['deduction_pct'] = num(raw.get('deduction_pct'))
            inputs['plan_pct'] = num(raw.get('plan_pct'))
            # Өөр ажилтны сувгаас шилжүүлж авсан дүн
            inputs['transfers'] = [
                {'channel': str(t.get('channel') or '')[:200], 'amount': num(t.get('amount')), 'rate': num(t.get('rate')) or '0'}
                for t in (raw.get('transfers') or []) if t.get('channel') and num(t.get('amount'))
            ]
    return inputs


def _sb_warehouses():
    """Агуулахын жагсаалт (Алтжин худалдагчийн тохиргоонд): [(код, нэр)]."""
    from django.db import connection
    with connection.cursor() as cursor:
        cursor.execute('SELECT "Id", "Name" FROM "OpenDataWarehouse" ORDER BY "Id"')
        return cursor.fetchall()


def _sb_month_editable(month_members, month_excluded, active_employees, employees, data):
    """Сараар нэмж/хасах хүснэгт бүрийн сонголт: {scheme: {'excluded': [(код, нэр)], 'candidates': [(код, нэр, тайлбар)]}}.
    Түгээгчийн нэмэх жагсаалтад тэр сард түгээлт хийсэн ажилчид (ажлаас гарсан ч) түгээлтийн дүнтэйгээ эхэнд,
    дараа нь бусад идэвхтэй ажилчид."""
    names = {e.id: (e.name, e.positionname) for e in active_employees} | {k: (e.name, e.positionname) for k, e in employees.items()}
    result = {}
    for scheme, people in month_members.items():
        in_table = {m.employee_code for m in people}
        candidates = []
        if scheme == 'distributor':
            for code, d in sorted(data['delivery'].items(), key=lambda kv: -kv[1]['amount']):
                if code in in_table:
                    continue
                name, position = names.get(code, (code, ''))
                candidates.append((code, name, f"{position or '-'} · түгээлт {d['amount']:,.0f}₮, {d['days']} өдөр"))
        listed = in_table | {c for c, _, _ in candidates}
        candidates += [(e.id, e.name, e.positionname or '-') for e in active_employees if e.id not in listed]
        result[scheme] = {
            'excluded': [(c, names.get(c, (c, ''))[0]) for c in month_excluded.get(scheme, [])],
            'candidates': candidates,
        }
    return result


@login_required
def sales_bonus(request):
    """Борлуулалтын нэмэгдэл бодох: Нярав, Худалдааны төлөөлөгч, Борлуулагч, Түгээгчийн хүснэгт (tab). Бодогдсон
    нэмэгдлийг карт / бэлэнд хуваарилж цалин бодолтын "Нэмэгдэл цалин /Борлуулалт/" баганад бичнэ (ажилтны бүх
    хүснэгтийн нийлбэр - жиш: ХТ нь няравыг орлосон бол хоёулаа). Няравыг орлосон / хамт ажилласан хүн, өдрийг
    сараар бүртгэж (SalesBonusSubstitute), тэр өдрүүдийн борлуулалт, буцаалтыг хуваарилна."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    import json
    from django.contrib import messages
    from shop.models import (
        EmployeePayProfile, PayrollEntry, PayrollSetting, SalesBonusEntry, SalesBonusMember, SalesBonusMonthMember,
        SalesBonusSubstitute,
    )
    from shop.services import sales_bonus as sb

    year, month = _month_from_request(request.POST if request.method == 'POST' else request.GET)
    setting = PayrollSetting.get()
    settings = sb.get_settings(setting)
    redirect_url = f"{reverse('shop:sales_bonus')}?year={year}&month={month}"

    if request.method == 'POST':
        action = request.POST.get('action')
        tab = request.POST.get('tab', '')
        if tab:
            redirect_url += f'#{tab}'

        if action == 'save':
            try:
                payload = json.loads(request.POST.get('payload') or '[]')
            except ValueError:
                payload = []
            members = {m.employee_code: m for m in SalesBonusMember.objects.filter(is_active=True).exclude(scheme='storekeeper')}
            storekeepers = {m.employee_code: (m, a) for m, a in sb.storekeeper_people(year, month, settings)}
            month_members = {(m.employee_code, sc): m for sc in sb.MONTH_EDITABLE_SCHEMES for m in sb.month_people(year, month, sc)}
            pay_types = dict(EmployeePayProfile.objects.values_list('employee_code', 'pay_type'))
            data = sb.month_data(year, month)
            cs_data = sb.month_data(year, month, *sb.period_range(year, month, 'cash_seller', settings))
            # Хадгалж буй мөрүүдийн шинэ утга - сувгаас шилжүүлгийг (transfers) эхлээд нэгтгэнэ
            prepared = []
            for raw in payload:
                code, scheme = str(raw.get('code')), raw.get('scheme')
                alloc = None
                if scheme == 'storekeeper':
                    if code not in storekeepers:
                        continue
                    member, alloc = storekeepers[code]
                else:
                    member = month_members.get((code, scheme)) if scheme in sb.MONTH_EDITABLE_SCHEMES else members.get(code)
                    if not member or member.scheme != scheme:
                        continue
                prepared.append((raw, code, scheme, member, alloc, _sb_clean_inputs(raw, scheme)))
            transfer_schemes = ('sales_rep', 'cash_seller')
            tmaps = {sc: sb.transfers_map(year, month, sc, {c: i for _, c, s_, _, _, i in prepared if s_ == sc})
                     for sc in transfer_schemes}
            saved = set()
            for raw, code, scheme, member, alloc, inputs in prepared:
                result = sb.compute(member, inputs, cs_data if scheme == 'cash_seller' else data, settings, alloc, tmaps.get(scheme))
                card, cash = sb.allocate(result['total'], pay_types.get(code), raw.get('cash_amount'))
                SalesBonusEntry.objects.update_or_create(
                    employee_code=code, year=year, month=month, scheme=scheme,
                    defaults={'inputs': inputs, 'total': result['total'], 'card_amount': card, 'cash_amount': cash,
                              'updated_by': request.user},
                )
                # Хувийг дараагийн саруудад анхдагч болгон хадгална (орлогчийнхоос бусад)
                if isinstance(member, SalesBonusMember) and inputs.get('lines'):
                    member.channels = [{'channel': l['channel'], 'rate': l['rate']} for l in inputs['lines']]
                    member.save(update_fields=['channels', 'updated_at'])
                saved.add(code)
            # Шилжүүлэг өөрчлөгдсөн бол сувгийн эзний хадгалсан бодолтыг дахин бодно
            for sc in transfer_schemes:
                for entry in SalesBonusEntry.objects.filter(year=year, month=month, scheme=sc).exclude(employee_code__in=saved):
                    member = members.get(entry.employee_code)
                    if not member or member.scheme != sc:
                        continue
                    result = sb.compute(member, entry.inputs, cs_data if sc == 'cash_seller' else data, settings, None, tmaps[sc])
                    if result['total'] != entry.total:
                        card, cash = sb.allocate(result['total'], pay_types.get(entry.employee_code), entry.cash_amount)
                        entry.total, entry.card_amount, entry.cash_amount, entry.updated_by = result['total'], card, cash, request.user
                        entry.save()
                        saved.add(entry.employee_code)
            for code in saved:
                sb.sync_payroll(code, year, month, request.user)
            messages.success(request, f'Борлуулалтын нэмэгдэл хадгалагдлаа ({len(saved)} ажилтан) - цалин бодолтод орсон.')

        elif action == 'member_rates':
            # Тохиргоо: ажилтан (ХТ, Борлуулагч) бүрийн суваг ба хувь. Сонголтоор тухайн сарын хадгалсан бодолтыг
            # шинэ хувиар дахин бодож цалин бодолтод шинэчилнэ (гараар зассан борлуулалт/буцаалт хэвээр)
            apply_month = request.POST.get('apply_month') == '1'
            known = set(sb.channel_names())
            pay_types = dict(EmployeePayProfile.objects.values_list('employee_code', 'pay_type'))
            data = None
            changed = 0
            for member in SalesBonusMember.objects.filter(is_active=True, scheme__in=['sales_rep', 'cash_seller']):
                code = member.employee_code
                if f'ch_{code}' not in request.POST:
                    continue
                channels = []
                for ch, rate in zip(request.POST.getlist(f'ch_{code}'), request.POST.getlist(f'rate_{code}')):
                    if ch in known and ch not in [c['channel'] for c in channels]:
                        channels.append({'channel': ch, 'rate': str(_parse_decimal(rate))})
                if channels == member.channels:
                    continue
                member.channels = channels
                member.save(update_fields=['channels', 'updated_at'])
                changed += 1
                entry = SalesBonusEntry.objects.filter(employee_code=code, year=year, month=month, scheme=member.scheme).first()
                if apply_month and entry:
                    data = sb.month_data(year, month, *sb.period_range(year, month, member.scheme, settings))
                    old = {l['channel']: l for l in (entry.inputs.get('lines') or [])}
                    entry.inputs['lines'] = [
                        {'channel': c['channel'], 'rate': c['rate'],
                         'sales': old.get(c['channel'], {}).get('sales'), 'returns': old.get(c['channel'], {}).get('returns')}
                        for c in channels
                    ]
                    result = sb.compute(member, entry.inputs, data, settings, None, sb.transfers_map(year, month, member.scheme))
                    card, cash = sb.allocate(result['total'], pay_types.get(code), entry.cash_amount)
                    entry.total, entry.card_amount, entry.cash_amount, entry.updated_by = result['total'], card, cash, request.user
                    entry.save()
                    sb.sync_payroll(code, year, month, request.user)
            messages.success(request, f'Сувгийн хувь хадгалагдлаа ({changed} ажилтан).'
                             + (' Энэ сарын хадгалсан бодолт шинэ хувиар дахин бодогдлоо.' if apply_month and changed else ''))

        elif action == 'substitute_save':
            # Нярав орлолт: үндсэн нярав байхгүй өдрүүдэд орлосон ажилтан ба өдрүүд (тухайн сарын, бусад орлолттой
            # давхцахгүй). Хадгалсан бодолтуудыг дахин бодно
            code = request.POST.get('employee_code', '')
            valid_days = set(sb.month_days(year, month))
            dates = sorted(set(d for d in request.POST.getlist('dates') if d in valid_days))
            sub_id = request.POST.get('substitute_id')
            others = SalesBonusSubstitute.objects.filter(year=year, month=month).exclude(id=sub_id or None)
            taken = {d: s.employee_code for s in others for d in (s.dates or [])}
            clash = [d for d in dates if d in taken]
            main_codes = set(SalesBonusMember.objects.filter(is_active=True, scheme='storekeeper').values_list('employee_code', flat=True))
            if not code or not dates:
                messages.error(request, 'Орлосон ажилтан болон дор хаяж нэг өдөр сонгоно уу.')
            elif code in main_codes:
                messages.error(request, 'Үндсэн няравыг орлогчоор бүртгэх боломжгүй.')
            elif clash:
                messages.error(request, f'{", ".join(d[-2:] for d in clash)}-ны өдөр өөр орлолтод бүртгэгдсэн байна - өдрүүд давхцахгүй.')
            else:
                sub = SalesBonusSubstitute.objects.filter(id=sub_id, year=year, month=month).first() if sub_id else None
                sub = sub or SalesBonusSubstitute(year=year, month=month)
                sub.employee_code, sub.mode, sub.dates = code, SalesBonusSubstitute.MODE_REPLACE, dates
                sub.note, sub.updated_by = request.POST.get('note', '').strip()[:200], request.user
                sub.save()
                sb.recompute_storekeepers(year, month, settings, request.user)
                emp = OpenDataEmployee.objects.filter(id=code).first()
                messages.success(request, f'{emp.name if emp else code}: {len(dates)} өдөр няравыг орлосон гэж бүртгэлээ.')

        elif action == 'substitute_delete':
            SalesBonusSubstitute.objects.filter(id=request.POST.get('substitute_id'), year=year, month=month).delete()
            sb.recompute_storekeepers(year, month, settings, request.user)
            messages.success(request, 'Орлолтыг устгалаа - өдрүүд нь үндсэн няравт буцлаа.')

        elif action in ('month_add', 'month_remove'):
            # Тухайн сард л хүснэгтэд ажилтан нэмэх / хасах (байнгын гишүүнчлэл өөрчлөгдөхгүй). Хасахад тэр сарын
            # хадгалсан бодолт устаж цалин бодолтоос хасагдана
            code = request.POST.get('employee_code', '')
            scheme = request.POST.get('scheme', '')
            if code and scheme in sb.MONTH_EDITABLE_SCHEMES:
                emp = OpenDataEmployee.objects.filter(id=code).first()
                emp_name = emp.name if emp else code
                is_member = SalesBonusMember.objects.filter(employee_code=code, scheme=scheme, is_active=True).exists()
                month_filter = {'year': year, 'month': month, 'scheme': scheme, 'employee_code': code}
                if action == 'month_add':
                    if is_member:
                        # Байнгын гишүүнийг энэ сард хассан байсан бол буцаана
                        SalesBonusMonthMember.objects.filter(**month_filter, excluded=True).delete()
                    else:
                        SalesBonusMonthMember.objects.update_or_create(
                            **month_filter, defaults={'excluded': False, 'updated_by': request.user})
                    messages.success(request, f'{emp_name}: {year}.{month:02d} сард нэмэгдлээ.')
                else:
                    if is_member:
                        SalesBonusMonthMember.objects.update_or_create(
                            **month_filter, defaults={'excluded': True, 'updated_by': request.user})
                    else:
                        SalesBonusMonthMember.objects.filter(**month_filter).delete()
                    SalesBonusEntry.objects.filter(employee_code=code, year=year, month=month, scheme=scheme).delete()
                    sb.sync_payroll(code, year, month, request.user)
                    messages.success(request, f'{emp_name}: {year}.{month:02d} сарын бодолтоос хасагдлаа.')

        elif action == 'add_member':
            code = request.POST.get('employee_code', '')
            scheme = request.POST.get('scheme', '')
            if code and scheme in dict(SalesBonusMember.SCHEME_CHOICES):
                emp = OpenDataEmployee.objects.filter(id=code).first()
                channels = sb.default_channels(scheme, emp.name if emp else '', sb.channel_names(), settings)
                member, created = SalesBonusMember.objects.get_or_create(
                    employee_code=code, defaults={'scheme': scheme, 'channels': channels},
                )
                if not created:
                    if member.scheme != scheme:
                        member.channels = channels
                    member.scheme, member.is_active = scheme, True
                    member.save()
                messages.success(request, f'{emp.name if emp else code} нэмэгдлээ.')

        elif action == 'remove_member':
            # Хүснэгтээс хасна; харж буй сарын тэр хүснэгтийн бодолт устаж цалин бодолтоос хасагдана
            # (өмнөх саруудын бодолт хэвээр)
            code = request.POST.get('employee_code', '')
            member = SalesBonusMember.objects.filter(employee_code=code).first()
            if member:
                member.is_active = False
                member.save(update_fields=['is_active', 'updated_at'])
                SalesBonusEntry.objects.filter(employee_code=code, year=year, month=month, scheme=member.scheme).delete()
                sb.sync_payroll(code, year, month, request.user)
            messages.success(request, f'Ажилтныг хүснэгтээс хаслаа - {year}.{month:02d} сарын нэмэгдэл нь цалин бодолтоос хасагдлаа.')

        elif action == 'settings':
            saved = dict(setting.sales_bonus_settings or {})
            saved['deduction_base_pct'] = float(_parse_decimal(request.POST.get('deduction_base_pct'), default=Decimal(settings['deduction_base_pct'])))
            saved['default_rates'] = {
                s: float(_parse_decimal(request.POST.get(f'rate_{s}'), default=Decimal(str(settings['default_rates'][s]))))
                for s in settings['default_rates']
            }
            try:
                saved['cash_seller_start_day'] = min(max(int(request.POST.get('cash_seller_start_day', 21)), 1), 28)
            except (TypeError, ValueError):
                pass
            dist = settings['distributor']
            saved['distributor'] = {
                'criteria': [
                    {**c, 'rate': float(_parse_decimal(request.POST.get(f"criteria_{c['key']}"), default=Decimal(str(c['rate']))))}
                    for c in dist['criteria']
                ],
                'solo_day': float(_parse_decimal(request.POST.get('solo_day'), default=Decimal(str(dist['solo_day'])))),
                'helper_day': float(_parse_decimal(request.POST.get('helper_day'), default=Decimal(str(dist['helper_day'])))),
            }
            # Алтжин худалдагч: шатлал (дээд хязгаар хоосон мөр - "түүнээс дээш", хувьгүй мөрийг алгасна)
            tiers = []
            for upto, rate in zip(request.POST.getlist('altjin_upto'), request.POST.getlist('altjin_rate')):
                if str(rate).strip() == '':
                    continue
                upto_value = _parse_decimal(upto.replace(',', ''), default=None) if str(upto).strip() else None
                tiers.append({'upto': float(upto_value) if upto_value is not None else None, 'rate': float(_parse_decimal(rate))})
            tiers.sort(key=lambda t: (t['upto'] is None, t['upto'] or 0))
            # Хамгийн сүүлийн шат үргэлж "түүнээс дээш" байна
            if tiers and tiers[-1]['upto'] is not None:
                tiers.append({'upto': None, 'rate': tiers[-1]['rate']})
            altjin = settings['altjin']
            saved['altjin'] = {
                'warehouse_id': request.POST.get('altjin_warehouse') or altjin['warehouse_id'],
                'mode': 'flat' if request.POST.get('altjin_mode') == 'flat' else 'marginal',
                'tiers': tiers or altjin['tiers'],
            }
            setting.sales_bonus_settings = saved
            setting.save(update_fields=['sales_bonus_settings', 'updated_at'])
            messages.success(request, 'Тохиргоо хадгалагдлаа. Хадгалаагүй сарын бодолтод шинэ утгаар бодогдоно.')

        if action != 'settings':
            from shop.models import PayrollMonthClose
            closed = PayrollMonthClose.objects.filter(year=year, month=month).values_list('sheet', flat=True)
            if closed:
                messages.warning(request, f'{year}.{month:02d} сарын цалин ({", ".join(dict(PayrollEntry.SHEET_CHOICES)[c].lower() for c in closed)}) '
                                          'хаагдсан тул энэ өөрчлөлт хаагдсан цалинд орохгүй - цалин бодолтоос сарыг дахин нээгээд хаана уу.')

        return redirect(redirect_url)

    # --- Харуулах ---
    sb.ensure_members(settings)
    data = sb.month_data(year, month)
    cs_range = sb.period_range(year, month, 'cash_seller', settings)
    cs_data = sb.month_data(year, month, *cs_range)
    members = list(SalesBonusMember.objects.filter(is_active=True))
    storekeepers = sb.storekeeper_people(year, month, settings)
    month_members = {sc: sb.month_people(year, month, sc) for sc in sb.MONTH_EDITABLE_SCHEMES}
    month_excluded = {
        sc: list(SalesBonusMonthMember.objects.filter(year=year, month=month, scheme=sc, excluded=True)
                 .values_list('employee_code', flat=True))
        for sc in sb.MONTH_EDITABLE_SCHEMES
    }
    codes = ({m.employee_code for m in members} | {m.employee_code for m, _ in storekeepers}
             | {m.employee_code for ms in month_members.values() for m in ms}
             | {c for cs in month_excluded.values() for c in cs} | set(data['delivery']))
    employees = {e.id: e for e in OpenDataEmployee.objects.filter(id__in=codes)}
    pay_types = dict(EmployeePayProfile.objects.values_list('employee_code', 'pay_type'))
    entries = {(e.employee_code, e.scheme): e for e in SalesBonusEntry.objects.filter(year=year, month=month)}
    # Хадгалаагүй сард карт / бэлэнгийн хуваарилалтыг (бэлэнд оруулах дүн) анхдагчаар өмнөх сарын хадгалсан бодолтоос
    prev_year, prev_month = (year, month - 1) if month > 1 else (year - 1, 12)
    prev_cash = {
        (e.employee_code, e.scheme): e.cash_amount
        for e in SalesBonusEntry.objects.filter(year=prev_year, month=prev_month)
    }
    pay_labels = dict(EmployeePayProfile.PAY_CHOICES)
    tmaps = {sc: sb.transfers_map(year, month, sc) for sc in ('sales_rep', 'cash_seller')}

    def person(m, alloc=None):
        emp = employees.get(m.employee_code)
        entry = entries.get((m.employee_code, m.scheme))
        inputs = entry.inputs if entry else {}
        result = sb.compute(m, inputs, cs_data if m.scheme == 'cash_seller' else data, settings, alloc, tmaps.get(m.scheme))
        pay_type = pay_types.get(m.employee_code) or ''
        return {
            'code': m.employee_code,
            'name': emp.name if emp else m.employee_code,
            'position': emp.positionname if emp else '',
            'pay_type': pay_type,
            'pay_label': pay_labels.get(pay_type, 'Олгох хэлбэргүй'),
            'is_substitute': getattr(m, 'is_substitute', False),
            'is_month_extra': getattr(m, 'is_month_extra', False),
            'saved': bool(entry),
            'cash_amount': entry.cash_amount if entry else prev_cash.get((m.employee_code, m.scheme), 0),
            'cash_from_prev': not entry and bool(prev_cash.get((m.employee_code, m.scheme))),
            'member_channels': m.channels or [],
            'inputs': inputs,
            'result': result,
        }

    tabs = []
    for scheme, label in SalesBonusMember.SCHEME_CHOICES:
        if scheme == 'storekeeper':
            people = [person(m, alloc) for m, alloc in storekeepers]
        else:
            scheme_members = month_members[scheme] if scheme in month_members else [m for m in members if m.scheme == scheme]
            ordered = sorted(scheme_members,
                             key=lambda m: (m.sort_order, employees[m.employee_code].name if m.employee_code in employees else ''))
            people = [person(m) for m in ordered]
        tabs.append({'key': scheme, 'label': label, 'people': people,
                     'total': sum((p['result']['total'] for p in people), Decimal(0))})

    member_codes = {m.employee_code for m in members}
    active_employees = list(OpenDataEmployee.objects.filter(isreclusion='N').order_by('name'))
    names = {e.id: e.name for e in active_employees} | {k: v.name for k, v in employees.items()}
    substitutes = [
        {'id': s.id, 'code': s.employee_code, 'name': names.get(s.employee_code, s.employee_code), 'mode': s.mode,
         'mode_label': s.get_mode_display(), 'dates': s.dates, 'note': s.note}
        for s in SalesBonusSubstitute.objects.filter(year=year, month=month)
    ]
    start, end = sb.month_range(year, month)
    days = [{'iso': d, 'day': int(d[-2:]), 'dow': datetime.fromisoformat(d).isoweekday()} for d in sb.month_days(year, month)]
    main_names = [names.get(m.employee_code, m.employee_code) for m, _ in storekeepers if not getattr(m, 'is_substitute', False)]
    context = {
        **_month_nav_context(year, month),
        'active_page': 'sales_bonus',
        'tabs': tabs,
        'settings': settings,
        'setting': setting,
        'available_employees': [(e.id, e.name, e.positionname) for e in active_employees if e.id not in member_codes],
        'all_employees': [(e.id, e.name, e.positionname) for e in active_employees],
        'month_editable': _sb_month_editable(month_members, month_excluded, active_employees, employees, data),
        'channel_list': sb.channel_names(),
        'rate_tabs': [t for t in tabs if t['key'] in ('sales_rep', 'cash_seller')],
        'warehouses': _sb_warehouses(),
        # Тохиргоонд хоосон нэг мөр нэмж шинэ шат оруулах боломжтой
        'altjin_tier_rows': settings['altjin']['tiers'] + [{'upto': '', 'rate': ''}],
        'substitutes': substitutes,
        'month_days': days,
        'main_storekeepers': main_names,
        'period_label': f'{year} он {start:%m.%d}-{end:%m.%d}',
        'state_json': _sb_json({
            'tabs': [{'key': t['key'], 'people': [
                {k: p[k] for k in ('code', 'name', 'pay_type', 'cash_amount', 'cash_from_prev', 'inputs', 'is_substitute', 'is_month_extra', 'saved')}
                | {'result': p['result']} for p in t['people']
            ]} for t in tabs],
            'channels': {
                name: {'sales': data['credits'].get(name, 0), 'returns': data['returns'].get(name, 0),
                       'accounts': data['channel_accounts'].get(name, [])}
                for name in sb.channel_names()
            },
            # Борлуулагчийн хүснэгт өөр хугацаагаар (өмнөх сарын 21 - энэ сарын 20) бодогдоно
            'channels_cash_seller': {
                name: {'sales': cs_data['credits'].get(name, 0), 'returns': cs_data['returns'].get(name, 0),
                       'accounts': cs_data['channel_accounts'].get(name, [])}
                for name in sb.channel_names()
            },
            'periods': {scheme: sb.period_label(*(cs_range if scheme == 'cash_seller' else sb.month_range(year, month)))
                        for scheme, _ in SalesBonusMember.SCHEME_CHOICES},
            'settings': settings,
            'substitutes': substitutes,
            'days': days,
        }),
    }
    return render(request, 'shop/sales_bonus.html', context)


# ---------------------------------------------------------------------------
# Өгөгдөл шинэчлэх - MSSQL-ээс нэг хүснэгт хурдан татах (datamigration.sync-ийг үз)
# ---------------------------------------------------------------------------

# Хамгийн их шинэчилдэг хүснэгтүүд - хуудасны дээд хэсэгт
DATA_SYNC_FAVORITES = [
    ('OpenDataEmployee', 'Ажилтны мэдээлэл (үндсэн цалин, НДШ, хадгаламж г.м.)'),
    ('OpenDataSale', 'Борлуулалт'),
    ('OpenDataRecPay', 'Авлага, өглөгийн гүйлгээ'),
    ('OpenDataCustomer', 'Харилцагч'),
    ('OpenDataItem', 'Бараа'),
    ('OpenDataDelivery', 'Түгээлт'),
    ('OpenDataSaleRefund', 'Борлуулалтын буцаалт'),
]


def _data_sync_run(log_id):
    """Ард (thread) нэг хүснэгт татаж, явц/үр дүнг OpenDataSyncLog-д бичнэ."""
    from django.db import connection as db_connection
    from shop.models import OpenDataSyncLog
    from datamigration.sync import sync_view

    log = OpenDataSyncLog.objects.get(id=log_id)
    try:
        result = sync_view(log.table_name, log=lambda text: OpenDataSyncLog.objects.filter(id=log_id).update(message=text))
        steps = result['steps']
        OpenDataSyncLog.objects.filter(id=log_id).update(
            status=OpenDataSyncLog.STATUS_SUCCESS, rows=result['rows'], seconds=result['seconds'],
            database=result['database'], finished_at=timezone.now(),
            message=f"Татах {steps['fetch']}с · бичих {steps['write']}с · индекс, кэш {steps['post']}с",
        )
    except Exception as e:  # noqa: BLE001 - алдааг хуудсанд харуулна
        OpenDataSyncLog.objects.filter(id=log_id).update(
            status=OpenDataSyncLog.STATUS_ERROR, message=str(e)[:1000], finished_at=timezone.now())
    finally:
        db_connection.close()


def _data_sync_log_json(log):
    return {
        'id': log.id, 'table': log.table_name, 'status': log.status, 'status_label': log.get_status_display(),
        'rows': log.rows, 'seconds': log.seconds, 'message': log.message, 'database': log.database,
        'started': timezone.localtime(log.started_at).strftime('%Y-%m-%d %H:%M:%S'),
        'finished': timezone.localtime(log.finished_at).strftime('%Y-%m-%d %H:%M:%S') if log.finished_at else '',
        'user': (log.user.get_full_name() or log.user.username) if log.user else 'Систем',
    }


@login_required
def data_sync(request):
    """Өгөгдөл шинэчлэх: MSSQL-ээс сонгосон нэг хүснэгтийг шаардлагатай үед хурдан татна (зөвхөн superuser).
    Өдөр тутмын хуваарьт бүрэн синк (import_views) хэвээр."""
    if not request.user.is_superuser:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн админд зориулагдсан.')
        return redirect('shop:dashboard')

    from django.db import connection
    from shop.models import OpenDataSyncLog

    with connection.cursor() as cursor:
        cursor.execute(
            '''
            SELECT c.relname, COALESCE(s.n_live_tup, c.reltuples::bigint), pg_total_relation_size(c.oid)
            FROM pg_class c LEFT JOIN pg_stat_user_tables s ON s.relid = c.oid
            WHERE c.relkind = 'r' AND c.relname LIKE 'OpenData%%' AND c.relname NOT LIKE '%%__sync'
            ORDER BY c.relname
            ''')
        tables = {name: {'name': name, 'rows': rows, 'size': size} for name, rows, size in cursor.fetchall()}
    # MSSQL-д байгаа ч PostgreSQL-д хараахан татаагүй view-ууд (жагсаалт кэштэй бол)
    from django.core.cache import cache
    for view, db_name in cache.get('opendata_sync:views') or []:
        tables.setdefault(view, {'name': view, 'rows': None, 'size': None, 'database': db_name})

    last = {}
    for log in OpenDataSyncLog.objects.filter(table_name__in=list(tables)).select_related('user').order_by('-started_at')[:500]:
        last.setdefault(log.table_name, _data_sync_log_json(log))
    favorites = dict(DATA_SYNC_FAVORITES)
    rows = [{**t, 'label': favorites.get(t['name'], ''), 'last': last.get(t['name'])} for t in tables.values()]
    rows.sort(key=lambda r: (r['name'] not in favorites, list(favorites).index(r['name']) if r['name'] in favorites else 0, r['name']))
    history = [_data_sync_log_json(l) for l in OpenDataSyncLog.objects.select_related('user')[:25]]
    return render(request, 'shop/data_sync.html', {
        'active_page': 'data_sync',
        'tables_json': rows,
        'history_json': history,
    })


@login_required
def data_sync_start(request):
    """Нэг хүснэгт татахыг ард эхлүүлнэ (JSON). Тухайн хүснэгт аль хэдийн татагдаж байвал тэр ажлыг буцаана."""
    import threading
    from datetime import timedelta as _td
    from django.http import JsonResponse
    from shop.models import OpenDataSyncLog

    if not request.user.is_superuser or request.method != 'POST':
        return JsonResponse({'error': 'forbidden'}, status=403)
    table = request.POST.get('table', '').strip()
    if not table or not table.replace('_', '').isalnum():
        return JsonResponse({'error': 'Хүснэгтийн нэр буруу'}, status=400)
    running = OpenDataSyncLog.objects.filter(
        table_name=table, status=OpenDataSyncLog.STATUS_RUNNING, started_at__gte=timezone.now() - _td(minutes=30)).first()
    if running:
        return JsonResponse({'log': _data_sync_log_json(running), 'already': True})
    log = OpenDataSyncLog.objects.create(table_name=table, user=request.user, message='Эхлүүлж байна...')
    threading.Thread(target=_data_sync_run, args=(log.id,), daemon=True).start()
    return JsonResponse({'log': _data_sync_log_json(log)})


@login_required
def data_sync_status(request):
    """Татах ажлын явц (JSON) - хуудас секунд тутам асууна."""
    from django.http import JsonResponse
    from shop.models import OpenDataSyncLog

    if not request.user.is_superuser:
        return JsonResponse({'error': 'forbidden'}, status=403)
    ids = [int(i) for i in request.GET.get('ids', '').split(',') if i.isdigit()]
    logs = OpenDataSyncLog.objects.filter(id__in=ids).select_related('user')
    return JsonResponse({'logs': [_data_sync_log_json(l) for l in logs]})


@login_required
def data_sync_views(request):
    """MSSQL-ийн бүх view-ийн жагсаалтыг шинэчилнэ (JSON) - шинээр нэмэгдсэн view-г харах."""
    from django.http import JsonResponse
    from django.core.cache import cache
    from datamigration.sync import list_views

    if not request.user.is_superuser:
        return JsonResponse({'error': 'forbidden'}, status=403)
    cache.delete('opendata_sync:views')
    try:
        views = list_views()
    except Exception as e:  # noqa: BLE001
        return JsonResponse({'error': str(e)}, status=500)
    return JsonResponse({'views': [{'name': v, 'database': d} for v, d in views]})


# ---------------------------------------------------------------------------
# Ачаа буулгалт - татан авалтын чингэлэг буулгалтад оролцсон хүмүүсийн ажлын хөлс
# ---------------------------------------------------------------------------

def _fee(value):
    return _parse_decimal((value or '').replace(',', '').strip())


@login_required
def unloading(request):
    """Ачаа буулгалт: сарын татан авалтууд (OpenDataLandedCost - илгээгч, бараа, тоо ширхэг), ачаа бүрт оролцсон
    ажилчид (үүрэг, хүн бүрийн ажлын хөлс), гаднаас хөлсөлсөн хүмүүс (тоо × хөлс); сарын нэгтгэл (буулгасан огноогоор)."""
    if not request.user.is_staff:
        from django.contrib import messages
        messages.warning(request, 'Энэ хуудас зөвхөн ажилтанд зориулагдсан.')
        return redirect('shop:home')

    import json
    from django.contrib import messages
    from shop.models import UnloadingEvent, UnloadingSettings, UnloadingWorker
    from shop.services import unloading as svc

    year, month = _month_from_request(request.POST if request.method == 'POST' else request.GET)
    redirect_url = f"{reverse('shop:unloading')}?year={year}&month={month}"
    settings_obj = UnloadingSettings.load()

    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'settings':
            for field in ('customs_code', 'check_code', 'receive_code'):
                setattr(settings_obj, field, request.POST.get(field, '').strip())
            for field in ('key_fee', 'worker_fee', 'outside_fee'):
                setattr(settings_obj, field, _fee(request.POST.get(field)))
            settings_obj.save()
            messages.success(request, 'Анхдагч тохиргоо хадгалагдлаа. Шинээр бүртгэх ачаанд хэрэглэгдэнэ.')
            return redirect(redirect_url)
        try:
            doc_id = int(request.POST.get('document_id') or 0)
        except ValueError:
            doc_id = 0
        if action == 'delete':
            UnloadingEvent.objects.filter(document_id=doc_id).delete()
            messages.success(request, f"{request.POST.get('document_number', '')} ачааны бүртгэлийг устгалаа.")
            return redirect(redirect_url)
        if action == 'save' and doc_id:
            number, doc_date, vendor = svc.document_info(doc_id)
            number = number or request.POST.get('document_number', '').strip()
            try:
                unload_date = datetime.strptime(request.POST.get('unload_date', ''), '%Y-%m-%d').date()
            except ValueError:
                unload_date = doc_date or timezone.localdate()
            try:
                outside = max(int(request.POST.get('outside_count') or 0), 0)
            except ValueError:
                outside = 0
            event, _ = UnloadingEvent.objects.update_or_create(document_id=doc_id, defaults={
                'document_number': number, 'document_date': doc_date, 'vendor_name': vendor or '', 'unload_date': unload_date,
                'outside_count': outside, 'outside_fee': _fee(request.POST.get('outside_fee')),
                'note': request.POST.get('note', '').strip()[:300], 'updated_by': request.user,
            })
            event.workers.all().delete()
            roles = dict(UnloadingWorker.ROLE_CHOICES)
            seen = set()
            for role, code, fee in zip(request.POST.getlist('worker_role'), request.POST.getlist('worker_code'),
                                       request.POST.getlist('worker_fee')):
                code = code.strip()
                if not code:
                    continue
                if code in seen:  # нэг ачаанд нэг хүн нэг л мөрөөр
                    messages.warning(request, f'{code} ажилтан давхар сонгогдсон тул эхний мөрийг авлаа.')
                    continue
                seen.add(code)
                UnloadingWorker.objects.create(event=event, employee_code=code, role=role if role in roles else 'unload', fee=_fee(fee))
            total = sum((w.fee for w in event.workers.all()), Decimal(0)) + event.outside_amount
            messages.success(request, f'{number}: {len(seen) + outside} хүн, нийт {money_mn(total)} ажлын хөлс хадгалагдлаа.')
        return redirect(redirect_url + (f'#doc-{doc_id}' if doc_id else ''))

    data = svc.month_data(year, month)
    employees = list(OpenDataEmployee.objects.filter(isreclusion='N').order_by('name').values('id', 'name', 'positionname'))
    active = {e['id'] for e in employees}
    # Өмнө бүртгэгдсэн, одоо гарсан ажилтныг сонголтод нэмнэ (засварлахад алга болохгүй)
    used = {w['code'] for d in data['docs'] if d['event'] for w in d['workers']} - active
    employees += [{'id': c, 'name': n, 'positionname': p} for c, (n, p) in svc.employee_names(used).items()]

    # Маягтын анхны мөрүүд: бүртгэлтэй бол хадгалсан, үгүй бол анхдагч (гааль, хяналт, орлого + нэг хоосон ажилтан)
    defaults = [{'role': role, 'code': code, 'fee': float(settings_obj.key_fee)} for role, code in settings_obj.role_codes().items()]
    defaults.append({'role': 'unload', 'code': '', 'fee': float(settings_obj.worker_fee)})
    forms = {}
    for d in data['docs']:
        if d['event']:
            forms[str(d['id'])] = {'rows': [{'role': w['role'], 'code': w['code'], 'fee': float(w['fee'])} for w in d['workers']],
                                   'outside_count': d['event'].outside_count, 'outside_fee': float(d['event'].outside_fee)}
        else:
            forms[str(d['id'])] = {'rows': defaults, 'outside_count': 0, 'outside_fee': float(settings_obj.outside_fee)}

    context = {
        **_month_nav_context(year, month),
        **data,
        'active_page': 'unloading',
        'roles': UnloadingWorker.ROLE_CHOICES,
        'settings_obj': settings_obj,
        'employees': employees,
        'unloading_js': {
            'employees': [{'id': e['id'], 'label': f"{e['name']} · {e['positionname'] or '-'}"} for e in employees],
            'roles': [{'id': r, 'label': l} for r, l in UnloadingWorker.ROLE_CHOICES],
            'role_fees': {r: float(settings_obj.role_fee(r)) for r, _ in UnloadingWorker.ROLE_CHOICES},
            'forms': forms,
        },
        'printed_at': timezone.localtime(),
    }
    return render(request, 'shop/unloading.html', context)


@login_required
def unloading_export(request):
    """Ачаа буулгалтын сарын тайлан Excel-ээр: Нэгтгэл, Ачаа (оролцогч бүрээр), Бараа."""
    if not request.user.is_staff:
        return redirect('shop:home')

    import openpyxl
    from io import BytesIO
    from django.http import HttpResponse
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from shop.services import unloading as svc

    year, month = _month_from_request(request.GET)
    data = svc.month_data(year, month)
    head_fill = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
    thin = Side(style='thin', color='D1D5DB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def sheet(ws, title, headers, widths, rows, total=None):
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
        for col, (h, w) in enumerate(zip(headers, widths), start=1):
            c = ws.cell(row=3, column=col, value=h)
            c.font, c.fill, c.border = Font(bold=True), head_fill, border
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = w
        r = 3
        for r, values in enumerate(rows, start=4):
            for col, value in enumerate(values, start=1):
                c = ws.cell(row=r, column=col, value=value)
                c.border = border
                if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
                    c.number_format = '#,##0.##' if headers[col - 1] == 'Тоо ширхэг' else '#,##0'
                c.alignment = Alignment(vertical='top', wrap_text=isinstance(value, str))
        if total is not None:
            label_col, value_col, value = total
            ws.cell(row=r + 1, column=label_col, value='Нийт').font = Font(bold=True)
            c = ws.cell(row=r + 1, column=value_col, value=value)
            c.font, c.number_format = Font(bold=True), '#,##0'
        ws.freeze_panes = 'A4'

    wb = openpyxl.Workbook()
    rows = [[i, r['name'], r['position'], r['loads'], r['amount']] for i, r in enumerate(data['summary_rows'], start=1)]
    if data['outside_people']:
        rows.append(['', 'Гаднаас хөлсөлсөн', f"{data['outside_people']} хүн", '', data['outside_amount']])
    sheet(wb.active, f"{data['period_label']} - ачаа буулгалтын ажлын хөлс (нэгтгэл)",
          ['№', 'Ажилтан', 'Албан тушаал', 'Ачааны тоо', 'Ажлын хөлс'], [6, 28, 24, 12, 16], rows, (2, 5, data['grand_total']))
    wb.active.title = 'Нэгтгэл'

    rows = []
    for d in data['month_docs']:
        ev = d['event']
        base = [ev.unload_date, d['number'], d['date'], d['vendor']]
        for w in d['workers']:
            rows.append(base + [w['role_label'], w['name'], w['position'], 1, w['fee'], w['fee']])
        if ev.outside_count:
            rows.append(base + ['Гаднаас хөлсөлсөн', '', '', ev.outside_count, ev.outside_fee, ev.outside_amount])
        rows.append(['', d['number'], '', 'Ачааны дүн', '', '', '', d['people'], '', d['total']])
    ws = wb.create_sheet('Ачаа')
    sheet(ws, f"{data['period_label']} - ачаа тус бүрийн оролцогчид",
          ['Буулгасан', 'Дугаар', 'Татан авалт', 'Илгээгч', 'Үүрэг', 'Ажилтан', 'Албан тушаал', 'Хүн', 'Нэг хүний хөлс', 'Дүн'],
          [12, 12, 12, 30, 20, 22, 20, 7, 14, 14], rows, (4, 10, data['grand_total']))
    for row in ws.iter_rows(min_row=4):
        for c in row:
            if hasattr(c.value, 'year'):
                c.number_format = 'yyyy-mm-dd'
        if row[3].value == 'Ачааны дүн':
            for c in row:
                c.font = Font(bold=True)

    rows = [[d['number'], d['date'], d['vendor'], it['name'], it['measure'], it['qty']]
            for d in data['month_docs'] for it in d['items']]
    ws = wb.create_sheet('Бараа')
    sheet(ws, f"{data['period_label']} - буулгасан ачааны бараа",
          ['Дугаар', 'Татан авалт', 'Илгээгч', 'Барааны нэр', 'Хэмжих нэгж', 'Тоо ширхэг'], [12, 12, 30, 50, 12, 14], rows)
    for row in ws.iter_rows(min_row=4, max_col=2):
        if hasattr(row[1].value, 'year'):
            row[1].number_format = 'yyyy-mm-dd'

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="unloading_{year}_{month:02d}.xlsx"'
    return response


def money_mn(value):
    """1234567 -> '1,234,567₮' (мессежид)."""
    return f'{Decimal(value or 0):,.0f}₮'
