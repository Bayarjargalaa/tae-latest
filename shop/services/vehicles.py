"""
Тээврийн хэрэгсэл (Хангамж): шатахуун, сэлбэг хэрэгслийн нийтлэг хэсэг - хугацааны шүүлт, машин/жолоочийн
сонголт, тайлангууд, Excel экспорт.

Тайлан бүр нэг бүтэцтэй ({'title', 'meta', 'tables': [{'title', 'columns', 'rows', 'totals'}]}) тул хуудсанд
хэвлэх хэлбэрээр болон Excel-д ижилхэн гарна. Багана: {'key', 'label', 'fmt'} - fmt нь text | int | money |
dec | date | datetime.
"""
import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal

from django.utils import timezone

ZERO = Decimal('0')
CENT = Decimal('0.01')

OUR_PLATES = '__ours'
PERSONAL = '__personal'
NO_PLATE = '__none'
PERSONAL_NOTE = 'Хувийн машин'


# --- Хугацаа ---

def _parse_date(value):
    try:
        return datetime.strptime(value or '', '%Y-%m-%d').date()
    except ValueError:
        return None


def month_end(d):
    nxt = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    return nxt - timedelta(days=1)


def add_months(d, n):
    m = d.year * 12 + d.month - 1 + n
    return date(m // 12, m % 12 + 1, 1)


def period_presets(today=None):
    today = today or timezone.localdate()
    month = today.replace(day=1)
    prev = add_months(month, -1)
    q_start = date(today.year, (today.month - 1) // 3 * 3 + 1, 1)
    return [
        ('Энэ сар', month, today),
        ('Өмнөх сар', prev, month_end(prev)),
        ('Энэ улирал', q_start, today),
        ('Энэ жил', date(today.year, 1, 1), today),
        ('Өмнөх жил', date(today.year - 1, 1, 1), date(today.year - 1, 12, 31)),
        ('Сүүлийн 12 сар', add_months(month, -11), today),
    ]


def period(params, default_from, default_to):
    """GET-ээс хугацаа: month=YYYY-MM эсвэл date_from/date_to. Буцаах: (эхлэх, дуусах)."""
    m = re.match(r'^(\d{4})-(\d{2})$', params.get('month', ''))
    if m:
        start = date(int(m.group(1)), int(m.group(2)), 1)
        return start, month_end(start)
    d_from = _parse_date(params.get('date_from')) or default_from
    d_to = _parse_date(params.get('date_to')) or default_to
    if d_to < d_from:
        d_from, d_to = d_to, d_from
    return d_from, d_to


def aware_range(d_from, d_to):
    """[эхлэх 00:00, дуусах+1 00:00) орон нутгийн цагаар."""
    tz = timezone.get_current_timezone()
    return (timezone.make_aware(datetime(d_from.year, d_from.month, d_from.day), tz),
            timezone.make_aware(datetime(d_to.year, d_to.month, d_to.day) + timedelta(days=1), tz))


def single_month(d_from, d_to):
    """Хугацаа яг нэг бүтэн сар бол тэр сарын 1-ний өдөр, үгүй бол None."""
    return d_from if d_from.day == 1 and d_to == month_end(d_from) else None


def period_label(d_from, d_to):
    if single_month(d_from, d_to):
        return f'{d_from:%Y} оны {d_from.month}-р сар'
    return f'{d_from:%Y.%m.%d} - {d_to:%Y.%m.%d}'


def months_between(d_from, d_to):
    out, m = [], d_from.replace(day=1)
    while m <= d_to:
        out.append(m)
        m = add_months(m, 1)
    return out


# --- Машин, жолооч ---

def vehicle_options():
    """[{'plate', 'label', 'ours'}] - манай машин (OpenDataAsset + Vehicle) эхэнд, дараа нь шатахуун/сэлбэгт
    бичигдсэн бусад дугаар. ours=True бол км заалт заавал."""
    from shop.models import FuelTransaction, Vehicle, VehiclePartRecord
    from shop.services.delivery_trips import vehicle_options as asset_vehicles

    options = {}
    for v in asset_vehicles():
        if not v['used']:
            options[v['plate']] = {'plate': v['plate'], 'label': v['label'], 'ours': True}
    for v in Vehicle.objects.all():
        label = f'{v.make} {v.model}'.strip()
        if v.is_active:
            options[v.plate] = {'plate': v.plate, 'label': label or options.get(v.plate, {}).get('label', ''), 'ours': True}
        else:
            options.pop(v.plate, None)
    for model in (FuelTransaction, VehiclePartRecord):
        for plate in model.objects.exclude(vehicle_plate='').values_list('vehicle_plate', flat=True).distinct():
            options.setdefault(plate, {'plate': plate, 'label': '', 'ours': False})
    return sorted(options.values(), key=lambda o: (not o['ours'], o['plate']))


def our_plates():
    return {v['plate'] for v in vehicle_options() if v['ours']}


def driver_options():
    """[{'code', 'name'}] - идэвхтэй ажилчид + шатахуун/сэлбэгт өмнө бичсэн бусад нэр."""
    from datamigration.models import OpenDataEmployee
    from shop.models import FuelTransaction, VehiclePartRecord

    options = [{'code': r['id'], 'name': (r['name'] or '').strip()}
               for r in OpenDataEmployee.objects.filter(isreclusion='N').order_by('name').values('id', 'name')
               if (r['name'] or '').strip()]
    known = {o['name'] for o in options}
    for model in (FuelTransaction, VehiclePartRecord):
        for name in model.objects.exclude(driver_name='').values_list('driver_name', flat=True).distinct():
            if name not in known:
                options.append({'code': '', 'name': name})
                known.add(name)
    return options


def driver_code(name, options=None):
    if not name:
        return ''
    return next((o['code'] for o in (options or driver_options()) if o['name'] == name and o['code']), '')


def multi(params, key):
    """Олон сонголттой шүүлтийн утгууд (?plate=a&plate=b), хоосныг хасна."""
    values = params.getlist(key) if hasattr(params, 'getlist') else [params.get(key, '')]
    return [v.strip() for v in values if v and v.strip()]


def plate_filter_label(plate):
    return {OUR_PLATES: 'Манай машинууд', PERSONAL: 'Хувийн машин', NO_PLATE: 'Машингүй'}.get(plate, plate)


# --- Форматлах ---

def fmt_value(value, fmt):
    if value is None or value == '':
        return ''
    if isinstance(value, str):  # "Нийт" г.м. шошго
        return value
    if fmt == 'money':
        return f'{value:,.0f}' if value == int(value) else f'{value:,.2f}'
    if fmt == 'int':
        return f'{value:,.0f}'
    if fmt == 'dec':
        return f'{value:,.2f}'
    if fmt == 'date':
        return f'{value:%Y.%m.%d}'
    if fmt == 'datetime':
        return f'{timezone.localtime(value):%Y.%m.%d %H:%M}' if isinstance(value, datetime) else str(value)
    return str(value)


def render_table(table):
    """Хуудсанд зориулж нүд бүрийг форматлана: table['cells'], table['total_cells']."""
    cols = table['columns']
    table['cells'] = [[{'v': fmt_value(row.get(c['key']), c.get('fmt', 'text')),
                        'num': c.get('fmt', 'text') in ('money', 'int', 'dec'),
                        'neg': isinstance(row.get(c['key']), (int, float, Decimal)) and row.get(c['key']) < 0}
                       for c in cols] for row in table['rows']]
    if table.get('totals'):
        table['total_cells'] = [{'v': fmt_value(table['totals'].get(c['key']), c.get('fmt', 'text')),
                                 'num': c.get('fmt', 'text') in ('money', 'int', 'dec')} for c in cols]
    return table


def render_report(report):
    for t in report['tables']:
        render_table(t)
    return report


# --- Excel ---

def excel_response(report, filename):
    """Тайланг (нэг хуудсанд хүснэгтүүдийг дараалуулан) Excel болгож буцаана - A4 хөндлөн, өргөнд багтаана."""
    import openpyxl
    from django.http import HttpResponse
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = re.sub(r'[\\/*?:\[\]]', ' ', report.get('sheet') or report['title'])[:31]
    side = Side(style='thin', color='9CA3AF')
    border = Border(left=side, right=side, top=side, bottom=side)
    head_fill = PatternFill('solid', fgColor='E5E7EB')
    total_fill = PatternFill('solid', fgColor='F3F4F6')
    number_formats = {'money': '#,##0.##', 'int': '#,##0', 'dec': '#,##0.00', 'date': 'yyyy.mm.dd',
                      'datetime': 'yyyy.mm.dd hh:mm'}

    ws.cell(row=1, column=1, value=report['title']).font = Font(bold=True, size=14)
    r = 2
    for label, value in report.get('meta', []):
        ws.cell(row=r, column=1, value=f'{label}: {value}').font = Font(color='374151')
        r += 1
    widths = defaultdict(lambda: 8)
    for table in report['tables']:
        r += 1
        if table.get('title'):
            ws.cell(row=r, column=1, value=table['title']).font = Font(bold=True, size=12)
            r += 1
        cols = table['columns']
        for ci, c in enumerate(cols, 1):
            cell = ws.cell(row=r, column=ci, value=c['label'])
            cell.font, cell.fill, cell.border = Font(bold=True), head_fill, border
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            widths[ci] = max(widths[ci], min(len(c['label']), 18) + 2)
        if table is report['tables'][0] and len(report['tables']) == 1:
            ws.freeze_panes = ws.cell(row=r + 1, column=1)
        for row in table['rows']:
            r += 1
            for ci, c in enumerate(cols, 1):
                v = row.get(c['key'])
                fmt = c.get('fmt', 'text')
                if isinstance(v, Decimal):
                    v = float(v)
                if fmt == 'datetime' and isinstance(v, datetime):
                    v = timezone.localtime(v).replace(tzinfo=None)
                cell = ws.cell(row=r, column=ci, value=v)
                cell.border = border
                if fmt in number_formats:
                    cell.number_format = number_formats[fmt]
                widths[ci] = max(widths[ci], min(len(fmt_value(row.get(c['key']), fmt)), 45) + 2)
        if table.get('totals'):
            r += 1
            for ci, c in enumerate(cols, 1):
                v = table['totals'].get(c['key'])
                cell = ws.cell(row=r, column=ci, value=float(v) if isinstance(v, Decimal) else v)
                cell.font, cell.fill, cell.border = Font(bold=True), total_fill, border
                if c.get('fmt') in number_formats:
                    cell.number_format = number_formats[c['fmt']]
        r += 1
    for ci, w in widths.items():
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.oddFooter.center.text = '&P / &N'

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    wb.save(response)
    return response


# --- Тайлангууд ---

REPORTS = [
    ('vehicle_summary', 'Машинаар нэгтгэл', 'Машин бүрийн шатахуун, км, 100км-ийн зарцуулалт, сэлбэг засвар, нийт зардал'),
    ('vehicle_detail', 'Нэг машины тайлан', 'Машины мэдээлэл, шатахуун, сэлбэг засварын дэлгэрэнгүй'),
    ('monthly', 'Сараар (машин × сар)', 'Машин бүрийн зардлыг сараар'),
    ('driver', 'Жолоочоор (шатахуун)', 'Жолооч бүрийн шатахуун, явсан км'),
    ('card', 'Картаар (үлдэгдэл)', 'Карт бүрийн эхний үлдэгдэл, орлого, зарлага, эцсийн үлдэгдэл'),
    ('fuel_type', 'Шатахууны төрлөөр', 'Төрөл бүрийн литр, дүн, дундаж үнэ'),
    ('parts_category', 'Сэлбэг төрлөөр (төрөл × машин)', 'Сэлбэг, засварын зардлыг төрөл, машинаар'),
]
MONTHLY_METRICS = [
    ('total', 'Нийт зардал (₮)', 'money'),
    ('fuel_amount', 'Шатахуун (₮)', 'money'),
    ('liters', 'Шатахуун (л)', 'dec'),
    ('parts_amount', 'Сэлбэг, засвар (₮)', 'money'),
    ('km', 'Явсан км', 'int'),
]


def _fuel_purchases(d_from, d_to, extra=None):
    from shop.models import FuelTransaction

    start, end = aware_range(d_from, d_to)
    qs = FuelTransaction.objects.filter(txn_type=FuelTransaction.PURCHASE, txn_at__gte=start, txn_at__lt=end)
    if extra:
        qs = qs.filter(**extra)
    return list(qs.order_by('txn_at', 'id'))


def _parts(d_from, d_to, extra=None):
    from shop.models import VehiclePartRecord

    qs = VehiclePartRecord.objects.filter(date__range=(d_from, d_to))
    if extra:
        qs = qs.filter(**extra)
    return list(qs.order_by('date', 'id'))


def _fuel_group(t):
    if t.vehicle_plate:
        return t.vehicle_plate
    return '— Хувийн машин' if PERSONAL_NOTE in t.note else '— Машин сонгоогүй'


def _per100(liters, km):
    return (liters * 100 / km).quantize(CENT) if km and liters else None


def vehicle_stats(d_from, d_to):
    """{машин: {...}} - шатахуун (тоо, литр, дүн, км, зарцуулсан литр), сэлбэг (тоо, дүн)."""
    from shop.services.fuel import plate_metrics

    fuel = _fuel_purchases(d_from, d_to)
    metrics = plate_metrics({t.vehicle_plate for t in fuel})
    stats = defaultdict(lambda: {'fills': 0, 'liters': ZERO, 'fuel_amount': ZERO, 'km': 0, 'used_liters': ZERO,
                                 'parts_count': 0, 'parts_amount': ZERO})
    for t in fuel:
        s = stats[_fuel_group(t)]
        s['fills'] += 1
        s['liters'] += t.quantity or ZERO
        s['fuel_amount'] += -t.amount
        m = metrics.get(t.id, {})
        if m.get('distance') and m['distance'] > 0:
            s['km'] += m['distance']
            s['used_liters'] += m.get('liters') or ZERO
    for p in _parts(d_from, d_to):
        s = stats[p.vehicle_plate or '— Машингүй']
        s['parts_count'] += 1
        s['parts_amount'] += p.amount
    return stats


def report_vehicle_summary(d_from, d_to, params):
    stats = vehicle_stats(d_from, d_to)
    labels = {v['plate']: v for v in vehicle_options()}
    ours = [p for p, v in labels.items() if v['ours']]
    rows = []
    for plate in ours + sorted(p for p in stats if p not in ours):
        s = stats.get(plate) or stats[plate]
        total = s['fuel_amount'] + s['parts_amount']
        rows.append({
            'plate': plate, 'label': labels.get(plate, {}).get('label', ''),
            'kind': 'Манай' if plate in ours else ('' if plate.startswith('—') else 'Бусад'),
            'fills': s['fills'], 'liters': s['liters'], 'fuel_amount': s['fuel_amount'], 'km': s['km'] or None,
            'per100': _per100(s['used_liters'], s['km']),
            'parts_count': s['parts_count'], 'parts_amount': s['parts_amount'], 'total': total,
            'per_km': (total / s['km']).quantize(CENT) if s['km'] else None,
        })
    rows = [r for r in rows if r['total'] or r['plate'] in ours]
    totals = {'plate': 'Нийт'}
    for k in ('fills', 'liters', 'fuel_amount', 'km', 'parts_count', 'parts_amount', 'total'):
        totals[k] = sum((r[k] or 0 for r in rows), 0)
    return {
        'title': 'Машинаар нэгтгэл',
        'tables': [{'columns': [
            {'key': 'plate', 'label': 'Машины дугаар'}, {'key': 'label', 'label': 'Марк'},
            {'key': 'kind', 'label': 'Төрөл'},
            {'key': 'fills', 'label': 'Шатахуун (удаа)', 'fmt': 'int'}, {'key': 'liters', 'label': 'Литр', 'fmt': 'dec'},
            {'key': 'fuel_amount', 'label': 'Шатахуун ₮', 'fmt': 'money'}, {'key': 'km', 'label': 'Явсан км', 'fmt': 'int'},
            {'key': 'per100', 'label': '100км-д (л)', 'fmt': 'dec'},
            {'key': 'parts_count', 'label': 'Сэлбэг (мөр)', 'fmt': 'int'},
            {'key': 'parts_amount', 'label': 'Сэлбэг, засвар ₮', 'fmt': 'money'},
            {'key': 'total', 'label': 'Нийт зардал ₮', 'fmt': 'money'}, {'key': 'per_km', 'label': '1 км-ийн зардал ₮', 'fmt': 'money'},
        ], 'rows': rows, 'totals': totals}],
        'notes': ['Явсан км - км заалттай цэнэглэлтүүдийн зөрүү; 100км-д - өмнө хийсэн шатахуун ÷ явсан км × 100.'],
    }


def report_vehicle_detail(d_from, d_to, params):
    from shop.models import Vehicle
    from shop.services.fuel import plate_metrics

    plate = params.get('plate', '')
    if not plate or plate.startswith('__'):
        return {'title': 'Нэг машины тайлан', 'tables': [], 'need_plate': True}
    v = Vehicle.objects.filter(plate=plate).first()
    label = next((o['label'] for o in vehicle_options() if o['plate'] == plate), '')
    meta = [('Машин', f'{plate} {label}'.strip())]
    if v:
        info = [('Марк, модел', f'{v.make} {v.model}'.strip()), ('Арлын дугаар', v.vin), ('Үйлдвэрлэсэн он', v.year),
                ('Монголд орж ирсэн', v.imported_year), ('Хөдөлгүүр', v.engine_cc), ('Даац (кг)', v.capacity_kg),
                ('Зориулалт', v.purpose)]
        meta += [(k, val) for k, val in info if val]

    fuel = _fuel_purchases(d_from, d_to, {'vehicle_plate': plate})
    metrics = plate_metrics({plate})
    parts = _parts(d_from, d_to, {'vehicle_plate': plate})
    fuel_rows = []
    km = 0
    used = ZERO
    for t in fuel:
        m = metrics.get(t.id, {})
        if m.get('distance') and m['distance'] > 0:
            km += m['distance']
            used += m.get('liters') or ZERO
        fuel_rows.append({'date': t.txn_at, 'driver': t.driver_name, 'card': f'{t.card_no} {t.holder}',
                          'station': t.merchant, 'fuel_type': t.fuel_type, 'price': t.unit_price,
                          'liters': t.quantity, 'amount': -t.amount, 'km': t.odometer_km,
                          'distance': m.get('distance'), 'per100': m.get('per100')})
    by_cat = defaultdict(lambda: {'count': 0, 'amount': ZERO, 'last': None})
    for p in parts:
        c = by_cat[p.category]
        c['count'] += 1
        c['amount'] += p.amount
        c['last'] = p.date
    fuel_amount = sum((r['amount'] for r in fuel_rows), ZERO)
    parts_amount = sum((p.amount for p in parts), ZERO)
    liters = sum((r['liters'] or ZERO for r in fuel_rows), ZERO)
    summary = [
        {'k': 'Шатахуун', 'count': len(fuel_rows), 'amount': fuel_amount, 'extra': f'{liters:,.2f} л'},
        {'k': 'Сэлбэг, засвар', 'count': len(parts), 'amount': parts_amount, 'extra': ''},
        {'k': 'Явсан км', 'count': None, 'amount': None, 'extra': f'{km:,} км' if km else ''},
        {'k': '100км-д', 'count': None, 'amount': None,
         'extra': f'{_per100(used, km)} л' if _per100(used, km) else ''},
        {'k': '1 км-ийн зардал', 'count': None,
         'amount': ((fuel_amount + parts_amount) / km).quantize(CENT) if km else None, 'extra': ''},
    ]
    return {
        'title': f'{plate} машины тайлан',
        'meta': meta,
        'tables': [
            {'title': 'Нэгтгэл', 'columns': [{'key': 'k', 'label': 'Үзүүлэлт'}, {'key': 'count', 'label': 'Тоо', 'fmt': 'int'},
                                              {'key': 'amount', 'label': 'Дүн ₮', 'fmt': 'money'}, {'key': 'extra', 'label': 'Хэмжээ'}],
             'rows': summary, 'totals': {'k': 'Нийт зардал', 'amount': fuel_amount + parts_amount}},
            {'title': 'Сэлбэг, засвар төрлөөр', 'columns': [
                {'key': 'category', 'label': 'Төрөл'}, {'key': 'count', 'label': 'Давтамж', 'fmt': 'int'},
                {'key': 'amount', 'label': 'Дүн ₮', 'fmt': 'money'}, {'key': 'last', 'label': 'Сүүлд', 'fmt': 'date'}],
             'rows': [{'category': k, **v} for k, v in sorted(by_cat.items(), key=lambda kv: -kv[1]['amount'])],
             'totals': {'category': 'Нийт', 'count': len(parts), 'amount': parts_amount}},
            {'title': 'Шатахуун', 'columns': [
                {'key': 'date', 'label': 'Огноо', 'fmt': 'datetime'}, {'key': 'driver', 'label': 'Жолооч'},
                {'key': 'card', 'label': 'Карт'}, {'key': 'station', 'label': 'ШТС'},
                {'key': 'fuel_type', 'label': 'Төрөл'}, {'key': 'price', 'label': 'Үнэ', 'fmt': 'money'},
                {'key': 'liters', 'label': 'Литр', 'fmt': 'dec'}, {'key': 'amount', 'label': 'Дүн ₮', 'fmt': 'money'},
                {'key': 'km', 'label': 'Км заалт', 'fmt': 'int'}, {'key': 'distance', 'label': 'Явсан км', 'fmt': 'int'},
                {'key': 'per100', 'label': '100км-д', 'fmt': 'dec'}],
             'rows': fuel_rows, 'totals': {'date': 'Нийт', 'liters': liters, 'amount': fuel_amount, 'distance': km or None}},
            {'title': 'Сэлбэг, засвар', 'columns': PART_COLUMNS_SHORT,
             'rows': [_part_row(p) for p in parts], 'totals': {'date': 'Нийт', 'amount': parts_amount}},
        ],
    }


PART_COLUMNS_SHORT = [
    {'key': 'date', 'label': 'Огноо', 'fmt': 'date'}, {'key': 'category', 'label': 'Төрөл'},
    {'key': 'description', 'label': 'Тайлбар'}, {'key': 'quantity', 'label': 'Тоо', 'fmt': 'dec'},
    {'key': 'unit_price', 'label': 'Нэгж үнэ', 'fmt': 'money'}, {'key': 'amount', 'label': 'Дүн ₮', 'fmt': 'money'},
    {'key': 'supplier', 'label': 'Нийлүүлэгч'}, {'key': 'driver', 'label': 'Жолооч'}, {'key': 'km', 'label': 'Км', 'fmt': 'int'},
]


def _part_row(p):
    return {'date': p.date, 'plate': p.vehicle_plate, 'category': p.category, 'description': p.description,
            'quantity': p.quantity, 'unit_price': p.unit_price, 'amount': p.amount, 'supplier': p.supplier,
            'driver': p.driver_name, 'km': p.odometer_km, 'note': p.note}


def report_monthly(d_from, d_to, params):
    metric = params.get('metric') if params.get('metric') in dict((m[0], m) for m in MONTHLY_METRICS) else 'total'
    label, fmt = next((m[1], m[2]) for m in MONTHLY_METRICS if m[0] == metric)
    months = months_between(d_from, d_to)
    grid = defaultdict(lambda: defaultdict(lambda: ZERO))
    for m in months:
        start, end = max(m, d_from), min(month_end(m), d_to)
        for plate, s in vehicle_stats(start, end).items():
            value = {'total': s['fuel_amount'] + s['parts_amount'], 'fuel_amount': s['fuel_amount'],
                     'liters': s['liters'], 'parts_amount': s['parts_amount'], 'km': s['km']}[metric]
            if value:
                grid[plate][m] += value
    ours = [v['plate'] for v in vehicle_options() if v['ours']]
    plates = ours + sorted(p for p in grid if p not in ours)
    rows = []
    for plate in plates:
        if plate not in grid and plate not in ours:
            continue
        row = {'plate': plate, **{m.isoformat(): grid[plate].get(m) for m in months}}
        row['total'] = sum(grid[plate].values(), ZERO)
        rows.append(row)
    totals = {'plate': 'Нийт', 'total': sum((r['total'] for r in rows), ZERO)}
    for m in months:
        totals[m.isoformat()] = sum((r[m.isoformat()] or ZERO for r in rows), ZERO)
    return {
        'title': f'Сараар - {label}',
        'tables': [{'columns': [{'key': 'plate', 'label': 'Машин'}]
                    + [{'key': m.isoformat(), 'label': f'{m:%Y.%m}', 'fmt': fmt} for m in months]
                    + [{'key': 'total', 'label': 'Нийт', 'fmt': fmt}], 'rows': rows, 'totals': totals}],
    }


def report_driver(d_from, d_to, params):
    from shop.services.fuel import plate_metrics

    fuel = _fuel_purchases(d_from, d_to)
    metrics = plate_metrics({t.vehicle_plate for t in fuel})
    by = defaultdict(lambda: {'fills': 0, 'liters': ZERO, 'amount': ZERO, 'km': 0, 'plates': set(), 'cards': set()})
    for t in fuel:
        d = by[t.driver_name or '— Тодорхойгүй']
        d['fills'] += 1
        d['liters'] += t.quantity or ZERO
        d['amount'] += -t.amount
        dist = metrics.get(t.id, {}).get('distance')
        if dist and dist > 0:
            d['km'] += dist
        if t.vehicle_plate:
            d['plates'].add(t.vehicle_plate)
        elif PERSONAL_NOTE in t.note:
            d['plates'].add('хувийн')
        d['cards'].add(t.holder or t.card_no)
    rows = [{'driver': k, 'fills': v['fills'], 'liters': v['liters'], 'amount': v['amount'], 'km': v['km'] or None,
             'plates': ', '.join(sorted(v['plates'])), 'cards': ', '.join(sorted(v['cards']))}
            for k, v in sorted(by.items(), key=lambda kv: -kv[1]['amount'])]
    return {
        'title': 'Жолоочоор - шатахуун',
        'tables': [{'columns': [
            {'key': 'driver', 'label': 'Жолооч'}, {'key': 'fills', 'label': 'Удаа', 'fmt': 'int'},
            {'key': 'liters', 'label': 'Литр', 'fmt': 'dec'}, {'key': 'amount', 'label': 'Дүн ₮', 'fmt': 'money'},
            {'key': 'km', 'label': 'Явсан км', 'fmt': 'int'}, {'key': 'plates', 'label': 'Машин'},
            {'key': 'cards', 'label': 'Карт эзэмшигч'}],
            'rows': rows, 'totals': {'driver': 'Нийт', 'fills': sum(r['fills'] for r in rows),
                                     'liters': sum((r['liters'] for r in rows), ZERO),
                                     'amount': sum((r['amount'] for r in rows), ZERO)}}],
    }


def report_card(d_from, d_to, params):
    from shop.services.fuel import card_movements

    data = card_movements(d_from, d_to)
    cols = [{'key': 'card', 'label': 'Карт'}, {'key': 'holder', 'label': 'Эзэмшигч'},
            {'key': 'opening', 'label': 'Эхний үлдэгдэл', 'fmt': 'money'}, {'key': 'income', 'label': 'Орлого', 'fmt': 'money'},
            {'key': 'transfer_out', 'label': 'Шилжүүлгийн зарлага', 'fmt': 'money'},
            {'key': 'purchase', 'label': 'Худалдан авалт', 'fmt': 'money'}, {'key': 'deduction', 'label': 'Хасалт', 'fmt': 'money'},
            {'key': 'closing', 'label': 'Эцсийн үлдэгдэл', 'fmt': 'money'}]
    if data['month']:
        cols += [{'key': 'report_closing', 'label': 'Тайлангийн эцсийн', 'fmt': 'money'}, {'key': 'diff', 'label': 'Зөрүү', 'fmt': 'money'},
                 {'key': 'status', 'label': 'Тулгалт'}]
    return {'title': 'Картын үлдэгдэл', 'tables': [{'columns': cols, 'rows': data['rows'], 'totals': data['totals']}],
            'notes': ['Эхний үлдэгдэл - картын хамгийн эхний үлдэгдлийн тайлангаас гүйлгээгээр тооцсон.']}


def report_fuel_type(d_from, d_to, params):
    by = defaultdict(lambda: {'fills': 0, 'liters': ZERO, 'amount': ZERO, 'prices': set()})
    for t in _fuel_purchases(d_from, d_to):
        if not t.amount:
            continue
        d = by[t.fuel_type or 'Тодорхойгүй']
        d['fills'] += 1
        d['liters'] += t.quantity or ZERO
        d['amount'] += -t.amount
        if t.unit_price:
            d['prices'].add(t.unit_price)
    rows = [{'type': k, 'fills': v['fills'], 'liters': v['liters'], 'amount': v['amount'],
             'avg': (v['amount'] / v['liters']).quantize(CENT) if v['liters'] else None,
             'min': min(v['prices']) if v['prices'] else None, 'max': max(v['prices']) if v['prices'] else None}
            for k, v in sorted(by.items(), key=lambda kv: -kv[1]['amount'])]
    return {
        'title': 'Шатахууны төрлөөр',
        'tables': [{'columns': [
            {'key': 'type', 'label': 'Төрөл'}, {'key': 'fills', 'label': 'Удаа', 'fmt': 'int'},
            {'key': 'liters', 'label': 'Литр', 'fmt': 'dec'}, {'key': 'amount', 'label': 'Дүн ₮', 'fmt': 'money'},
            {'key': 'avg', 'label': 'Дундаж үнэ', 'fmt': 'money'}, {'key': 'min', 'label': 'Бага үнэ', 'fmt': 'money'},
            {'key': 'max', 'label': 'Их үнэ', 'fmt': 'money'}],
            'rows': rows, 'totals': {'type': 'Нийт', 'fills': sum(r['fills'] for r in rows),
                                     'liters': sum((r['liters'] for r in rows), ZERO),
                                     'amount': sum((r['amount'] for r in rows), ZERO)}}],
    }


def report_parts_category(d_from, d_to, params):
    parts = _parts(d_from, d_to)
    grid = defaultdict(lambda: defaultdict(lambda: ZERO))
    for p in parts:
        grid[p.category][p.vehicle_plate or '—'] += p.amount
    plates = sorted({k for row in grid.values() for k in row}, key=lambda p: (p == '—', p))
    rows = []
    for cat in sorted(grid, key=lambda c: -sum(grid[c].values())):
        rows.append({'category': cat, **{p: grid[cat].get(p) for p in plates}, 'total': sum(grid[cat].values(), ZERO)})
    totals = {'category': 'Нийт', 'total': sum((r['total'] for r in rows), ZERO)}
    for p in plates:
        totals[p] = sum((r[p] or ZERO for r in rows), ZERO)
    return {
        'title': 'Сэлбэг, засвар - төрөл × машин',
        'tables': [{'columns': [{'key': 'category', 'label': 'Төрөл'}]
                    + [{'key': p, 'label': 'Машингүй' if p == '—' else p, 'fmt': 'money'} for p in plates]
                    + [{'key': 'total', 'label': 'Нийт', 'fmt': 'money'}], 'rows': rows, 'totals': totals}],
    }


REPORT_BUILDERS = {
    'vehicle_summary': report_vehicle_summary, 'vehicle_detail': report_vehicle_detail, 'monthly': report_monthly,
    'driver': report_driver, 'card': report_card, 'fuel_type': report_fuel_type,
    'parts_category': report_parts_category,
}


def build_report(kind, d_from, d_to, params):
    kind = kind if kind in REPORT_BUILDERS else 'vehicle_summary'
    report = REPORT_BUILDERS[kind](d_from, d_to, params)
    report['kind'] = kind
    report['meta'] = [('Хугацаа', period_label(d_from, d_to))] + [(k, str(v)) for k, v in report.get('meta', [])]
    return report


# --- Машины тайлан (шатахуун): машин бүр 5 баганатай блок - хуучин "Машины_Тайлан" Excel-ийн загвар ---

MATRIX_COLS = [('km', 'Заалт'), ('qty', 'Хэмжээ'), ('amount', 'Дүн'), ('distance', 'Явсан км'), ('per100', '100км-д')]
MATRIX_CHUNK = 5  # нэг хэсэгт багтах машины тоо (A4 хөндлөн)


def fuel_matrix(rows, d_from, d_to, chunk=MATRIX_CHUNK):
    """fuel.fuel_data()['rows'] (шүүсэн) -> машинаар блоклосон тайлан (хуучин "Машины_Тайлан" Excel-ийн хэлбэр).
    Машинтай худалдан авалт л орно. Тооцоо хуудасны хүснэгттэй ижил: Заалт = тухайн цэнэглэлтэд оруулсан заалт,
    Явсан км = өмнөх заалтаас хойш, 100км-д = өмнө хийсэн шатахуун ÷ явсан км × 100. Блокийн толгойд хугацааны
    өмнөх сүүлийн заалттай цэнэглэлтийн огноо, заалт."""
    from shop.models import FuelTransaction

    rows = [r for r in rows if r['t'].is_purchase and r['t'].vehicle_plate and r['t'].amount]
    labels = {v['plate']: v for v in vehicle_options()}
    ours = [p for p, v in labels.items() if v['ours']]
    used = {r['t'].vehicle_plate for r in rows}
    plates = [p for p in ours if p in used] + sorted(used - set(ours))
    start = aware_range(d_from, d_to)[0]
    heads = []
    for p in plates:
        prev = (FuelTransaction.objects.filter(vehicle_plate=p, txn_type=FuelTransaction.PURCHASE, txn_at__lt=start,
                                               odometer_km__isnull=False).order_by('-txn_at', '-id').first())
        heads.append({'plate': p, 'label': labels.get(p, {}).get('label', ''),
                      'prev_date': timezone.localtime(prev.txn_at).date() if prev else None,
                      'prev_km': prev.odometer_km if prev else None})

    sections = []
    grand = ZERO
    for i in range(0, len(heads), chunk):
        part = heads[i:i + chunk]
        index = {h['plate']: n for n, h in enumerate(part)}
        totals = [{'qty': ZERO, 'amount': ZERO, 'distance': 0, 'used': ZERO} for _ in part]
        lines = []
        for r in rows:
            t, m = r['t'], r['m']
            n = index.get(t.vehicle_plate)
            if n is None:
                continue
            cells = [None] * len(part)
            dist = m.get('distance')
            cells[n] = {'km': t.odometer_km, 'qty': t.quantity, 'amount': -t.amount, 'distance': dist,
                        'per100': m.get('per100')}
            tot = totals[n]
            tot['qty'] += t.quantity or ZERO
            tot['amount'] += -t.amount
            if dist and dist > 0:
                tot['distance'] += dist
                tot['used'] += m.get('liters') or ZERO
            lines.append({'date': timezone.localtime(t.txn_at).date(), 'card': t.card_no, 'holder': t.holder,
                          'driver': t.driver_name, 'cells': cells})
        for tot in totals:
            tot['per100'] = _per100(tot['used'], tot['distance'])
            grand += tot['amount']
        sections.append({'heads': part, 'lines': lines, 'totals': totals})
    # Шатахууны төрлөөр: удаа, литр, дүн, дундаж үнэ, аль машинууд
    types = defaultdict(lambda: {'count': 0, 'qty': ZERO, 'amount': ZERO, 'plates': set()})
    for r in rows:
        t = r['t']
        d = types[t.fuel_type or 'Тодорхойгүй']
        d['count'] += 1
        d['qty'] += t.quantity or ZERO
        d['amount'] += -t.amount
        d['plates'].add(t.vehicle_plate)
    by_type = [{'type': k, 'count': v['count'], 'qty': v['qty'], 'amount': v['amount'],
                'avg': (v['amount'] / v['qty']).quantize(CENT) if v['qty'] else None,
                'plates': ', '.join(sorted(v['plates']))}
               for k, v in sorted(types.items(), key=lambda kv: -kv[1]['amount'])]
    return {'sections': sections, 'grand_total': grand, 'plates': len(heads), 'count': len(rows), 'by_type': by_type,
            'total_qty': sum((d['qty'] for d in by_type), ZERO)}


def fuel_matrix_excel(matrix, meta, filename):
    """Машины тайлан (шатахуун) Excel: хэсэг бүр - машины дугаар (нэгтгэсэн нүд), өмнөх огноо/заалт, багана,
    мөрүүд, нийлбэр (SUM томьёо); эцэст нийт дүн."""
    import openpyxl
    from django.http import HttpResponse
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as col_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Машины тайлан'
    side = Side(style='thin', color='808080')
    border = Border(left=side, right=side, top=side, bottom=side)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    head_fill = PatternFill('solid', fgColor='E5E7EB')
    plate_fill = PatternFill('solid', fgColor='FEF3C7')
    total_fill = PatternFill('solid', fgColor='F3F4F6')
    fmts = {'km': '#,##0', 'qty': '#,##0.00', 'amount': '#,##0', 'distance': '#,##0', 'per100': '0.00'}
    positive = chr(34) + '>0' + chr(34)

    ws.cell(row=1, column=1, value='Машины тайлан (шатахуун)').font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=' · '.join(f'{k}: {v}' for k, v in meta)).font = Font(color='374151')
    r = 4
    amount_cells = []
    for sec in matrix['sections']:
        # 1-р мөр: машины дугаар, 2-р мөр: өмнөх огноо, заалт
        for n, h in enumerate(sec['heads']):
            c0 = 4 + n * 5
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 4)
            cell = ws.cell(row=r, column=c0, value=f"{h['plate']} {h['label']}".strip())
            cell.font, cell.alignment, cell.fill = Font(bold=True, size=12), center, plate_fill
            for c in range(c0, c0 + 5):
                ws.cell(row=r, column=c).border = border
            ws.cell(row=r + 1, column=c0, value='огноо').font = Font(color='6B7280')
            ws.cell(row=r + 1, column=c0 + 1, value=h['prev_date'].strftime('%Y.%m.%d') if h['prev_date'] else '')
            ws.cell(row=r + 1, column=c0 + 2, value='Өмнөх заалт').font = Font(color='6B7280')
            km = ws.cell(row=r + 1, column=c0 + 4, value=h['prev_km'])
            km.number_format, km.font = '#,##0', Font(bold=True)
        hr = r + 3
        for c, label in enumerate(['Огноо', 'Карт', 'Жолооч'], 1):
            cell = ws.cell(row=hr, column=c, value=label)
            cell.font, cell.alignment, cell.fill, cell.border = Font(bold=True), center, head_fill, border
        for n in range(len(sec['heads'])):
            for k, (_, label) in enumerate(MATRIX_COLS):
                cell = ws.cell(row=hr, column=4 + n * 5 + k, value=label)
                cell.font, cell.alignment, cell.fill, cell.border = Font(bold=True), center, head_fill, border
        first = hr + 1
        r = hr
        for line in sec['lines']:
            r += 1
            for c, v in enumerate([line['date'].strftime('%Y.%m.%d'), line['card'], line['driver']], 1):
                ws.cell(row=r, column=c, value=v).border = border
            for n, data in enumerate(line['cells']):
                for k, (key, _) in enumerate(MATRIX_COLS):
                    v = data.get(key) if data else None
                    cell = ws.cell(row=r, column=4 + n * 5 + k, value=float(v) if isinstance(v, Decimal) else v)
                    cell.border = border
                    cell.number_format = fmts[key]
                    if key == 'distance' and isinstance(v, int) and v < 0:
                        cell.font = Font(color='DC2626')
        last = r
        r += 1
        ws.cell(row=r, column=1, value='Нийт').font = Font(bold=True)
        for c in range(1, 4):
            ws.cell(row=r, column=c).fill = total_fill
            ws.cell(row=r, column=c).border = border
        for n, tot in enumerate(sec['totals']):
            for k, (key, _) in enumerate(MATRIX_COLS):
                col = 4 + n * 5 + k
                rng = f'{col_letter(col)}{first}:{col_letter(col)}{last}'
                if key in ('qty', 'amount') and last >= first:
                    value = f'=SUM({rng})'
                elif key == 'distance' and last >= first:
                    value = f'=SUMIF({rng},{positive})'
                elif key == 'per100':
                    value = float(tot['per100']) if tot['per100'] else None
                else:
                    value = None
                cell = ws.cell(row=r, column=col, value=value)
                cell.font, cell.fill, cell.border, cell.number_format = Font(bold=True), total_fill, border, fmts[key]
                if key == 'amount':
                    amount_cells.append(f'{col_letter(col)}{r}')
        r += 3
    if amount_cells:
        ws.cell(row=r, column=1, value='Нийт дүн').font = Font(bold=True, size=12)
        cell = ws.cell(row=r, column=3, value='=' + '+'.join(amount_cells))
        cell.font, cell.number_format = Font(bold=True, size=12), '#,##0'
    if matrix.get('by_type'):
        r += 2
        ws.cell(row=r, column=1, value='Шатахууны төрлөөр').font = Font(bold=True, size=12)
        r += 1
        for c, label in enumerate(['Төрөл', 'Удаа', 'Литр', 'Дүн', 'Дундаж үнэ', 'Машин'], 1):
            cell = ws.cell(row=r, column=c, value=label)
            cell.font, cell.alignment, cell.fill, cell.border = Font(bold=True), center, head_fill, border
        first = r + 1
        for d in matrix['by_type']:
            r += 1
            for c, (v, fm) in enumerate([(d['type'], None), (d['count'], '#,##0'), (float(d['qty']), '#,##0.00'),
                                         (float(d['amount']), '#,##0'), (float(d['avg']) if d['avg'] else None, '#,##0.00'),
                                         (d['plates'], None)], 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.border = border
                if fm:
                    cell.number_format = fm
        r += 1
        ws.cell(row=r, column=1, value='Нийт')
        for c, fm in ((2, '#,##0'), (3, '#,##0.00'), (4, '#,##0')):
            ws.cell(row=r, column=c, value=f'=SUM({col_letter(c)}{first}:{col_letter(c)}{r - 1})').number_format = fm
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.fill, cell.border = Font(bold=True), total_fill, border

    ws.column_dimensions['A'].width = 11
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    widths = [9, 8, 10, 8, 8]
    for col in range(4, 4 + MATRIX_CHUNK * 5):
        ws.column_dimensions[col_letter(col)].width = widths[(col - 4) % 5]
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.oddFooter.center.text = '&P / &N'

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    wb.save(response)
    return response
